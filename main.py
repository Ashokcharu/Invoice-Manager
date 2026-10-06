import sys
import os
import shutil
import sqlite3
import subprocess
import textwrap
import re
import json
import tempfile
from datetime import datetime, timedelta
from copy import copy
from urllib.parse import quote

from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP

from PySide6.QtCore import (
    QDate, QEvent, QTimer, Qt, QStandardPaths, QUrl, QRectF, QSize,
    QRegularExpression,
)
from PySide6.QtCore import QMimeData
from PySide6.QtGui import (
    QFont, QAction, QImage, QImageReader, QColor, QBrush, QDesktopServices,
    QPainter, QPen, QIcon, QPixmap, QRegularExpressionValidator,
)
from PySide6.QtWidgets import (
    QApplication,
    QMainWindow,
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QGridLayout,
    QFormLayout,
    QLabel,
    QLineEdit,
    QComboBox,
    QDoubleSpinBox,
    QAbstractSpinBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QMessageBox,
    QCompleter,
    QToolTip,
    QDateEdit,
    QHeaderView,
    QGroupBox,
    QToolButton,
    QScrollArea,
    QTabWidget,
    QFileDialog,
    QTextEdit,
    QCheckBox,
    QListWidget,
    QListWidgetItem,
    QDialog,
    QDialogButtonBox,
    QButtonGroup,
    QStackedWidget,
    QFrame,
    QStyle,
)

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.drawing.spreadsheet_drawing import (
    AnchorMarker,
    OneCellAnchor,
)
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.pagebreak import Break
import qrcode


# ============================================================
# PATHS
# ============================================================

IS_FROZEN = bool(getattr(sys, "frozen", False))
BUNDLE_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
APP_DIR = BUNDLE_DIR if IS_FROZEN else Path(__file__).resolve().parent
RESOURCE_DIR = BUNDLE_DIR

DATA_DIR = APP_DIR / "data"
INVOICE_DIR = APP_DIR / "invoices"
PRIVATE_INVOICE_DIR = INVOICE_DIR / "Private"
EXTERNAL_PDF_DIR = (
    Path(sys.executable).resolve().parent / "invoices"
    if IS_FROZEN and sys.platform == "win32" else None
)
TEMPLATE_DIR = RESOURCE_DIR / "data"

DB_PATH = DATA_DIR / "invoices.db"
SIGNATURE_PATH = DATA_DIR / "invoice_signature.png"
START_INVOICE_NUMBER = 101
START_PRIVATE_INVOICE_NUMBER = 100001
MAX_INVOICE_ITEMS = 10

DATA_DIR.mkdir(parents=True, exist_ok=True)
INVOICE_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# MONEY
# ============================================================

def money(value):
    return Decimal(str(value)).quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP
    )


def invoice_item_row_shift(sheet):
    item_count = 0
    while sheet.cell(18 + item_count, 2).value:
        item_count += 1
    if (
        item_count <= 15
        and isinstance(sheet["H33"].value, (int, float))
        and sheet["B33"].value is None
    ):
        return 5
    return max(0, item_count - 10)


def normalize_indian_mobile(value):
    digits = re.sub(r"\D", "", value or "")
    if len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    if len(digits) != 10 or digits[0] not in "6789":
        return ""
    return f"91{digits}"


def number_to_words(value):
    value = int(value)
    if value == 0:
        return "Zero"

    ones = (
        "Zero", "One", "Two", "Three", "Four", "Five", "Six", "Seven",
        "Eight", "Nine", "Ten", "Eleven", "Twelve", "Thirteen",
        "Fourteen", "Fifteen", "Sixteen", "Seventeen", "Eighteen",
        "Nineteen",
    )
    tens = (
        "", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty",
        "Seventy", "Eighty", "Ninety",
    )

    def under_thousand(number):
        words = []
        if number >= 100:
            words.extend((ones[number // 100], "Hundred"))
            number %= 100
        if number >= 20:
            words.append(tens[number // 10])
            number %= 10
        if number:
            words.append(ones[number])
        return " ".join(words)

    words = []
    for divisor, label in (
        (10**11, "Kharab"), (10**9, "Arab"), (10**7, "Crore"),
        (10**5, "Lakh"), (1000, "Thousand"), (100, "Hundred"),
    ):
        if value >= divisor:
            group, value = divmod(value, divisor)
            words.append(f"{under_thousand(group)} {label}")
    if value:
        words.append(under_thousand(value))
    return " ".join(words)


def amount_to_words(value):
    value = money(value)
    rupees = int(value)
    paise = int((value - rupees) * 100)
    words = f"Rupees {number_to_words(rupees)}"
    if paise:
        words += f" and {number_to_words(paise)} Paise"
    return f"{words} Only"


def invoice_filename(invoice_no, buyer_name, extension, document_kind="INVOICE"):
    buyer_part = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", buyer_name or "")
    buyer_part = re.sub(r"\s+", "_", buyer_part).strip(" ._")[:80]
    buyer_part = buyer_part.rstrip(" ._") or "Buyer"
    prefix = "QTN" if document_kind == "QUOTATION" else "INV"
    return f"{prefix}-{invoice_no:04d}_{buyer_part}{extension}"


def document_kind_for_selection(index):
    return "INVOICE" if index == 0 else "QUOTATION"


def document_kind_selection_for_kind(document_kind):
    return 1 if document_kind == "QUOTATION" else 0


def document_folder_for_selection(index):
    return "Invoices" if index == 0 else "Quotations"


class NoWheelDoubleSpinBox(QDoubleSpinBox):

    def wheelEvent(self, event):
        event.ignore()


class AmountEntryDoubleSpinBox(NoWheelDoubleSpinBox):

    def __init__(self):
        super().__init__()
        self.lineEdit().installEventFilter(self)

    def focusInEvent(self, event):
        super().focusInEvent(event)
        self.selectAll()

    def eventFilter(self, watched, event):
        if (
            watched is self.lineEdit()
            and event.type() == QEvent.Type.MouseButtonPress
        ):
            QTimer.singleShot(0, self.selectAll)
        return super().eventFilter(watched, event)


class NoWheelComboBox(QComboBox):

    def wheelEvent(self, event):
        event.ignore()


class NoWheelDateEdit(QDateEdit):

    def wheelEvent(self, event):
        event.ignore()


class FinanceDonutChart(QWidget):

    def __init__(self, parent=None):
        super().__init__(parent)
        self.collected = 0.0
        self.outstanding = 0.0
        self.setMinimumSize(220, 190)

    def set_values(self, collected, outstanding):
        self.collected = max(0.0, collected)
        self.outstanding = max(0.0, outstanding)
        self.update()

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        side = min(self.width() - 24, self.height() - 24)
        chart = QRectF(
            (self.width() - side) / 2,
            (self.height() - side) / 2,
            side,
            side,
        )
        total = self.collected + self.outstanding
        if total <= 0:
            painter.setPen(QPen(QColor("#d8e1e8"), 18))
            painter.drawArc(chart, 0, 360 * 16)
            return

        collected_angle = round(self.collected / total * 360 * 16)
        painter.setPen(QPen(QColor("#21866b"), 18))
        painter.drawArc(chart, 90 * 16, -collected_angle)
        painter.setPen(QPen(QColor("#e7a34d"), 18))
        painter.drawArc(chart, 90 * 16 - collected_angle, -(360 * 16 - collected_angle))
        painter.setPen(QColor("#253b4a"))
        painter.setFont(QFont("Segoe UI", 10, QFont.Weight.DemiBold))
        painter.drawText(chart, Qt.AlignmentFlag.AlignCenter, "INVOICED")


class FinanceMonthlyChart(QWidget):

    MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")

    def __init__(self, parent=None):
        super().__init__(parent)
        self.collected_values = [0.0] * 12
        self.outstanding_values = [0.0] * 12
        self.setMinimumHeight(270)
        self.setMouseTracking(True)

    def set_values(self, collected, outstanding):
        self.collected_values = list(collected[:12]) + [0.0] * max(0, 12 - len(collected))
        self.outstanding_values = list(outstanding[:12]) + [0.0] * max(0, 12 - len(outstanding))
        self.update()

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        left, right, top, bottom = 48, 12, 34, 34
        chart = self.rect().adjusted(left, top, -right, -bottom)
        max_value = max(
            max(self.collected_values, default=0),
            max(self.outstanding_values, default=0),
        )
        max_value = max_value or 1
        painter.setPen(QPen(QColor("#dce4e9"), 1))
        for fraction in (0, 0.5, 1):
            y = chart.bottom() - chart.height() * fraction
            painter.drawLine(chart.left(), round(y), chart.right(), round(y))
            value = max_value * fraction
            label = f"{value / 1000:.0f}k" if value >= 1000 else f"{value:.0f}"
            painter.setPen(QColor("#74828a"))
            painter.setFont(QFont("Segoe UI", 7))
            painter.drawText(
                QRectF(0, y - 8, left - 7, 16),
                Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter,
                label,
            )
            painter.setPen(QPen(QColor("#dce4e9"), 1))

        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(QColor("#21866b"))
        painter.drawRoundedRect(QRectF(chart.left(), 8, 9, 9), 2, 2)
        painter.setPen(QColor("#52646d"))
        painter.setFont(QFont("Segoe UI", 8))
        painter.drawText(QRectF(chart.left() + 13, 3, 92, 20), "Collected")
        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(QColor("#e7a34d"))
        painter.drawRoundedRect(QRectF(chart.left() + 100, 8, 9, 9), 2, 2)
        painter.setPen(QColor("#52646d"))
        painter.drawText(QRectF(chart.left() + 113, 3, 110, 20), "Outstanding")

        slot = chart.width() / 12
        bar_width = max(3, min(15, slot * 0.29))
        for index, (collected, outstanding) in enumerate(zip(
            self.collected_values, self.outstanding_values
        )):
            for offset, value, color in (
                (-bar_width / 2 - 1, collected, "#21866b"),
                (bar_width / 2 + 1, outstanding, "#e7a34d"),
            ):
                bar_height = chart.height() * value / max_value
                bar = QRectF(
                    chart.left() + slot * index + slot / 2 + offset,
                    chart.bottom() - bar_height,
                    bar_width,
                    bar_height,
                )
                painter.setPen(Qt.PenStyle.NoPen)
                painter.setBrush(QColor(color if value else "#e3e9ed"))
                painter.drawRoundedRect(bar, 2, 2)

            painter.setPen(QColor("#667783"))
            painter.setFont(QFont("Segoe UI", 8))
            painter.drawText(
                QRectF(chart.left() + slot * index, chart.bottom() + 7, slot, 20),
                Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignTop,
                self.MONTHS[index],
            )

    def mouseMoveEvent(self, event):
        chart = self.rect().adjusted(48, 34, -12, -34)
        position = event.position().toPoint()
        if not chart.contains(position):
            QToolTip.hideText()
            return

        slot = chart.width() / 12
        month_index = min(11, int((position.x() - chart.left()) / slot))
        collected = self.collected_values[month_index]
        outstanding = self.outstanding_values[month_index]
        details = [
            f"{self.MONTHS[month_index]}",
            f"Collected to date on invoices issued: ₹{collected:,.2f}",
            f"Outstanding on invoices issued: ₹{outstanding:,.2f}",
        ]
        if month_index and self.collected_values[month_index - 1] > 0:
            previous = self.collected_values[month_index - 1]
            growth = (collected - previous) / previous * 100
            details.append(f"Collected-to-date change vs prior invoice month: {growth:+.1f}%")
        QToolTip.showText(event.globalPosition().toPoint(), "\n".join(details), self)


# ============================================================
# APPLICATION
# ============================================================

class InvoiceApp(QMainWindow):

    def __init__(self):
        super().__init__()

        self.pending_private_move_paths = []
        self.private_move_cancelled = False
        self.active_document_kind = None
        self.setWindowTitle("Invoice Manager")
        window_icon_path = TEMPLATE_DIR / "icon.png"
        if not window_icon_path.is_file():
            window_icon_path = RESOURCE_DIR / "app.ico"
        self.setWindowIcon(QIcon(str(window_icon_path)))
        self.resize(1280, 820)
        self.setMinimumSize(1050, 700)

        self.init_database()

        self.windows_style()

        self.create_menu()

        self.create_ui()

        self.delete_document_draft("INVOICE")
        self.delete_document_draft("QUOTATION")
        self.new_invoice()


    # ========================================================
    # WINDOWS STYLE
    # ========================================================

    def windows_style(self):

        self.setStyleSheet("""

        QMainWindow, QWidget {
            background: #f4f7fb;
            color: #1f2937;
            font-family: "Segoe UI";
            font-size: 10pt;
        }

        QLineEdit, QTextEdit, QComboBox, QDateEdit, QDoubleSpinBox,
        QTableWidget, QListWidget {
            background: #ffffff;
            color: #1f2937;
            border: 1px solid #bdc9d8;
            border-radius: 4px;
            padding: 6px;
            selection-background-color: #dceeff;
            selection-color: #123b5d;
        }
        QScrollBar:vertical {
            background: #e8eef4;
            width: 12px;
            margin: 2px;
            border-radius: 6px;
        }
        QScrollBar::handle:vertical {
            background: #7893a8;
            min-height: 34px;
            border-radius: 6px;
        }
        QScrollBar::handle:vertical:hover {
            background: #4e718d;
        }
        QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {
            height: 0;
        }
        QScrollBar::horizontal {
            background: #e8eef4;
            height: 12px;
            margin: 2px;
            border-radius: 6px;
        }
        QScrollBar::handle:horizontal {
            background: #7893a8;
            min-width: 34px;
            border-radius: 6px;
        }
        QScrollBar::handle:horizontal:hover {
            background: #4e718d;
        }
        QScrollBar::add-line:horizontal, QScrollBar::sub-line:horizontal {
            width: 0;
        }
        QTableWidget {
            alternate-background-color: #f4f7fb;
            background-color: #ffffff;
            color: #1f2937;
        }
        QTableWidget::item {
            color: #1f2937;
        }
        QLineEdit:focus, QTextEdit:focus, QComboBox:focus, QDateEdit:focus,
        QDoubleSpinBox:focus, QTableWidget:focus, QListWidget:focus {
            border: 1px solid #0078d4;
        }
        QGroupBox {
            background: #ffffff;
            border: 1px solid #cbd6e2;
            border-radius: 6px;
            margin-top: 16px;
            padding: 14px 12px 10px;
            font-weight: 600;
        }
        QGroupBox::title {
            subcontrol-origin: margin;
            left: 12px;
            padding: 0 6px;
            color: #173b57;
        }
        QCheckBox {
            spacing: 8px;
            padding: 6px 12px;
            border: 1px solid #8aa1b3;
            border-radius: 4px;
            font-weight: 600;
        }
        QCheckBox::indicator {
            width: 22px;
            height: 22px;
        }
        QCheckBox:checked {
            background: #e8f3fb;
            border-color: #36749e;
        }
        QHeaderView::section {
            background: #eaf1f8;
            color: #173b57;
            border: 1px solid #cbd6e2;
            padding: 7px 6px;
            font-weight: bold;
        }
        QPushButton {
            background: #ffffff;
            color: #1f2937;
            border: 1px solid #aebdcd;
            border-radius: 4px;
            padding: 7px 14px;
        }
        QPushButton:hover {
            background: #edf5fc;
            border-color: #6d9fc7;
        }
        QPushButton:pressed {
            background: #dceeff;
        }
        QPushButton#primary {
            background: #21866b;
            border-color: #21866b;
            color: #ffffff;
            font-weight: 600;
        }
        QPushButton#primary:hover {
            background: #176b56;
            border-color: #176b56;
        }
        QListWidget::item {
            padding: 5px 8px;
            border-bottom: 1px solid #edf1f5;
        }
        QListWidget::item:selected {
            background: #dceeff;
            color: #123b5d;
        }
        """)


    def spinbox_with_arrows(self, spinbox, step_size=1):

        spinbox.setButtonSymbols(QAbstractSpinBox.ButtonSymbols.NoButtons)
        container = QWidget()
        container_layout = QHBoxLayout(container)
        container_layout.setContentsMargins(0, 0, 0, 0)
        container_layout.setSpacing(0)
        container_layout.addWidget(spinbox, 1)

        buttons = QWidget()
        buttons_layout = QVBoxLayout(buttons)
        buttons_layout.setContentsMargins(0, 0, 0, 0)
        buttons_layout.setSpacing(0)
        for arrow_type, step, tooltip in (
            (Qt.ArrowType.UpArrow, step_size, "Increase"),
            (Qt.ArrowType.DownArrow, -step_size, "Decrease"),
        ):
            button = QToolButton()
            button.setArrowType(arrow_type)
            button.setToolTip(tooltip)
            button.setFixedSize(22, 16)
            button.setStyleSheet(
                "QToolButton { background: #eaf1f8; color: #173b57; "
                "border: 1px solid #bdc9d8; padding: 0; }"
                "QToolButton:hover { background: #dceeff; }"
            )
            button.clicked.connect(
                lambda checked=False, amount=step: spinbox.stepBy(amount)
            )
            buttons_layout.addWidget(button)
        container_layout.addWidget(buttons)
        return container


    def create_menu(self):

        file_menu = self.menuBar().addMenu("File")
        new_action = QAction("New Invoice", self)
        new_action.triggered.connect(self.new_invoice)
        file_menu.addAction(new_action)
        save_action = QAction("Save XLSX and Download PDF", self)
        save_action.triggered.connect(
            lambda checked=False: self.save_document_bundle()
        )
        file_menu.addAction(save_action)
        file_menu.addSeparator()
        exit_action = QAction("Exit", self)
        exit_action.triggered.connect(self.close)
        file_menu.addAction(exit_action)


    # ========================================================
    # DATABASE
    # ========================================================

    def init_database(self):

        with sqlite3.connect(DB_PATH) as con:

            con.execute("""
            
                CREATE TABLE IF NOT EXISTS invoices (

                    id INTEGER PRIMARY KEY AUTOINCREMENT,

                    invoice_no INTEGER UNIQUE NOT NULL,

                    invoice_date TEXT NOT NULL,

                    invoice_type TEXT NOT NULL,

                    customer_name TEXT,

                    grand_total REAL,

                    xlsx_path TEXT,

                    pdf_path TEXT,

                    created_at
                    TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )

            """)
            columns = {
                row[1] for row in con.execute("PRAGMA table_info(invoices)")
            }
            if "document_kind" not in columns:
                con.execute(
                    "ALTER TABLE invoices ADD COLUMN document_kind TEXT NOT NULL DEFAULT 'INVOICE'"
                )
            if "qr_scanner_path" not in columns:
                con.execute(
                    "ALTER TABLE invoices ADD COLUMN qr_scanner_path TEXT"
                )
            if "payment_address" not in columns:
                con.execute(
                    "ALTER TABLE invoices ADD COLUMN payment_address TEXT NOT NULL DEFAULT ''"
                )
            if "is_private" not in columns:
                con.execute(
                    "ALTER TABLE invoices ADD COLUMN is_private INTEGER NOT NULL DEFAULT 0"
                )
                private_root = PRIVATE_INVOICE_DIR.resolve()
                for invoice_no, document_kind, xlsx_path, pdf_path in con.execute(
                    "SELECT invoice_no, document_kind, xlsx_path, pdf_path FROM invoices"
                ).fetchall():
                    if any(
                        path and Path(path).resolve().is_relative_to(private_root)
                        for path in (xlsx_path, pdf_path)
                    ):
                        con.execute(
                            "UPDATE invoices SET is_private = 1 "
                            "WHERE invoice_no = ? AND document_kind = ?",
                            (invoice_no, document_kind),
                        )
            single_number_unique_index = False
            for index in con.execute("PRAGMA index_list(invoices)"):
                if index[2]:
                    indexed_columns = [
                        column[2]
                        for column in con.execute(
                            f"PRAGMA index_info('{index[1]}')"
                        )
                    ]
                    if indexed_columns == ["invoice_no"]:
                        single_number_unique_index = True
                        break
            if single_number_unique_index:
                con.execute(
                    """
                    CREATE TABLE invoices_migrated (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        invoice_no INTEGER NOT NULL,
                        invoice_date TEXT NOT NULL,
                        invoice_type TEXT NOT NULL,
                        customer_name TEXT,
                        grand_total REAL,
                        xlsx_path TEXT,
                        pdf_path TEXT,
                        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                        document_kind TEXT NOT NULL DEFAULT 'INVOICE',
                        qr_scanner_path TEXT,
                        payment_address TEXT NOT NULL DEFAULT '',
                        is_private INTEGER NOT NULL DEFAULT 0,
                        UNIQUE(document_kind, invoice_no)
                    )
                    """
                )
                con.execute(
                    """
                    INSERT INTO invoices_migrated (
                        id, invoice_no, invoice_date, invoice_type, customer_name,
                        grand_total, xlsx_path, pdf_path, created_at, document_kind,
                                qr_scanner_path, payment_address, is_private
                    )
                    SELECT id, invoice_no, invoice_date, invoice_type, customer_name,
                           grand_total, xlsx_path, pdf_path, created_at, document_kind,
                                    qr_scanner_path, payment_address, is_private
                    FROM invoices
                    """
                )
                con.execute("DROP TABLE invoices")
                con.execute("ALTER TABLE invoices_migrated RENAME TO invoices")
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS bank_profiles (
                    name TEXT PRIMARY KEY,
                    account TEXT NOT NULL DEFAULT '',
                    ifsc TEXT NOT NULL DEFAULT '',
                    payment_address TEXT NOT NULL DEFAULT '',
                    scanner_path TEXT
                )
                """
            )
            bank_columns = {
                row[1] for row in con.execute("PRAGMA table_info(bank_profiles)")
            }
            if "scanner_path" not in bank_columns:
                con.execute("ALTER TABLE bank_profiles ADD COLUMN scanner_path TEXT")
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS qr_scanners (
                    id TEXT PRIMARY KEY,
                    name TEXT NOT NULL,
                    image_path TEXT NOT NULL
                )
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS seller_profile (
                    id INTEGER PRIMARY KEY CHECK (id = 1),
                    seller_name TEXT NOT NULL,
                    seller_business TEXT NOT NULL,
                    address_1 TEXT NOT NULL,
                    address_2 TEXT NOT NULL,
                    gstin TEXT NOT NULL,
                    state TEXT NOT NULL,
                    website_phone TEXT NOT NULL,
                    phone_numbers TEXT NOT NULL DEFAULT '',
                    email_phone TEXT NOT NULL,
                    signatory TEXT NOT NULL,
                    jurisdiction TEXT NOT NULL
                )
                """
            )
            seller_columns = {
                row[1] for row in con.execute("PRAGMA table_info(seller_profile)")
            }
            if "phone_numbers" not in seller_columns:
                con.execute(
                    "ALTER TABLE seller_profile ADD COLUMN "
                    "phone_numbers TEXT NOT NULL DEFAULT ''"
                )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS buyer_profiles (
                    name TEXT PRIMARY KEY,
                    address_1 TEXT NOT NULL DEFAULT '',
                    address_2 TEXT NOT NULL DEFAULT '',
                    mobile TEXT NOT NULL DEFAULT '',
                    gstin TEXT NOT NULL DEFAULT ''
                )
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS document_drafts (
                    document_kind TEXT PRIMARY KEY,
                    state_json TEXT NOT NULL
                )
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS invoice_payments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    document_kind TEXT NOT NULL DEFAULT 'INVOICE',
                    invoice_no INTEGER NOT NULL,
                    paid_on TEXT NOT NULL,
                    amount REAL NOT NULL CHECK (amount > 0),
                    method TEXT NOT NULL,
                    reference TEXT NOT NULL DEFAULT '',
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            con.execute(
                """
                CREATE INDEX IF NOT EXISTS invoice_payments_document
                ON invoice_payments (document_kind, invoice_no)
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS invoice_requests (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    invoice_no INTEGER NOT NULL,
                    requested_at TEXT NOT NULL,
                    request_percent REAL NOT NULL,
                    request_amount REAL NOT NULL,
                    pdf_path TEXT NOT NULL,
                    qr_path TEXT NOT NULL DEFAULT '',
                    payment_update_path TEXT NOT NULL DEFAULT '',
                    final_pdf_path TEXT NOT NULL DEFAULT '',
                    paid_total_at_request REAL
                )
                """
            )
            request_columns = {
                row[1] for row in con.execute("PRAGMA table_info(invoice_requests)")
            }
            if "paid_total_at_request" not in request_columns:
                con.execute(
                    "ALTER TABLE invoice_requests ADD COLUMN paid_total_at_request REAL"
                )
            if "payment_update_path" not in request_columns:
                con.execute(
                    "ALTER TABLE invoice_requests ADD COLUMN "
                    "payment_update_path TEXT NOT NULL DEFAULT ''"
                )
            if "final_pdf_path" not in request_columns:
                con.execute(
                    "ALTER TABLE invoice_requests ADD COLUMN "
                    "final_pdf_path TEXT NOT NULL DEFAULT ''"
                )
            con.execute(
                """
                CREATE INDEX IF NOT EXISTS invoice_requests_document
                ON invoice_requests (invoice_no, requested_at)
                """
            )


    def prune_missing_invoice_records(self):

        with sqlite3.connect(DB_PATH) as con:
            records = con.execute(
                "SELECT id, invoice_no, document_kind, xlsx_path, pdf_path FROM invoices"
            ).fetchall()
            saved_request_numbers = {
                invoice_no
                for invoice_no, pdf_path in con.execute(
                    "SELECT invoice_no, pdf_path FROM invoice_requests"
                )
                if pdf_path and Path(pdf_path).is_file()
            }
            missing_invoice_numbers = [
                (record_id, invoice_no, document_kind)
                for record_id, invoice_no, document_kind, xlsx_path, pdf_path in records
                if not any(
                    path and Path(path).is_file()
                    for path in (xlsx_path, pdf_path)
                )
                and not (
                    document_kind == "INVOICE"
                    and invoice_no in saved_request_numbers
                )
            ]
            con.executemany(
                "DELETE FROM invoice_payments WHERE document_kind = ? AND invoice_no = ?",
                ((kind, invoice_no) for _, invoice_no, kind in missing_invoice_numbers),
            )
            con.executemany(
                "DELETE FROM invoice_requests WHERE invoice_no = ?",
                (
                    (invoice_no,)
                    for _, invoice_no, kind in missing_invoice_numbers
                    if kind == "INVOICE"
                ),
            )
            con.executemany(
                "DELETE FROM invoices WHERE id = ?",
                ((record_id,) for record_id, _, _ in missing_invoice_numbers),
            )


    def next_invoice_number(self, document_kind=None, is_private=None):

        self.prune_missing_invoice_records()
        if document_kind is None:
            document_kind = document_kind_for_selection(
                self.document_kind.currentIndex()
            )
        if is_private is None:
            is_private = (
                document_kind == "INVOICE"
                and self.private_invoice_checkbox.isChecked()
            )

        with sqlite3.connect(DB_PATH) as con:
            numbers = con.execute(
                "SELECT invoice_no, is_private FROM invoices "
                "WHERE document_kind = ?",
                (document_kind,),
            ).fetchall()

        used_numbers = {number for number, _ in numbers}
        series_numbers = [
            number for number, saved_private in numbers
            if bool(saved_private) == bool(is_private)
        ]
        if document_kind == "INVOICE" and is_private:
            manual_private_numbers = [
                number for number in series_numbers
                if number < START_PRIVATE_INVOICE_NUMBER
            ]
            series_numbers = manual_private_numbers
        next_number = (
            max(series_numbers) + 1
            if series_numbers else START_INVOICE_NUMBER
        )
        while next_number in used_numbers:
            next_number += 1
        return next_number


    def saved_invoice_file(self, invoice_no, document_kind=None):

        if document_kind is None:
            document_kind = document_kind_for_selection(
                self.document_kind.currentIndex()
            )

        with sqlite3.connect(DB_PATH) as con:
            row = con.execute(
                """SELECT xlsx_path, customer_name FROM invoices
                   WHERE invoice_no = ? AND document_kind = ?""",
                (invoice_no, document_kind),
            ).fetchone()

        if not row or not row[0]:
            return None, None

        return Path(row[0]), row[1]


    def seller_profile_fields(self):

        return (
            self.seller_name,
            self.seller_business,
            self.seller_address_1,
            self.seller_address_2,
            self.seller_gstin,
            self.seller_state,
            self.seller_website,
            self.seller_phone,
            self.seller_email_phone,
            self.seller_signatory,
            self.jurisdiction,
        )


    def load_seller_profile(self):

        with sqlite3.connect(DB_PATH) as con:
            profile = con.execute(
                """
                SELECT seller_name, seller_business, address_1, address_2,
                       gstin, state, website_phone, phone_numbers, email_phone,
                       signatory, jurisdiction
                FROM seller_profile WHERE id = 1
                """
            ).fetchone()
        if not profile:
            return

        profile = list(profile)
        legacy_website = profile[6]
        if not profile[7] and re.search(r"\s+PHONE\s*[:\-]*\s*", legacy_website, re.IGNORECASE):
            website_parts = re.split(
                r"\s+PHONE\s*[:\-]*\s*",
                legacy_website,
                maxsplit=1,
                flags=re.IGNORECASE,
            )
            profile[6] = website_parts[0].strip()
            profile[7] = website_parts[1].strip() if len(website_parts) > 1 else ""
        for field, value in zip(self.seller_profile_fields(), profile):
            field.setText(value)


    def save_seller_profile(self):

        with sqlite3.connect(DB_PATH) as con:
            con.execute(
                """
                INSERT INTO seller_profile (
                    id, seller_name, seller_business, address_1, address_2,
                    gstin, state, website_phone, phone_numbers, email_phone, signatory,
                    jurisdiction
                ) VALUES (1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(id) DO UPDATE SET
                    seller_name = excluded.seller_name,
                    seller_business = excluded.seller_business,
                    address_1 = excluded.address_1,
                    address_2 = excluded.address_2,
                    gstin = excluded.gstin,
                    state = excluded.state,
                    website_phone = excluded.website_phone,
                    phone_numbers = excluded.phone_numbers,
                    email_phone = excluded.email_phone,
                    signatory = excluded.signatory,
                    jurisdiction = excluded.jurisdiction
                """,
                tuple(field.text().strip() for field in self.seller_profile_fields()),
            )


    def refresh_buyer_profiles(self, selected_name=""):

        with sqlite3.connect(DB_PATH) as con:
            profiles = con.execute(
                "SELECT name FROM buyer_profiles ORDER BY name COLLATE NOCASE"
            ).fetchall()

        self.buyer_selector.blockSignals(True)
        self.buyer_selector.clear()
        self.buyer_selector.addItem("Select saved buyer", "")
        for (name,) in profiles:
            self.buyer_selector.addItem(name, name)
        selected_index = self.buyer_selector.findData(selected_name)
        self.buyer_selector.setCurrentIndex(max(0, selected_index))
        self.buyer_selector.blockSignals(False)


    def load_buyer_profile(self, index):

        profile_name = self.buyer_selector.itemData(index)
        if not profile_name:
            return

        with sqlite3.connect(DB_PATH) as con:
            profile = con.execute(
                """
                SELECT name, address_1, address_2, mobile, gstin
                FROM buyer_profiles WHERE name = ?
                """,
                (profile_name,),
            ).fetchone()
        if not profile:
            return

        for field, value in zip(
            (
                self.customer_name,
                self.customer_address,
                self.customer_address_2,
                self.customer_mobile,
                self.customer_gstin,
            ),
            profile,
        ):
            field.setText(value)


    def save_buyer_profile(self):

        name = self.customer_name.text().strip()
        if not name:
            QMessageBox.warning(self, "Buyer Name Required", "Enter a buyer name first.")
            return
        mobile = self.customer_mobile.text().strip()
        if mobile and not normalize_indian_mobile(mobile):
            QMessageBox.warning(
                self,
                "Invalid Buyer Mobile",
                "Enter a valid Indian mobile number with 10 digits, optionally prefixed by +91.",
            )
            self.customer_mobile.setFocus()
            return

        with sqlite3.connect(DB_PATH) as con:
            con.execute(
                """
                INSERT INTO buyer_profiles (name, address_1, address_2, mobile, gstin)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(name) DO UPDATE SET
                    address_1 = excluded.address_1,
                    address_2 = excluded.address_2,
                    mobile = excluded.mobile,
                    gstin = excluded.gstin
                """,
                (
                    name,
                    self.customer_address.text().strip(),
                    self.customer_address_2.text().strip(),
                    mobile,
                    self.customer_gstin.text().strip(),
                ),
            )
        self.refresh_buyer_profiles(name)
        self.statusBar().showMessage(f"Buyer details saved: {name}")


    def delete_buyer_profile(self):

        profile_name = self.buyer_selector.currentData()
        if not profile_name:
            QMessageBox.information(
                self, "Select a Saved Buyer", "Choose a saved buyer to delete."
            )
            return

        answer = QMessageBox.question(
            self,
            "Delete Buyer Details",
            f"Delete the saved buyer '{profile_name}'?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return

        with sqlite3.connect(DB_PATH) as con:
            con.execute("DELETE FROM buyer_profiles WHERE name = ?", (profile_name,))
        self.refresh_buyer_profiles()


    def new_buyer_profile(self):

        self.buyer_selector.setCurrentIndex(0)
        for field in (
            self.customer_name,
            self.customer_address,
            self.customer_address_2,
            self.customer_mobile,
            self.customer_gstin,
        ):
            field.clear()


    def refresh_bank_profiles(self, selected_name=None):

        if selected_name is None:
            selected_name = self.bank_name.text().strip()
        with sqlite3.connect(DB_PATH) as con:
            profiles = con.execute(
                "SELECT name FROM bank_profiles ORDER BY name COLLATE NOCASE"
            ).fetchall()

        self.bank_selector.blockSignals(True)
        self.bank_selector.clear()
        self.bank_selector.addItem("Manual / not saved", "")
        for (name,) in profiles:
            self.bank_selector.addItem(name, name)
        selected_index = self.bank_selector.findData(selected_name)
        self.bank_selector.setCurrentIndex(max(0, selected_index))
        self.bank_selector.blockSignals(False)
        if selected_index > 0:
            self.load_bank_profile(selected_index)


    def load_bank_profile(self, index):

        profile_name = self.bank_selector.itemData(index)
        if not profile_name:
            return

        with sqlite3.connect(DB_PATH) as con:
            profile = con.execute(
                """
                SELECT name, account, ifsc, payment_address
                FROM bank_profiles WHERE name = ?
                """,
                (profile_name,),
            ).fetchone()
        if not profile:
            return

        self.bank_name.setText(profile[0])
        self.bank_account.setText(profile[1])
        self.bank_ifsc.setText(profile[2])
        self.bank_upi_id.setText(profile[3])


    def save_bank_profile(self):

        name = self.bank_name.text().strip()
        if not name:
            QMessageBox.warning(self, "Bank Name Required", "Enter a bank name first.")
            return

        with sqlite3.connect(DB_PATH) as con:
            con.execute(
                """
                INSERT INTO bank_profiles (name, account, ifsc, payment_address)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(name) DO UPDATE SET
                    account = excluded.account,
                    ifsc = excluded.ifsc,
                    payment_address = excluded.payment_address
                """,
                (
                    name,
                    self.bank_account.text().strip(),
                    self.bank_ifsc.text().strip(),
                    self.bank_upi_id.text().strip(),
                ),
            )
        self.refresh_bank_profiles(name)
        self.statusBar().showMessage(f"Bank profile saved: {name}")


    def new_bank_profile(self):

        self.bank_selector.setCurrentIndex(0)
        self.bank_name.clear()
        self.bank_account.clear()
        self.bank_ifsc.clear()
        self.bank_upi_id.clear()


    def delete_bank_profile(self):

        profile_name = self.bank_selector.currentData()
        if not profile_name:
            QMessageBox.information(
                self, "Select a Saved Bank", "Choose a saved bank profile to delete."
            )
            return

        answer = QMessageBox.question(
            self,
            "Delete Bank Profile",
            f"Delete the saved bank profile '{profile_name}'?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return

        with sqlite3.connect(DB_PATH) as con:
            con.execute("DELETE FROM bank_profiles WHERE name = ?", (profile_name,))
        self.refresh_bank_profiles("")
        self.new_bank_profile()


    def select_bank_profile(self, name):

        profile_index = self.bank_selector.findData(name)
        self.bank_selector.setCurrentIndex(max(0, profile_index))


    def payment_uri(self, invoice_no, amount=None):

        uri = (
            "upi://pay?"
            f"pa={quote(self.bank_upi_id.text().strip())}"
            f"&pn={quote(self.seller_name.text())}"
            "&cu=INR"
            f"&tn={quote(str(invoice_no))}"
        )
        if amount is not None:
            uri += f"&am={money(amount):.2f}"
        return uri


    def invoice_storage_root(self, document_kind, invoice_no, grand_total=None):

        return (
            self.custom_output_dir
            or (
                PRIVATE_INVOICE_DIR
                if self.private_invoice_checkbox.isChecked()
                else INVOICE_DIR
            )
        )


    def invoice_pdf_folder(self, document_kind, invoice_no, invoice_date):

        document_root = self.invoice_storage_root(document_kind, invoice_no)
        document_folder = document_folder_for_selection(
            document_kind_selection_for_kind(document_kind)
        )
        return (
            document_root / document_folder / "PDF" / str(invoice_date.year)
        )


    def external_final_pdf_path(
        self, document_kind, invoice_no, invoice_date, filename
    ):

        if EXTERNAL_PDF_DIR is None:
            return None
        if document_kind == "QUOTATION":
            series_folder = "Quotation"
        elif self.private_invoice_checkbox.isChecked():
            series_folder = "Private"
        else:
            series_folder = "Invoices"
        return (
            EXTERNAL_PDF_DIR / f"{series_folder}_PDF_{invoice_date.year}"
            / invoice_date.strftime("%B") / filename
        )


    def mirror_external_final_pdf(
        self, source_pdf, document_kind, invoice_no, invoice_date
    ):

        destination = self.external_final_pdf_path(
            document_kind, invoice_no, invoice_date, Path(source_pdf).name
        )
        if destination is None:
            return
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source_pdf, destination)
        prefix = "QTN" if document_kind == "QUOTATION" else "INV"
        filename_pattern = f"{prefix}-{invoice_no:04d}_*.pdf"
        for old_pdf in EXTERNAL_PDF_DIR.rglob(filename_pattern):
            if old_pdf.resolve() != destination.resolve():
                old_pdf.unlink(missing_ok=True)


    def quotation_payment_qr_path(self, invoice_no, invoice_date, grand_total=None):

        document_kind = document_kind_for_selection(
            self.document_kind.currentIndex()
        )
        document_root = self.invoice_storage_root(
            document_kind, invoice_no, grand_total
        )
        return (
            document_root
            / document_folder_for_selection(self.document_kind.currentIndex())
            / "QR"
            / str(invoice_date.year)
            / invoice_filename(
                invoice_no,
                self.customer_name.text(),
                "_payment.png",
                document_kind_for_selection(
                    self.document_kind.currentIndex()
                ),
            )
        )


    def load_invoice(self):

        try:
            invoice_no = int(self.invoice_no.text())
        except ValueError:
            return

        with sqlite3.connect(DB_PATH) as con:
            row = con.execute(
                """
                SELECT invoice_date, invoice_type, xlsx_path, grand_total,
                    document_kind, qr_scanner_path, payment_address, is_private
                FROM invoices
                WHERE invoice_no = ? AND document_kind = ?
                """,
                (
                    invoice_no,
                    document_kind_for_selection(self.document_kind.currentIndex()),
                ),
            ).fetchone()

        if not row:
            self.loaded_invoice_no = None
            self.refresh_invoice_payments()
            self.refresh_invoice_request_history()
            self.statusBar().showMessage(
                f"No saved invoice found for number {invoice_no}"
            )
            return

        workbook_path = Path(row[2])
        if not workbook_path.is_file():
            QMessageBox.warning(
                self,
                "Invoice File Missing",
                f"The saved workbook could not be found:\n{workbook_path}",
            )
            return
        self.private_invoice_checkbox.blockSignals(True)
        self.private_invoice_checkbox.setChecked(bool(row[7]))
        self.private_invoice_checkbox.blockSignals(False)

        sheet = load_workbook(workbook_path, data_only=True).active
        item_row_shift = invoice_item_row_shift(sheet)

        def cell_text(address, prefix=""):
            value = sheet[address].value
            text = "" if value is None else str(value)
            if prefix and text.startswith(prefix):
                return text[len(prefix):].strip()
            return text

        def shifted_cell_text(column, row_number, prefix=""):
            return cell_text(
                f"{column}{row_number + item_row_shift}", prefix
            )

        self.loaded_invoice_no = invoice_no
        self.document_kind.setCurrentIndex(
            document_kind_selection_for_kind(row[4])
        )
        self.buyer_selector.setCurrentIndex(0)
        if row[4] in ("INVOICE", "QUOTATION"):
            payment_address = (
                row[6]
                or shifted_cell_text("G", 49, "UPI ID:")
                or cell_text("G3", "UPI ID:")
            )
            if not payment_address:
                payment_address = shifted_cell_text("A", 54, "Payee address:")
            if payment_address.startswith("UPI ID:"):
                payment_address = payment_address.removeprefix("UPI ID:").strip()
            elif payment_address.startswith("Payee address:"):
                payment_address = payment_address.removeprefix("Payee address:").strip()
            if payment_address == "Uploaded scanner selected":
                payment_address = ""
            self.bank_upi_id.setText(payment_address)
            payment_text = (
                shifted_cell_text("G", 44) or shifted_cell_text("H", 44)
                or shifted_cell_text("G", 48) or cell_text("E2")
                or shifted_cell_text("A", 53)
            )
            payment_match = re.search(
                r"(?:Advance|Scan to pay|Collect|Balance|Payment Request)\s+([\d.]+)%",
                payment_text,
            )
            self.payment_percent.setValue(
                float(payment_match.group(1)) if payment_match else 50
            )
            amount_match = re.search(
                r"(?:INR|₹)\s*([\d,]+(?:\.\d{1,2})?)", payment_text
            )
            if amount_match:
                self.payment_request_amount.setValue(
                    float(amount_match.group(1).replace(",", ""))
                )
        self.invoice_date.setDate(QDate.fromString(row[0], "yyyy-MM-dd"))
        self.invoice_type.setCurrentIndex(0 if row[1] == "IGST" else 1)
        self.customer_name.setText(cell_text("A11"))
        customer_address = sheet["A12"].value
        if not customer_address:
            legacy_address = sheet["A14"].value
            legacy_address_text = str(legacy_address or "").strip().lower()
            if legacy_address_text and not legacy_address_text.startswith(
                ("mobile:", "gst no:")
            ):
                customer_address = legacy_address
        address_lines = str(customer_address or "").splitlines()
        self.customer_address.setText(address_lines[0] if address_lines else "")
        self.customer_address_2.setText("\n".join(address_lines[1:]))
        self.po_number.setText(cell_text("G14") or cell_text("F13"))
        def prefixed_cell_text(address, prefix):
            value = cell_text(address)
            return (
                value[len(prefix):].strip()
                if value.lower().startswith(prefix.lower()) else ""
            )

        mobile = (
            prefixed_cell_text("A13", "Mobile:")
            or prefixed_cell_text("A14", "Mobile:")
        )
        gstin = (
            prefixed_cell_text("A14", "Gst no:")
            or prefixed_cell_text("A15", "Gst no:")
        )
        legacy_contact = cell_text("A13")
        if legacy_contact:
            mobile_match = re.search(r"Mobile\s*:\s*(.*)$", legacy_contact, re.IGNORECASE)
            gstin_match = re.search(
                r"GSTIN\s*:\s*(.*?)(?:\s+Mobile\s*:|$)",
                legacy_contact,
                re.IGNORECASE,
            )
            if mobile_match:
                mobile = mobile or mobile_match.group(1).strip()
            if gstin_match:
                gstin = gstin or gstin_match.group(1).strip()
        self.customer_mobile.setText(mobile)
        self.customer_gstin.setText(gstin)
        if row[4] in ("INVOICE", "QUOTATION"):
            bank_label_column = "C" if row[4] == "INVOICE" else "D"
            bank_value_columns = (
                ("D", "F") if row[4] == "INVOICE" else ("F",)
            )
            for field, address, prefix in (
                (self.bank_name, 44, "Bank Name:"),
                (self.bank_account, 45, "A/C No:"),
                (self.bank_ifsc, 46, "IFSC:"),
            ):
                value = next(
                    (
                        shifted_cell_text(column, address)
                        for column in bank_value_columns
                        if shifted_cell_text(column, address)
                    ),
                    shifted_cell_text(bank_label_column, address, prefix),
                )
                if value:
                    field.setText(value)
        column_a_values = [
            str(row[0].value or "").strip()
            for row in sheet.iter_rows(min_col=1, max_col=1)
        ]
        for field, label in (
            (self.bundle_count, "NO OF BUNDLE'S"),
            (self.through, "THROUGH"),
        ):
            value = next(
                (
                    text
                    for text in column_a_values
                    if text.casefold().startswith(label.casefold())
                ),
                "",
            )
            if value:
                field.setText(value.split(":", 1)[-1].strip())
        self.select_bank_profile(self.bank_name.text())
        if row[6]:
            self.bank_upi_id.setText(row[6])

        declaration_row = next(
            (
                index
                for index, value in enumerate(column_a_values, start=1)
                if value.casefold() == "declaration"
            ),
            None,
        )
        if declaration_row:
            declaration = sheet.cell(declaration_row + 1, 1).value
            if declaration:
                self.declaration.setPlainText(str(declaration).strip())

        tax_row = next(
            (
                candidate + item_row_shift
                for candidate in (38, 37)
                if isinstance(
                    sheet[f"D{candidate + item_row_shift}"].value,
                    (int, float),
                )
            ),
            None,
        )
        rate_cell = sheet[f"D{tax_row}"].value if tax_row else None
        if rate_cell is None:
            tax_label = str(sheet[f"B{29 + item_row_shift}"].value or "")
            tax_match = re.search(r"(\d+(?:\.\d+)?)\s*%", tax_label)
            subtotal_value = sheet[f"H{28 + item_row_shift}"].value or 0
            if tax_match:
                gst_rate = float(tax_match.group(1))
            elif subtotal_value:
                multiplier = 100 if row[1] == "IGST" else 200
                gst_rate = max(0, (float(row[3]) / float(subtotal_value) - 1) * multiplier)
            else:
                gst_rate = self.gst_rate.value()
        else:
            gst_rate = float(rate_cell) * (100 if row[1] == "IGST" else 200)
        self.gst_rate.setValue(gst_rate)
        hsn_code = (
            (sheet[f"A{tax_row}"].value if tax_row else None)
            or sheet[f"A{38 + item_row_shift}"].value
            or sheet[f"A{37 + item_row_shift}"].value
        )
        if hsn_code:
            self.hsn_code.setText(str(hsn_code))

        self.items.setRowCount(0)
        for source_row in range(18, 28 + item_row_shift):
            description = sheet.cell(source_row, 2).value
            if not description:
                continue

            self.add_item_row()
            target_row = self.items.rowCount() - 1
            self.items.blockSignals(True)
            for column, source_column in ((1, 2), (2, 4), (3, 5), (4, 6), (5, 7)):
                value = sheet.cell(source_row, source_column).value
                self.items.item(target_row, column).setText(
                    "PCS" if column == 5 and not str(value or "").strip()
                    else "" if value is None else str(value)
                )
            self.items.blockSignals(False)

        self.refresh_item_row_controls()
        self.refresh()
        self.refresh_invoice_payments()
        self.refresh_invoice_request_history()
        self.statusBar().showMessage(f"Invoice {invoice_no} loaded for editing")


    # ========================================================
    # UI
    # ========================================================

    def sidebar_action_icon(self, kind):
        pixmap = QPixmap(24, 24)
        pixmap.fill(Qt.GlobalColor.transparent)
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        if kind == "plus":
            painter.setPen(QPen(QColor("#21866b"), 2.4))
            painter.drawLine(12, 5, 12, 19)
            painter.drawLine(5, 12, 19, 12)
        else:
            painter.setPen(QPen(QColor("#36515d"), 1.7))
            painter.setBrush(Qt.BrushStyle.NoBrush)
            painter.drawRoundedRect(QRectF(3.5, 2, 15, 20), 2, 2)

        if kind == "plus":
            pass
        elif kind == "create":
            painter.drawLine(7, 8, 14, 8)
            painter.drawLine(7, 11, 13, 11)
            painter.setPen(QPen(QColor("#21866b"), 2.2))
            painter.drawLine(18, 15, 18, 22)
            painter.drawLine(14.5, 18.5, 21.5, 18.5)
        elif kind == "finance":
            painter.setPen(QColor("#21866b"))
            painter.setFont(QFont("Segoe UI", 12, QFont.Weight.Bold))
            painter.drawText(
                QRectF(3, 1, 18, 22),
                Qt.AlignmentFlag.AlignCenter,
                "₹",
            )
        else:
            painter.setPen(QColor("#21866b"))
            painter.setFont(QFont("Segoe UI", 12, QFont.Weight.Bold))
            painter.drawText(
                QRectF(5, 2, 13, 13),
                Qt.AlignmentFlag.AlignCenter,
                '"',
            )
            painter.setPen(QPen(QColor("#36515d"), 1.5))
            painter.drawLine(7, 15, 15, 15)
            painter.drawLine(7, 18, 15, 18)

        painter.end()
        return QIcon(pixmap)


    def create_ui(self):
        central = QWidget()
        root_layout = QHBoxLayout(central)
        root_layout.setContentsMargins(0, 0, 0, 0)
        root_layout.setSpacing(0)

        sidebar = QWidget()
        sidebar.setFixedWidth(196)
        sidebar.setStyleSheet("background: #e9eff2;")
        sidebar_layout = QVBoxLayout(sidebar)
        sidebar_layout.setContentsMargins(14, 20, 14, 16)
        sidebar_layout.setSpacing(10)
        brand = QLabel("ADOODS\nINVOICE MANAGER")
        brand.setFont(QFont("Segoe UI", 12, QFont.Weight.Bold))
        brand.setStyleSheet("color: #193f4c; padding: 8px 6px 18px;")
        sidebar_layout.addWidget(brand)

        self.sidebar_navigation = QListWidget()
        navigation_items = (
            ("Create", self.sidebar_action_icon("plus")),
            ("invoices", self.sidebar_action_icon("create")),
            ("Quotations", self.sidebar_action_icon("quotation")),
            ("Finance", self.sidebar_action_icon("finance")),
        )
        for label, icon in navigation_items:
            item = QListWidgetItem(icon, label)
            self.sidebar_navigation.addItem(item)
        self.sidebar_navigation.setSpacing(4)
        self.sidebar_navigation.setIconSize(QSize(18, 18))
        self.sidebar_navigation.setCurrentRow(0)
        self.sidebar_navigation.setStyleSheet(
            "QListWidget { border: 0; background: transparent; padding: 2px; }"
            "QListWidget::item { padding: 12px 10px; border-radius: 4px; "
            "color: #36515d; font-weight: 600; }"
            "QListWidget::item:selected { background: #d0e6dd; color: #155b46; }"
            "QListWidget::item:hover { background: #dce7e8; }"
        )
        sidebar_layout.addWidget(self.sidebar_navigation, 1)
        sidebar_note = QLabel("GST invoices\nReceipts & reporting")
        sidebar_note.setStyleSheet("color: #6f8188; padding: 8px 6px;")
        sidebar_layout.addWidget(sidebar_note)
        root_layout.addWidget(sidebar)

        content = QWidget()
        content_layout = QVBoxLayout(content)
        content_layout.setContentsMargins(18, 16, 18, 16)
        content_layout.setSpacing(8)
        title = QLabel("Invoice Manager")
        title.setFont(QFont("Segoe UI", 20, QFont.Weight.DemiBold))
        content_layout.addWidget(title)
        subtitle = QLabel("Create and manage GST invoices")
        subtitle.setStyleSheet("color:#667783;")
        content_layout.addWidget(subtitle)

        self.main_tabs = QStackedWidget()
        self.invoice_tab = QWidget()
        self.create_invoice_tab()
        self.main_tabs.addWidget(self.invoice_tab)
        self.create_pending_invoices_tab()
        self.main_tabs.addWidget(self.pending_invoices_tab)
        self.create_quotation_register_tab()
        self.main_tabs.addWidget(self.quotation_register_tab)
        self.create_finance_tab()
        self.main_tabs.addWidget(self.finance_tab)
        self.sidebar_navigation.currentRowChanged.connect(
            self.main_tabs.setCurrentIndex
        )
        self.main_tabs.currentChanged.connect(self.sync_sidebar_navigation)
        self.main_tabs.currentChanged.connect(self.refresh_finance_if_visible)
        content_layout.addWidget(self.main_tabs, 1)
        root_layout.addWidget(content, 1)
        self.setCentralWidget(central)
        return


    def finance_metric_card(self, title):

        card = QFrame()
        card.setStyleSheet(
            "QFrame { background: #ffffff; border: 1px solid #d9e2e6; "
            "border-radius: 5px; }"
        )
        card_layout = QVBoxLayout(card)
        card_layout.setContentsMargins(15, 12, 15, 12)
        card_layout.setSpacing(5)
        title_label = QLabel(title.upper())
        title_label.setStyleSheet(
            "color: #657780; border: 0; font-size: 8pt; font-weight: 600;"
        )
        value_label = QLabel("₹0.00")
        value_label.setFont(QFont("Segoe UI", 17, QFont.Weight.DemiBold))
        value_label.setStyleSheet("color: #193f4c; border: 0;")
        detail_label = QLabel("")
        detail_label.setStyleSheet("color: #72828a; border: 0; font-size: 9pt;")
        card_layout.addWidget(title_label)
        card_layout.addWidget(value_label)
        card_layout.addWidget(detail_label)
        return card, value_label, detail_label


    def create_finance_tab(self):

        self.finance_tab = QWidget()
        tab_layout = QVBoxLayout(self.finance_tab)
        tab_layout.setContentsMargins(0, 0, 0, 0)
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        finance_content = QWidget()
        scroll_area.setWidget(finance_content)
        tab_layout.addWidget(scroll_area)

        layout = QVBoxLayout(finance_content)
        layout.setContentsMargins(18, 14, 18, 16)
        layout.setSpacing(12)

        header = QHBoxLayout()
        heading = QLabel("Financial overview")
        heading.setFont(QFont("Segoe UI", 17, QFont.Weight.DemiBold))
        header.addWidget(heading)
        header.addStretch()
        self.finance_period_mode = "year"
        self.finance_period_buttons = QButtonGroup(self.finance_tab)
        self.finance_period_buttons.setExclusive(True)
        period_controls = QWidget()
        period_layout = QHBoxLayout(period_controls)
        period_layout.setContentsMargins(0, 0, 0, 0)
        period_layout.setSpacing(0)
        for button_id, label in ((0, "Year"), (1, "Month")):
            button = QPushButton(label)
            button.setCheckable(True)
            button.setFixedSize(72, 32)
            button.setStyleSheet(
                "QPushButton { background: #f2f5f6; color: #536873; "
                "border: 1px solid #cbd6da; padding: 5px 10px; }"
                "QPushButton:checked { background: #d0e6dd; color: #155b46; "
                "font-weight: 600; border-color: #21866b; }"
            )
            self.finance_period_buttons.addButton(button, button_id)
            period_layout.addWidget(button)
            if button_id == 0:
                button.setChecked(True)
        self.finance_period_buttons.idClicked.connect(
            lambda button_id: self.set_finance_period_mode(
                "month" if button_id else "year"
            )
        )
        header.addWidget(period_controls)
        self.finance_year_label = QLabel("Year")
        header.addWidget(self.finance_year_label)
        self.finance_year_selector = NoWheelComboBox()
        self.finance_year_selector.setMinimumWidth(110)
        header.addWidget(self.finance_year_selector)
        self.finance_month_label = QLabel("Month")
        header.addWidget(self.finance_month_label)
        self.finance_month_selector = NoWheelComboBox()
        self.finance_month_selector.addItems(FinanceMonthlyChart.MONTHS)
        self.finance_month_selector.setCurrentIndex(datetime.now().month - 1)
        header.addWidget(self.finance_month_selector)
        self.finance_month_label.hide()
        self.finance_month_selector.hide()
        refresh_button = QPushButton("Refresh")
        refresh_button.clicked.connect(self.refresh_finance_dashboard)
        header.addWidget(refresh_button)
        layout.addLayout(header)

        metrics = QHBoxLayout()
        metrics.setSpacing(10)
        self.finance_collected_card, self.finance_collected_value, self.finance_collected_detail = (
            self.finance_metric_card("Collected")
        )
        self.finance_sales_card, self.finance_sales_value, self.finance_sales_detail = (
            self.finance_metric_card("Invoiced")
        )
        self.finance_outstanding_card, self.finance_outstanding_value, self.finance_outstanding_detail = (
            self.finance_metric_card("Outstanding")
        )
        self.finance_growth_card, self.finance_growth_value, self.finance_growth_detail = (
            self.finance_metric_card("Collection growth")
        )
        for card in (
            self.finance_collected_card,
            self.finance_sales_card,
            self.finance_outstanding_card,
            self.finance_growth_card,
        ):
            metrics.addWidget(card, 1)
        layout.addLayout(metrics)

        charts = QHBoxLayout()
        trend_box = QGroupBox("Monthly collected vs outstanding")
        trend_layout = QVBoxLayout(trend_box)
        self.finance_monthly_chart = FinanceMonthlyChart()
        trend_layout.addWidget(self.finance_monthly_chart)
        self.finance_monthly_insight = QLabel(
            "Collected-to-date and outstanding amounts are grouped by invoice month."
        )
        self.finance_monthly_insight.setWordWrap(True)
        self.finance_monthly_insight.setStyleSheet("color: #657780;")
        trend_layout.addWidget(self.finance_monthly_insight)
        charts.addWidget(trend_box, 2)

        mix_box = QGroupBox("Invoice position")
        mix_layout = QVBoxLayout(mix_box)
        self.finance_donut_chart = FinanceDonutChart()
        mix_layout.addWidget(self.finance_donut_chart, 1)
        self.finance_mix_collected = QLabel("Collected: ₹0.00")
        self.finance_mix_pending = QLabel("Outstanding: ₹0.00")
        self.finance_mix_collected.setStyleSheet("color: #21866b; font-weight: 600;")
        self.finance_mix_pending.setStyleSheet("color: #b7792f; font-weight: 600;")
        mix_layout.addWidget(self.finance_mix_collected)
        mix_layout.addWidget(self.finance_mix_pending)
        charts.addWidget(mix_box, 1)
        layout.addLayout(charts, 1)

        detail_row = QHBoxLayout()
        debtors_box = QGroupBox("Largest outstanding customer balances")
        debtors_layout = QVBoxLayout(debtors_box)
        self.finance_debtors_table = QTableWidget(0, 3)
        self.finance_debtors_table.setHorizontalHeaderLabels(
            ["Customer", "Open invoices", "Balance due"]
        )
        self.finance_debtors_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.finance_debtors_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.finance_debtors_table.verticalHeader().setVisible(False)
        self.finance_debtors_table.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.ResizeMode.Stretch
        )
        self.finance_debtors_table.setMinimumHeight(170)
        debtors_layout.addWidget(self.finance_debtors_table)
        detail_row.addWidget(debtors_box, 3)

        insights_box = QGroupBox("At a glance")
        insights_layout = QVBoxLayout(insights_box)
        self.finance_overdue_label = QLabel("Overdue invoices: 0")
        self.finance_overdue_label.setWordWrap(True)
        self.finance_overdue_label.setStyleSheet("font-weight: 600; color: #193f4c;")
        self.finance_year_insight = QLabel("Year-on-year comparison will appear here.")
        self.finance_year_insight.setWordWrap(True)
        self.finance_year_insight.setStyleSheet("color: #657780;")
        self.finance_top_month_insight = QLabel("Most collected invoice month: --")
        self.finance_top_month_insight.setWordWrap(True)
        insights_layout.addWidget(self.finance_overdue_label)
        insights_layout.addWidget(self.finance_year_insight)
        insights_layout.addWidget(self.finance_top_month_insight)
        insights_layout.addStretch()
        detail_row.addWidget(insights_box, 2)
        layout.addLayout(detail_row)

        self.finance_year_selector.currentTextChanged.connect(
            self.refresh_finance_dashboard
        )
        self.finance_month_selector.currentIndexChanged.connect(
            lambda _index: self.refresh_finance_dashboard()
        )


    def sync_sidebar_navigation(self, index):

        if self.sidebar_navigation.currentRow() != index:
            self.sidebar_navigation.setCurrentRow(index)


    def refresh_finance_if_visible(self, index):

        if index == 3:
            self.refresh_finance_dashboard()


    def set_finance_period_mode(self, mode):

        self.finance_period_mode = mode
        is_month_mode = mode == "month"
        self.finance_month_label.setVisible(is_month_mode)
        self.finance_month_selector.setVisible(is_month_mode)
        self.refresh_finance_dashboard()


    def refresh_finance_dashboard(self, selected_year=None):

        if not hasattr(self, "finance_year_selector"):
            return
        if selected_year is None:
            selected_year = self.finance_year_selector.currentText()

        with sqlite3.connect(DB_PATH) as con:
            invoice_rows = con.execute(
                """
                SELECT i.invoice_no, i.invoice_date, COALESCE(i.customer_name, ''),
                       COALESCE(i.grand_total, 0), COALESCE(SUM(p.amount), 0)
                FROM invoices AS i
                LEFT JOIN invoice_payments AS p
                  ON p.document_kind = i.document_kind
                 AND p.invoice_no = i.invoice_no
                WHERE i.document_kind = 'INVOICE'
                GROUP BY i.id
                ORDER BY i.invoice_date DESC
                """
            ).fetchall()
            payment_rows = con.execute(
                """
                SELECT p.paid_on, p.amount
                FROM invoice_payments AS p
                INNER JOIN invoices AS i
                  ON i.document_kind = p.document_kind
                 AND i.invoice_no = p.invoice_no
                WHERE p.document_kind = 'INVOICE'
                """
            ).fetchall()

        available_years = {datetime.now().year}
        available_years.update(
            int(invoice_date[:4])
            for _, invoice_date, _, _, _ in invoice_rows
            if invoice_date and invoice_date[:4].isdigit()
        )
        available_years.update(
            int(paid_on[:4])
            for paid_on, _ in payment_rows
            if paid_on and paid_on[:4].isdigit()
        )
        year_text = str(selected_year or datetime.now().year)
        self.finance_year_selector.blockSignals(True)
        self.finance_year_selector.clear()
        self.finance_year_selector.addItems(
            [str(year) for year in sorted(available_years, reverse=True)]
        )
        if self.finance_year_selector.findText(year_text) >= 0:
            self.finance_year_selector.setCurrentText(year_text)
        self.finance_year_selector.blockSignals(False)
        year_text = self.finance_year_selector.currentText()
        year = int(year_text) if year_text.isdigit() else datetime.now().year
        today = datetime.now().date()
        comparison_cutoff = (
            f"{year - 1}-{today:%m-%d}"
            if year == today.year else f"{year - 1}-12-31"
        )
        comparison_period = (
            f"{year - 1} through {today:%b %d}"
            if year == today.year else str(year - 1)
        )
        month_mode = self.finance_period_mode == "month"
        month_number = self.finance_month_selector.currentIndex() + 1
        month_name = FinanceMonthlyChart.MONTHS[month_number - 1]
        selected_month_start = datetime(year, month_number, 1).date()
        next_month_start = (
            selected_month_start.replace(day=28) + timedelta(days=4)
        ).replace(day=1)
        previous_month_start = (
            selected_month_start - timedelta(days=1)
        ).replace(day=1)
        if year == today.year and month_number == today.month:
            selected_month_end = today + timedelta(days=1)
            previous_month_days = min(
                today.day,
                (selected_month_start - timedelta(days=1)).day,
            )
            previous_month_end = previous_month_start + timedelta(
                days=previous_month_days
            )
        else:
            selected_month_end = next_month_start
            previous_month_end = selected_month_start

        yearly_sales = sum(
            float(total or 0)
            for _, invoice_date, _, total, _ in invoice_rows
            if invoice_date and invoice_date.startswith(year_text)
        )
        monthly_receipts = [0.0] * 12
        monthly_outstanding = [0.0] * 12
        yearly_receipts = 0.0
        previous_year_receipts = 0.0
        month_receipts = 0.0
        previous_month_receipts = 0.0
        for paid_on, amount in payment_rows:
            if not paid_on or len(paid_on) < 7:
                continue
            payment_year = paid_on[:4]
            if payment_year == year_text:
                yearly_receipts += float(amount or 0)
            elif payment_year == str(year - 1) and paid_on <= comparison_cutoff:
                previous_year_receipts += float(amount or 0)
            payment_day = paid_on[:10]
            if selected_month_start.isoformat() <= payment_day < selected_month_end.isoformat():
                month_receipts += float(amount or 0)
            elif previous_month_start.isoformat() <= payment_day < previous_month_end.isoformat():
                previous_month_receipts += float(amount or 0)

        month_invoiced = 0.0
        month_outstanding = 0.0
        month_open_count = 0
        for _, invoice_date, _, total, paid in invoice_rows:
            if not invoice_date or not (
                selected_month_start.isoformat()
                <= invoice_date[:10]
                < selected_month_end.isoformat()
            ):
                continue
            invoice_total = max(0.0, float(total or 0))
            received = max(0.0, float(paid or 0))
            month_invoiced += invoice_total
            balance = max(0.0, invoice_total - received)
            month_outstanding += balance
            month_open_count += balance > 0.005

        for _, invoice_date, _, total, paid in invoice_rows:
            if not invoice_date or not invoice_date.startswith(year_text):
                continue
            try:
                month_index = int(invoice_date[5:7]) - 1
            except ValueError:
                continue
            if 0 <= month_index < 12:
                invoice_total = max(0.0, float(total or 0))
                received = max(0.0, float(paid or 0))
                monthly_receipts[month_index] += min(invoice_total, received)
                monthly_outstanding[month_index] += max(
                    0.0, invoice_total - received
                )

        total_collected = 0.0
        total_outstanding = 0.0
        overdue_count = 0
        pending_by_customer = {}
        for _, invoice_date, customer, total, paid in invoice_rows:
            invoice_total = max(0.0, float(total or 0))
            received = max(0.0, float(paid or 0))
            allocated = min(invoice_total, received)
            balance = max(0.0, invoice_total - received)
            total_collected += allocated
            total_outstanding += balance
            if balance > 0.005:
                buyer = customer.strip() or "Unnamed buyer"
                entry = pending_by_customer.setdefault(buyer, [0, 0.0])
                entry[0] += 1
                entry[1] += balance
                try:
                    invoice_day = datetime.strptime(invoice_date, "%Y-%m-%d").date()
                    if (today - invoice_day).days > 30:
                        overdue_count += 1
                except (TypeError, ValueError):
                    pass

        open_count = sum(count for count, _ in pending_by_customer.values())
        if month_mode:
            self.finance_collected_value.setText(f"₹{month_receipts:,.2f}")
            self.finance_collected_detail.setText(
                f"Cash received in {month_name} {year}"
            )
            self.finance_sales_value.setText(f"₹{month_invoiced:,.2f}")
            self.finance_sales_detail.setText(
                f"Invoices dated in {month_name} {year}"
            )
            self.finance_outstanding_value.setText(f"₹{month_outstanding:,.2f}")
            self.finance_outstanding_detail.setText(
                f"{month_open_count} open invoice(s) from {month_name} {year}"
            )
        else:
            self.finance_collected_value.setText(f"₹{yearly_receipts:,.2f}")
            self.finance_collected_detail.setText(f"Cash received during {year}")
            self.finance_sales_value.setText(f"₹{yearly_sales:,.2f}")
            self.finance_sales_detail.setText(f"Invoices dated in {year}")
            self.finance_outstanding_value.setText(f"₹{total_outstanding:,.2f}")
            self.finance_outstanding_detail.setText(
                f"{open_count} invoice(s) still open overall"
            )

        comparison_receipts = (
            previous_month_receipts if month_mode else previous_year_receipts
        )
        current_receipts = month_receipts if month_mode else yearly_receipts
        comparison_label = (
            f"vs {FinanceMonthlyChart.MONTHS[previous_month_start.month - 1]} "
            f"{previous_month_start.year}"
            if month_mode else f"vs {comparison_period}"
        )
        if comparison_receipts > 0:
            growth = (current_receipts - comparison_receipts) / comparison_receipts * 100
            growth_prefix = "+" if growth >= 0 else ""
            self.finance_growth_value.setText(f"{growth_prefix}{growth:.1f}%")
            self.finance_growth_detail.setText(
                f"{comparison_label} (₹{comparison_receipts:,.2f})"
            )
            self.finance_year_insight.setText(
                f"Receipts are {growth_prefix}{growth:.1f}% {comparison_label}."
            )
        else:
            self.finance_growth_value.setText("N/A")
            no_comparison = (
                f"No receipts in "
                f"{FinanceMonthlyChart.MONTHS[previous_month_start.month - 1]} "
                f"{previous_month_start.year}"
                if month_mode else f"No {year - 1} receipt history"
            )
            self.finance_growth_detail.setText(no_comparison)
            self.finance_year_insight.setText(
                f"No comparable receipt history. {no_comparison}."
            )

        self.finance_monthly_chart.set_values(
            monthly_receipts, monthly_outstanding
        )
        nonzero_months = [
            (amount, FinanceMonthlyChart.MONTHS[index])
            for index, amount in enumerate(monthly_receipts)
            if amount > 0
        ]
        if nonzero_months:
            best_amount, best_month = max(nonzero_months)
            self.finance_top_month_insight.setText(
                f"Most collected invoice month: {best_month} · ₹{best_amount:,.2f}"
            )
        else:
            self.finance_top_month_insight.setText("Best collection month: --")

        active_months = [
            index
            for index, (collected, outstanding) in enumerate(zip(
                monthly_receipts, monthly_outstanding
            ))
            if collected > 0 or outstanding > 0
        ]
        if active_months:
            latest_index = max(active_months)
            month_name = FinanceMonthlyChart.MONTHS[latest_index]
            month_summary = (
                f"{month_name}: collected ₹{monthly_receipts[latest_index]:,.2f}; "
                f"outstanding ₹{monthly_outstanding[latest_index]:,.2f}"
            )
            if latest_index > 0 and monthly_receipts[latest_index - 1] > 0:
                prior = monthly_receipts[latest_index - 1]
                change = (
                    monthly_receipts[latest_index] - prior
                ) / prior * 100
                month_summary += (
                    f"; collected-to-date change {change:+.1f}% "
                    "vs prior invoice month"
                )
            else:
                month_summary += "; no prior-month receipt comparison"
            self.finance_monthly_insight.setText(month_summary)
        else:
            self.finance_monthly_insight.setText(
                f"No receipts or outstanding invoices recorded in {year}."
            )

        self.finance_donut_chart.set_values(total_collected, total_outstanding)
        self.finance_mix_collected.setText(f"Collected: ₹{total_collected:,.2f}")
        self.finance_mix_pending.setText(f"Outstanding: ₹{total_outstanding:,.2f}")
        self.finance_overdue_label.setText(
            f"Open more than 30 days: {overdue_count} invoice(s)"
        )
        sorted_customers = sorted(
            pending_by_customer.items(), key=lambda item: item[1][1], reverse=True
        )[:5]
        self.finance_debtors_table.setRowCount(len(sorted_customers))
        for row, (customer, (invoice_count, balance)) in enumerate(sorted_customers):
            for column, value in enumerate((
                customer,
                str(invoice_count),
                f"₹{balance:,.2f}",
            )):
                self.finance_debtors_table.setItem(
                    row, column, QTableWidgetItem(value)
                )

        return

        self.invoice_type = NoWheelComboBox()

        self.invoice_type.addItems([
            "IGST",
            "CGST + SGST"
        ])

        self.invoice_type.currentIndexChanged.connect(
            self.refresh
        )
        self.invoice_date = NoWheelDateEdit()

        self.invoice_date.setCalendarPopup(
            True
        )

        self.invoice_date.setDisplayFormat(
            "dd/MM/yyyy"
        )

        self.invoice_date.dateChanged.connect(
            self.refresh
        )

        self.invoice_no.textChanged.connect(
            self.refresh
        )


        grid.addWidget(
            QLabel("Invoice Template"),
            0, 0
        )

        grid.addWidget(
            self.invoice_type,
            0, 1
        )


        grid.addWidget(
            QLabel("Invoice No"),
            0, 2
        )

        grid.addWidget(
            self.invoice_no,
            0, 3
        )


        grid.addWidget(
            QLabel("Date"),
            1, 0
        )

        grid.addWidget(
            self.invoice_date,
            1, 1
        )


        layout.addWidget(
            invoice_box
        )


        seller_box = QGroupBox(
            "Seller Details"
        )

        seller_form = QFormLayout(
            seller_box
        )

        self.seller_name = QLineEdit(
            "ADOODS GARMENTS"
        )

        self.seller_address = QLineEdit(
            "88-1/50M, 1st FLOOR NEAR SRINIVASA THEATRE, "
            "ULAVARSANDHAI ROAD, HOSUR - 635 109"
        )


        self.seller_gstin = QLineEdit(
            "33DJRPP6135G1ZW"
        )


        seller_form.addRow(
            "Company Name",
            self.seller_name
        )

        seller_form.addRow(
            "Address",
            self.seller_address
        )

        seller_form.addRow(
            "GSTIN",
            self.seller_gstin
        )


        layout.addWidget(
            seller_box
        )


        # ----------------------------------------------------
        # CUSTOMER
        # ----------------------------------------------------

        customer_box = QGroupBox(
            "Customer Details"
        )

        customer_form = QFormLayout(
            customer_box
        )


        self.customer_name = QLineEdit()

        self.customer_gstin = QLineEdit()

        self.customer_address = QLineEdit()

        self.customer_mobile = QLineEdit()

        customer_form.addRow(
            "Customer Name",
            self.customer_name
        )

        customer_form.addRow(
            "GSTIN",
            self.customer_gstin
        )

        customer_form.addRow(
            "Address",
            self.customer_address
        )

        customer_form.addRow(
            "Mobile",
            self.customer_mobile
        )

        self.customer_name.textChanged.connect(
            lambda value: self.sync_buyer_field(
                self.invoice_customer_name,
                value
            )
        )

        self.customer_gstin.textChanged.connect(
            lambda value: self.sync_buyer_field(
                self.invoice_customer_gstin,
                value
            )
        )

        self.customer_address.textChanged.connect(
            lambda value: self.sync_buyer_field(
                self.invoice_customer_address,
                value
            )
        )

        self.customer_mobile.textChanged.connect(
            lambda value: self.sync_buyer_field(
                self.invoice_customer_mobile,
                value
            )
        )


        layout.addWidget(
            customer_box
        )


        # ----------------------------------------------------
        # GST
        # ----------------------------------------------------

        tax_box = QGroupBox(
            "Tax"
        )

        tax_form = QFormLayout(
            tax_box
        )


        self.gst_rate = NoWheelDoubleSpinBox()

        self.gst_rate.setRange(
            0,
            100
        )

        self.gst_rate.setDecimals(
            2
        )

        self.gst_rate.setValue(
            5
        )

        self.gst_rate.setSuffix(
            " %"
        )

        self.gst_rate.valueChanged.connect(
            self.refresh
        )


        tax_form.addRow(
            "GST Rate",
            self.spinbox_with_arrows(self.gst_rate)
        )


        layout.addWidget(
            tax_box
        )


        # ----------------------------------------------------
        # BUTTONS
        # ----------------------------------------------------

        buttons = QHBoxLayout()


        new_button = QPushButton(
            "New Invoice"
        )

        new_button.clicked.connect(
            self.new_invoice
        )


        save_button = QPushButton(
            "Save Invoice"
        )

        save_button.setObjectName(
            "primary"
        )

        save_button.clicked.connect(
            lambda checked=False: self.save_invoice()
        )


        pdf_button = QPushButton(
            "Download PDF"
        )

        pdf_button.setObjectName(
            "primary"
        )

        pdf_button.clicked.connect(
            self.save_document_bundle
        )


        buttons.addWidget(
            new_button
        )

        buttons.addWidget(
            save_button
        )

        buttons.addWidget(
            pdf_button
        )

        buttons.addStretch()


        layout.addLayout(
            buttons
        )

        layout.addStretch()


    # ========================================================
    # INVOICE TAB
    # ========================================================

    def create_legacy_invoice_tab(self):

        layout = QVBoxLayout(
            self.invoice_tab
        )

        buyer_box = QGroupBox(
            "Buyer Details"
        )

        buyer_form = QFormLayout(
            buyer_box
        )

        self.invoice_customer_name = QLineEdit()
        self.invoice_customer_gstin = QLineEdit()
        self.invoice_customer_address = QLineEdit()
        self.invoice_customer_mobile = QLineEdit()

        buyer_form.addRow(
            "Customer Name",
            self.invoice_customer_name
        )
        buyer_form.addRow(
            "GSTIN",
            self.invoice_customer_gstin
        )
        buyer_form.addRow(
            "Address",
            self.invoice_customer_address
        )
        buyer_form.addRow(
            "Mobile",
            self.invoice_customer_mobile
        )

        for main_field, invoice_field in (
            (self.customer_name, self.invoice_customer_name),
            (self.customer_gstin, self.invoice_customer_gstin),
            (self.customer_address, self.invoice_customer_address),
            (self.customer_mobile, self.invoice_customer_mobile),
        ):
            invoice_field.textChanged.connect(
                lambda value, target=main_field: self.sync_buyer_field(
                    target,
                    value
                )
            )

        layout.addWidget(
            buyer_box
        )


        box = QGroupBox(
            "Items"
        )

        box_layout = QVBoxLayout(
            box
        )


        self.items = QTableWidget(
            10,
            7
        )


        self.items.setHorizontalHeaderLabels([
            "No.",
            "Description of Goods",
            "HSN / SAC",
            "Quantity",
            "Rate",
            "Per",
            "Amount"
        ])


        self.items.horizontalHeader().setSectionResizeMode(
            1,
            QHeaderView.ResizeMode.Stretch
        )


        for column in (0, 2, 3, 4, 5, 6):

            self.items.horizontalHeader().setSectionResizeMode(
                column,
                QHeaderView.ResizeMode.ResizeToContents
            )

        for row in range(self.items.rowCount()):
            for column in range(self.items.columnCount()):
                item = QTableWidgetItem("")
                if column not in (1, 3, 4):
                    item.setFlags(
                        item.flags()
                        & ~Qt.ItemFlag.ItemIsEditable
                    )
                self.items.setItem(row, column, item)


        self.items.itemChanged.connect(
            self.refresh
        )


        box_layout.addWidget(
            self.items
        )


        layout.addWidget(
            box
        )


        # ----------------------------------------------------
        # TOTAL
        # ----------------------------------------------------

        total_box = QGroupBox(
            "Invoice Total"
        )

        total_layout = QVBoxLayout(
            total_box
        )


        self.summary = QLabel()

        self.summary.setFont(
            QFont(
                "Segoe UI",
                11,
                QFont.Weight.DemiBold
            )
        )


        total_layout.addWidget(
            self.summary
        )


        layout.addWidget(
            total_box
        )


        # ----------------------------------------------------
        # BUTTONS
        # ----------------------------------------------------

        buttons = QHBoxLayout()


        save_button = QPushButton(
            "Save Invoice"
        )

        save_button.setObjectName(
            "primary"
        )

        save_button.clicked.connect(
            lambda checked=False: self.save_invoice()
        )


        pdf_button = QPushButton(
            "Download PDF"
        )

        pdf_button.setObjectName(
            "primary"
        )

        pdf_button.clicked.connect(
            self.save_document_bundle
        )


        buttons.addWidget(
            save_button
        )

        buttons.addWidget(
            pdf_button
        )

        buttons.addStretch()


        layout.addLayout(
            buttons
        )


    def sync_buyer_field(self, target, value):

        if target.text() != value:
            target.setText(value)


    def create_invoice_tab(self):

        layout = QVBoxLayout(self.invoice_tab)
        self.create_mode_tabs = QTabWidget()
        self.create_mode_tabs.addTab(QWidget(), "+ Invoice")
        self.create_mode_tabs.addTab(QWidget(), "+ Quotation")
        self.create_mode_tabs.setFixedHeight(58)
        self.create_mode_tabs.setStyleSheet(
            "QTabWidget::pane { border: 0; }"
            "QTabBar::tab { min-width: 150px; padding: 9px 18px; "
            "color: #536873; background: #e9eff2; border: 0; }"
            "QTabBar::tab:selected { color: #155b46; background: #d0e6dd; "
            "font-weight: 600; border-bottom: 3px solid #21866b; }"
        )
        layout.addWidget(self.create_mode_tabs)
        top_actions = QHBoxLayout()
        self.new_document_button = QPushButton("+ New Invoice")
        self.new_document_button.clicked.connect(self.new_invoice)
        self.private_invoice_checkbox = QCheckBox("Private invoice")
        self.private_invoice_checkbox.toggled.connect(
            self.private_invoice_changed
        )
        top_actions.addWidget(self.new_document_button)
        top_actions.addWidget(self.private_invoice_checkbox)
        top_actions.addStretch()
        layout.addLayout(top_actions)
        scroll_area = QScrollArea()
        self.invoice_scroll_area = scroll_area
        scroll_area.setWidgetResizable(True)
        page = QWidget()
        page_layout = QVBoxLayout(page)
        page_layout.setContentsMargins(18, 18, 18, 18)

        self.invoice_request_history_box = QGroupBox("Payment Request History")
        history_layout = QVBoxLayout(self.invoice_request_history_box)
        history_actions = QHBoxLayout()
        self.open_invoice_request_button = QPushButton("Open PDF")
        self.open_invoice_request_button.clicked.connect(
            self.open_selected_invoice_request
        )
        self.download_invoice_request_button = QPushButton("Download PDF")
        self.download_invoice_request_button.clicked.connect(
            self.download_selected_invoice_request
        )
        self.send_paid_invoice_button = QPushButton("Send PDF")
        self.send_paid_invoice_button.clicked.connect(
            self.send_selected_paid_invoice
        )
        self.open_invoice_request_button.setEnabled(False)
        self.download_invoice_request_button.setEnabled(False)
        self.send_paid_invoice_button.setEnabled(False)
        history_actions.addWidget(self.open_invoice_request_button)
        history_actions.addWidget(self.download_invoice_request_button)
        history_actions.addWidget(self.send_paid_invoice_button)
        history_actions.addStretch()
        history_layout.addLayout(history_actions)
        self.invoice_request_history_table = QTableWidget(0, 4)
        self.invoice_request_history_table.setHorizontalHeaderLabels(
            ["Date / Time", "Payment Request", "Status", "PDF"]
        )
        self.invoice_request_history_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.invoice_request_history_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.invoice_request_history_table.setSelectionMode(
            QTableWidget.SelectionMode.SingleSelection
        )
        self.invoice_request_history_table.verticalHeader().setVisible(False)
        self.invoice_request_history_table.horizontalHeader().setSectionResizeMode(
            3, QHeaderView.ResizeMode.Stretch
        )
        self.invoice_request_history_table.setMinimumHeight(240)
        self.invoice_request_history_table.setMaximumHeight(360)
        self.invoice_request_history_table.itemSelectionChanged.connect(
            self.update_invoice_request_actions
        )
        self.invoice_request_history_table.cellDoubleClicked.connect(
            lambda row, column: self.open_selected_invoice_request()
        )
        history_layout.addWidget(self.invoice_request_history_table)
        self.invoice_request_history_box.setVisible(False)

        seller_box = QGroupBox("Seller Details")
        seller_grid = QGridLayout(seller_box)
        self.invoice_heading = QLineEdit("TAX INVOICE")
        self.seller_name = QLineEdit("ADOODS GARMENTS")
        self.seller_business = QLineEdit("ALL KINDS OF T-SHIRT MANUFACTURE")
        self.seller_address_1 = QLineEdit("88-1/50M, 1st FLOOR NEAR SRINIVASA THEATRE,")
        self.seller_address_2 = QLineEdit("ULAVARSANDHAI ROAD, HOSUR - 635 109")
        self.seller_gstin = QLineEdit("33DJRPP6135G1ZW")
        self.seller_state = QLineEdit("STATE: TAMIL NADU   CODE: 33")
        self.seller_website = QLineEdit("adoods.com")
        self.seller_phone = QLineEdit("+91 9514109080, +91 6382208004")
        self.seller_email_phone = QLineEdit("Email: adoodshosur@gmail.com   +91 6382208004")
        self.invoice_heading.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.invoice_heading.setStyleSheet("font: 20pt 'Arial Black';")
        self.seller_name.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.seller_name.setStyleSheet("font: bold 26pt 'Arial Black';")
        seller_rows = (
            ("Title", self.invoice_heading),
            ("Company", self.seller_name),
            ("Business", self.seller_business),
            ("Address", self.seller_address_1),
            ("Address", self.seller_address_2),
            ("GSTIN", self.seller_gstin),
            ("State", self.seller_state),
            ("Website", self.seller_website),
            ("Phone", self.seller_phone),
            ("Email / Phone", self.seller_email_phone),
        )
        for row, (label, field) in enumerate(seller_rows):
            seller_grid.addWidget(QLabel(label), row, 0)
            seller_grid.addWidget(field, row, 1)
        page_layout.addWidget(seller_box)

        party_grid = QGridLayout()
        buyer_box = QGroupBox("Buyer Details")
        buyer_form = QFormLayout(buyer_box)
        self.customer_name = QLineEdit()
        self.customer_address = QLineEdit()
        self.customer_address_2 = QLineEdit()
        self.customer_mobile = QLineEdit()
        self.customer_mobile.setPlaceholderText("10-digit Indian mobile number")
        self.customer_mobile.setValidator(
            QRegularExpressionValidator(
                QRegularExpression(r"[+0-9() -]{0,20}"),
                self.customer_mobile,
            )
        )
        self.customer_gstin = QLineEdit()
        self.buyer_selector = NoWheelComboBox()
        self.buyer_selector.setEditable(True)
        self.buyer_selector.setInsertPolicy(QComboBox.InsertPolicy.NoInsert)
        self.buyer_selector.lineEdit().setPlaceholderText(
            "Search or select a saved buyer"
        )
        buyer_completer = QCompleter(
            self.buyer_selector.model(), self.buyer_selector
        )
        buyer_completer.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
        buyer_completer.setFilterMode(Qt.MatchFlag.MatchContains)
        self.buyer_selector.setCompleter(buyer_completer)
        self.buyer_selector.currentIndexChanged.connect(self.load_buyer_profile)
        buyer_profile_row = QHBoxLayout()
        buyer_profile_row.addWidget(self.buyer_selector, 1)
        new_buyer_button = QPushButton("New Buyer")
        new_buyer_button.clicked.connect(self.new_buyer_profile)
        save_buyer_button = QPushButton("Save / Update")
        save_buyer_button.clicked.connect(self.save_buyer_profile)
        delete_buyer_button = QPushButton("Delete")
        delete_buyer_button.clicked.connect(self.delete_buyer_profile)
        buyer_profile_row.addWidget(new_buyer_button)
        buyer_profile_row.addWidget(save_buyer_button)
        buyer_profile_row.addWidget(delete_buyer_button)
        buyer_form.addRow("Saved buyer details", buyer_profile_row)
        buyer_form.addRow("Name", self.customer_name)
        buyer_form.addRow("Address line 1", self.customer_address)
        buyer_form.addRow("Address line 2", self.customer_address_2)
        buyer_form.addRow("Mobile", self.customer_mobile)
        buyer_form.addRow("GSTIN", self.customer_gstin)

        invoice_box = QGroupBox("Invoice Details")
        self.document_details_box = invoice_box
        invoice_form = QFormLayout(invoice_box)
        self.invoice_no = QLineEdit()
        self.document_kind = NoWheelComboBox()
        self.document_kind.addItems(["Invoice", "Quotation"])
        self.document_kind.hide()
        self.create_mode_tabs.currentChanged.connect(
            self.document_kind.setCurrentIndex
        )
        self.document_number_label = QLabel("Invoice No")
        self.invoice_date = NoWheelDateEdit(QDate.currentDate())
        self.invoice_date.setCalendarPopup(True)
        self.invoice_date.setDisplayFormat("dd/MM/yyyy")
        self.invoice_date.setKeyboardTracking(False)
        self.po_number = QLineEdit()
        self.invoice_type = NoWheelComboBox()
        self.invoice_type.addItems(["IGST", "CGST + SGST"])
        self.gst_rate = NoWheelDoubleSpinBox()
        self.gst_rate.setRange(0, 100)
        self.gst_rate.setDecimals(2)
        self.gst_rate.setValue(5)
        self.gst_rate.setSuffix(" %")
        invoice_form.addRow(self.document_number_label, self.invoice_no)
        invoice_form.addRow("Date", self.invoice_date)
        invoice_form.addRow("P.O No", self.po_number)
        invoice_form.addRow("Tax Type", self.invoice_type)
        invoice_form.addRow("GST Rate", self.spinbox_with_arrows(self.gst_rate))
        self.document_kind.currentIndexChanged.connect(self.document_kind_changed)
        self.document_kind.currentIndexChanged.connect(
            self.create_mode_tabs.setCurrentIndex
        )
        self.invoice_no.editingFinished.connect(self.load_invoice)
        party_grid.addWidget(buyer_box, 0, 0)
        party_grid.addWidget(invoice_box, 0, 1)
        page_layout.addLayout(party_grid)

        self.payment_percent = NoWheelDoubleSpinBox()
        self.payment_percent.setRange(0, 99999999999999.99)
        self.payment_percent.setDecimals(2)
        self.payment_percent.setValue(50)
        self.payment_percent.setSuffix(" % of invoice")
        self.payment_percent.setToolTip(
            "Percentage of the full invoice total. The request cannot exceed the balance due."
        )
        self._updating_payment_request = False
        self._payment_request_source = "percent"
        self.payment_percent.valueChanged.connect(
            lambda value: self.refresh_payment_request(source="percent")
        )

        self.invoice_payments_box = QGroupBox("Invoice Receipts")
        invoice_payments_layout = QVBoxLayout(self.invoice_payments_box)
        receipt_form = QGridLayout()
        receipt_form.addWidget(QLabel("Request %"), 0, 0)
        receipt_form.addWidget(
            self.spinbox_with_arrows(self.payment_percent, step_size=5), 0, 1
        )
        receipt_form.addWidget(QLabel("Request amount"), 0, 2)
        self.payment_request_amount = AmountEntryDoubleSpinBox()
        self.payment_request_amount.setRange(0, 999999999.99)
        self.payment_request_amount.setDecimals(2)
        self.payment_request_amount.setPrefix("₹ ")
        self.payment_request_amount.setToolTip(
            "Enter the amount to request, up to the balance due. Its invoice percentage updates automatically."
        )
        self.payment_request_amount.valueChanged.connect(
            lambda value: self.refresh_payment_request(source="amount")
        )
        receipt_form.addWidget(
            self.spinbox_with_arrows(self.payment_request_amount), 0, 3
        )
        self.remaining_percent_label = QLabel("Balance remaining: 100%")
        receipt_form.addWidget(QLabel("Invoice balance"), 1, 0)
        receipt_form.addWidget(self.remaining_percent_label, 1, 1, 1, 3)
        self.payment_request_warning = QLabel()
        self.payment_request_warning.setStyleSheet(
            "color: #b42318; font-weight: 600;"
        )
        self.payment_request_warning.setWordWrap(True)
        self.payment_request_warning.hide()
        receipt_form.addWidget(self.payment_request_warning, 2, 0, 1, 4)
        self._payment_request_exceeds_balance = False
        self._payment_request_balance_due = Decimal("0.00")
        self.receipt_amount = AmountEntryDoubleSpinBox()
        self.receipt_amount.setRange(0, 999999999.99)
        self.receipt_amount.setDecimals(2)
        self.receipt_amount.setPrefix("₹ ")
        self._auto_receipt_amount = Decimal("0.00")
        self._receipt_amount_manual_override = False
        self._updating_receipt_amount = False
        self.receipt_amount.valueChanged.connect(
            self.track_receipt_amount_override
        )
        self.receipt_method = NoWheelComboBox()
        self.receipt_method.addItems(
            ["Cash", "UPI", "Bank Transfer", "Cheque", "Other"]
        )
        self.receipt_date = NoWheelDateEdit(QDate.currentDate())
        self.receipt_date.setCalendarPopup(True)
        self.receipt_date.setDisplayFormat("dd/MM/yyyy")
        self.receipt_reference = QLineEdit()
        self.receipt_reference.setPlaceholderText("Optional transaction / receipt reference")
        receipt_form.addWidget(QLabel("Amount received"), 3, 0)
        receipt_form.addWidget(
            self.spinbox_with_arrows(self.receipt_amount), 3, 1
        )
        receipt_form.addWidget(QLabel("Payment method"), 3, 2)
        receipt_form.addWidget(self.receipt_method, 3, 3)
        receipt_form.addWidget(QLabel("Payment date"), 4, 0)
        receipt_form.addWidget(self.receipt_date, 4, 1)
        receipt_form.addWidget(QLabel("Reference"), 4, 2)
        receipt_form.addWidget(self.receipt_reference, 4, 3)
        receipt_form.setColumnStretch(1, 1)
        receipt_form.setColumnStretch(3, 1)
        invoice_payments_layout.addLayout(receipt_form)
        self.record_payment_button = QPushButton("Record Payment")
        self.record_payment_button.clicked.connect(self.record_invoice_payment)
        self.record_payment_button.setMaximumWidth(160)
        self.record_payment_button.setMinimumHeight(34)
        self.edit_payment_button = QPushButton("Edit Selected")
        self.edit_payment_button.clicked.connect(self.edit_selected_invoice_payment)
        self.delete_payment_button = QPushButton("Delete Selected")
        self.delete_payment_button.clicked.connect(self.delete_selected_invoice_payment)
        self.edit_payment_button.setEnabled(False)
        self.delete_payment_button.setEnabled(False)
        payment_actions = QHBoxLayout()
        payment_actions.addWidget(self.record_payment_button)
        payment_actions.addWidget(self.edit_payment_button)
        payment_actions.addWidget(self.delete_payment_button)
        payment_actions.addStretch()
        self.invoice_payment_summary = QLabel("Save the invoice before recording payments.")
        payment_actions.addWidget(self.invoice_payment_summary, 1)
        invoice_payments_layout.addLayout(payment_actions)
        self.invoice_payments_table = QTableWidget(0, 4)
        self.invoice_payments_table.setHorizontalHeaderLabels(
            ["Date", "Amount", "Method", "Reference"]
        )
        self.invoice_payments_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.invoice_payments_table.verticalHeader().setVisible(False)
        self.invoice_payments_table.horizontalHeader().setSectionResizeMode(
            3, QHeaderView.ResizeMode.Stretch
        )
        self.invoice_payments_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.invoice_payments_table.setSelectionMode(
            QTableWidget.SelectionMode.SingleSelection
        )
        self.invoice_payments_table.setMinimumHeight(160)
        self.invoice_payments_table.setMaximumHeight(240)
        self.invoice_payments_table.itemSelectionChanged.connect(
            self.update_invoice_payment_actions
        )
        invoice_payments_layout.addWidget(self.invoice_payments_table)
        self.record_payment_button.setEnabled(False)
        self.invoice_payments_box.setVisible(False)

        items_box = QGroupBox("Items")
        items_layout = QVBoxLayout(items_box)
        self.items = QTableWidget(0, 8)
        self.items.setFixedHeight(560)
        self.items.setWordWrap(True)
        self.items.setAlternatingRowColors(True)
        self.items.setStyleSheet(
            "QTableWidget { background: #ffffff; alternate-background-color: #f4f7fb; }"
            " QTableWidget::item { color: #1f2937; }"
        )
        self.items.verticalHeader().setDefaultSectionSize(52)
        self.items.verticalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Fixed
        )
        self.items.setHorizontalHeaderLabels([
            "SL.NO", "Description of Goods", "HSN/SAC", "Quantity",
            "Rate", "Per", "Amount", "Remove"
        ])
        self.items.horizontalHeader().setSectionResizeMode(
            1, QHeaderView.ResizeMode.Stretch
        )
        for column in (0, 2, 3, 4, 5, 6, 7):
            self.items.horizontalHeader().setSectionResizeMode(
                column, QHeaderView.ResizeMode.ResizeToContents
            )
        item_buttons = QHBoxLayout()
        add_item_button = QPushButton("Add Item")
        add_item_button.clicked.connect(self.add_item_row)
        item_buttons.addWidget(add_item_button)
        item_buttons.addStretch()
        items_layout.addLayout(item_buttons)
        items_layout.addWidget(self.items)
        page_layout.addWidget(items_box)

        totals_grid = QGridLayout()
        self.summary = QLabel()
        self.summary.setFont(QFont("Arial Black", 12, QFont.Weight.Bold))
        self.hsn_code = QLineEdit("6006")
        self.amount_words = QLineEdit()
        self.tax_words = QLineEdit()
        self.amount_words.setReadOnly(True)
        self.tax_words.setReadOnly(True)
        totals_grid.addWidget(QLabel("Amount Chargeable (in words)"), 0, 0)
        totals_grid.addWidget(self.amount_words, 0, 1, 1, 3)
        totals_grid.addWidget(QLabel("Tax Amount (in words)"), 1, 0)
        totals_grid.addWidget(self.tax_words, 1, 1, 1, 3)
        totals_grid.addWidget(QLabel("HSN/SAC"), 2, 0)
        totals_grid.addWidget(self.hsn_code, 2, 1)

        self.hsn_table = QTableWidget(2, 8)
        self.hsn_table.setHorizontalHeaderLabels([
            "HSN/SAC", "", "Taxable Value", "Tax Rate", "Tax Amount",
            "Tax Rate", "Tax Amount", "Total Tax"
        ])
        hsn_header = self.hsn_table.horizontalHeader()
        hsn_header.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        hsn_header.setMinimumSectionSize(72)
        hsn_header.setMinimumHeight(58)
        hsn_header.setStretchLastSection(True)
        self.hsn_table.setWordWrap(True)
        self.hsn_table.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAsNeeded
        )
        self.hsn_table.setMinimumHeight(124)
        for column, width in enumerate((82, 48, 116, 82, 102, 82, 102, 102)):
            hsn_header.resizeSection(column, width)
        self.hsn_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.hsn_table.verticalHeader().setVisible(False)
        totals_grid.addWidget(self.hsn_table, 3, 0, 1, 4)
        page_layout.addWidget(self.summary)
        page_layout.addLayout(totals_grid)

        bank_box = QGroupBox("Bank Profiles")
        self.bank_box = bank_box
        bank_grid = QGridLayout(bank_box)
        self.bank_name = QLineEdit("K V B")
        self.bank_account = QLineEdit("1607135000011387")
        self.bank_ifsc = QLineEdit("KVBL0001607")
        self.bank_upi_id = QLineEdit()
        self.bank_upi_id.setPlaceholderText("Required for PhonePe / Google Pay QR")
        self.bank_upi_label = QLabel("UPI ID for amount-specific QR")
        self.bank_selector = NoWheelComboBox()
        self.bank_selector.currentIndexChanged.connect(self.load_bank_profile)
        new_bank_button = QPushButton("New Bank")
        new_bank_button.clicked.connect(self.new_bank_profile)
        save_bank_button = QPushButton("Save / Update")
        save_bank_button.clicked.connect(self.save_bank_profile)
        delete_bank_button = QPushButton("Delete Selected")
        delete_bank_button.clicked.connect(self.delete_bank_profile)
        bank_profile_row = QHBoxLayout()
        bank_profile_row.addWidget(self.bank_selector, 1)
        bank_profile_row.addWidget(new_bank_button)
        bank_profile_row.addWidget(save_bank_button)
        bank_profile_row.addWidget(delete_bank_button)
        bank_grid.addWidget(QLabel("Saved profiles"), 0, 0)
        bank_grid.addLayout(bank_profile_row, 0, 1)
        for row, (label, field) in enumerate((
            ("Bank name", self.bank_name),
            ("Account number", self.bank_account),
            ("IFSC", self.bank_ifsc),
        ), start=1):
            bank_grid.addWidget(QLabel(label), row, 0)
            bank_grid.addWidget(field, row, 1)
        bank_grid.addWidget(self.bank_upi_label, 4, 0)
        bank_grid.addWidget(self.bank_upi_id, 4, 1)
        page_layout.addWidget(self.invoice_payments_box)
        page_layout.addWidget(bank_box)

        footer_box = QGroupBox("Invoice Footer")
        footer_grid = QGridLayout(footer_box)
        self.bundle_count = QLineEdit()
        self.through = QLineEdit()
        self.declaration = QTextEdit(
            "We declare that the Invoice shows the actual price of the goods "
            "described and that all particulars are true and correct."
        )
        self.declaration.setMinimumHeight(72)
        self.declaration.setMaximumHeight(96)
        self.seller_signatory = QLineEdit("for ADOODS GARMENTS")
        self.jurisdiction = QLineEdit(
            "SUBJECT TO HOSUR JURISDICTION. THIS IS COMPUTER GENERATED INVOICE"
        )
        for field in self.seller_profile_fields():
            field.editingFinished.connect(self.save_seller_profile)
        self.signature_image_path = (
            SIGNATURE_PATH if SIGNATURE_PATH.is_file() else None
        )
        footer_fields = (
            ("No of Bundles", self.bundle_count),
            ("Through", self.through),
            ("Declaration", self.declaration),
            ("Authorised Signatory", self.seller_signatory),
            ("Footer", self.jurisdiction),
        )
        for row, (label, field) in enumerate(footer_fields):
            footer_grid.addWidget(QLabel(label), row, 0)
            footer_grid.addWidget(field, row, 1)
        signature_row = QHBoxLayout()
        self.signature_preview = QLabel("No signature")
        self.signature_preview.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.signature_preview.setFixedSize(180, 70)
        self.signature_preview.setStyleSheet(
            "QLabel { background: #f5f8f7; color: #71817b; "
            "border: 1px dashed #9db4aa; border-radius: 4px; }"
        )
        self.signature_image_label = QLabel(
            "Signature ready for invoices"
            if self.signature_image_path else "No signature uploaded"
        )
        self.signature_image_label.setStyleSheet("color: #657780;")
        self.upload_signature_button = QPushButton(
            "Replace signature" if self.signature_image_path else "Upload signature PNG"
        )
        self.upload_signature_button.clicked.connect(self.upload_signature)
        signature_row.addWidget(self.signature_preview)
        signature_row.addWidget(self.signature_image_label)
        signature_row.addWidget(self.upload_signature_button)
        signature_row.addStretch()
        footer_grid.addLayout(signature_row, len(footer_fields), 0, 1, 2)
        self.refresh_signature_display()
        page_layout.addWidget(footer_box)

        action_row = QHBoxLayout()
        self.whatsapp_button = QPushButton("WhatsApp Invoice")
        self.whatsapp_button.setObjectName("primary")
        self.whatsapp_button.clicked.connect(self.share_document_whatsapp)
        self.whatsapp_button.setVisible(False)
        self.attach_invoice_checkbox = QCheckBox("Attach invoice when sharing")
        self.attach_invoice_checkbox.setChecked(True)
        self.attach_invoice_checkbox.setVisible(False)
        self.save_document_button = QPushButton("Save Invoice")
        self.save_document_button.setObjectName("primary")
        self.save_document_button.clicked.connect(
            lambda checked=False: self.save_document_bundle()
        )
        action_row.addWidget(self.save_document_button)
        action_row.addWidget(self.whatsapp_button)
        action_row.addWidget(self.attach_invoice_checkbox)
        action_row.addStretch()
        page_layout.addLayout(action_row)
        page_layout.addWidget(self.invoice_request_history_box)

        scroll_area.setWidget(page)
        layout.addWidget(scroll_area)
        self.items.itemChanged.connect(self.refresh)
        self.invoice_type.currentIndexChanged.connect(self.refresh)
        self.invoice_date.dateChanged.connect(self.refresh)
        self.gst_rate.valueChanged.connect(self.refresh)
        self.load_seller_profile()
        self.refresh_buyer_profiles()
        self.refresh_bank_profiles()
        self.refresh_invoice_history()
        self.document_kind_changed(self.document_kind.currentIndex())


    def create_pending_invoices_tab(self):

        self.pending_invoices_tab = QWidget()
        layout = QVBoxLayout(self.pending_invoices_tab)
        toolbar = QHBoxLayout()
        self.pending_invoice_private = False
        self.pending_invoice_series_group = QButtonGroup(
            self.pending_invoices_tab
        )
        series_controls = QWidget()
        series_layout = QHBoxLayout(series_controls)
        series_layout.setContentsMargins(0, 0, 0, 0)
        series_layout.setSpacing(0)
        for button_id, label in ((0, "Invoices"), (1, "Private")):
            button = QPushButton(label)
            button.setCheckable(True)
            button.setMinimumWidth(96)
            button.setStyleSheet(
                "QPushButton { background: #eef2f3; color: #405761; "
                "border: 1px solid #b9c7cb; padding: 7px 12px; }"
                "QPushButton:checked { background: #155b46; color: white; "
                "border-color: #155b46; font-weight: 600; }"
                "QPushButton:hover:!checked { background: #dce8e5; }"
            )
            self.pending_invoice_series_group.addButton(button, button_id)
            series_layout.addWidget(button)
            if button_id == 0:
                button.setChecked(True)
        self.pending_invoice_series_group.idClicked.connect(
            lambda button_id: self.set_pending_invoice_series(button_id == 1)
        )
        toolbar.addWidget(series_controls)
        self.pending_invoice_search = QLineEdit()
        self.pending_invoice_search.setPlaceholderText(
            "Search customer or invoice number"
        )
        self.pending_invoice_search.textChanged.connect(
            lambda value: self.reset_pending_invoice_page()
        )
        toolbar.addWidget(self.pending_invoice_search, 1)
        self.include_paid_invoices = QCheckBox("Include paid invoices")
        self.include_paid_invoices.toggled.connect(
            lambda checked: self.reset_pending_invoice_page()
        )
        toolbar.addWidget(self.include_paid_invoices)
        self.open_pending_invoice_button = QPushButton("Open Selected Invoice")
        self.open_pending_invoice_button.clicked.connect(
            self.open_selected_pending_invoice
        )
        toolbar.addWidget(self.open_pending_invoice_button)
        layout.addLayout(toolbar)

        self.pending_invoices_table = QTableWidget(0, 6)
        self.pending_invoices_table.setHorizontalHeaderLabels(
            ["Invoice No", "Date", "Customer", "Total", "Received", "Balance Due"]
        )
        self.pending_invoices_table.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.pending_invoices_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.pending_invoices_table.setSelectionMode(
            QTableWidget.SelectionMode.SingleSelection
        )
        self.pending_invoices_table.verticalHeader().setVisible(False)
        self.pending_invoices_table.horizontalHeader().setSectionResizeMode(
            2, QHeaderView.ResizeMode.Stretch
        )
        self.pending_invoices_table.cellDoubleClicked.connect(
            lambda row, column: self.open_selected_pending_invoice()
        )
        layout.addWidget(self.pending_invoices_table, 1)

        pagination = QHBoxLayout()
        self.pending_previous_button = QPushButton("Previous")
        self.pending_previous_button.clicked.connect(
            lambda checked=False: self.change_pending_invoice_page(-1)
        )
        self.pending_page_label = QLabel("Page 1 of 1")
        self.pending_page_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.pending_next_button = QPushButton("Next")
        self.pending_next_button.clicked.connect(
            lambda checked=False: self.change_pending_invoice_page(1)
        )
        pagination.addWidget(self.pending_previous_button)
        pagination.addWidget(self.pending_page_label, 1)
        pagination.addWidget(self.pending_next_button)
        layout.addLayout(pagination)

        self.pending_invoice_page = 0
        self.pending_invoice_page_size = 20
        self.refresh_pending_invoices()


    def reset_pending_invoice_page(self):

        self.pending_invoice_page = 0
        self.refresh_pending_invoices()


    def set_pending_invoice_series(self, is_private):

        self.pending_invoice_private = is_private
        self.reset_pending_invoice_page()


    def change_pending_invoice_page(self, offset):

        self.pending_invoice_page += offset
        self.refresh_pending_invoices()


    def refresh_pending_invoices(self):

        if not hasattr(self, "pending_invoices_table"):
            return

        search = self.pending_invoice_search.text().strip()
        search_pattern = f"%{search}%"
        paid_filter = "" if self.include_paid_invoices.isChecked() else (
            "COALESCE(i.grand_total, 0) - COALESCE(SUM(p.amount), 0) > 0.005 AND "
        )
        with sqlite3.connect(DB_PATH) as con:
            rows = con.execute(
                f"""
                SELECT i.invoice_no, i.invoice_date, i.customer_name,
                       i.grand_total, COALESCE(SUM(p.amount), 0) AS received
                FROM invoices AS i
                LEFT JOIN invoice_payments AS p
                  ON p.document_kind = i.document_kind
                 AND p.invoice_no = i.invoice_no
                WHERE i.document_kind = 'INVOICE' AND i.is_private = ?
                GROUP BY i.id
                HAVING {paid_filter}
                   (COALESCE(i.customer_name, '') LIKE ?
                    OR CAST(i.invoice_no AS TEXT) LIKE ?)
                ORDER BY i.invoice_date DESC, i.invoice_no DESC
                """,
                (
                    int(self.pending_invoice_private),
                    search_pattern,
                    search_pattern,
                ),
            ).fetchall()

        page_count = max(
            1,
            (len(rows) + self.pending_invoice_page_size - 1)
            // self.pending_invoice_page_size,
        )
        self.pending_invoice_page = min(
            max(0, self.pending_invoice_page), page_count - 1
        )
        start = self.pending_invoice_page * self.pending_invoice_page_size
        page_rows = rows[start:start + self.pending_invoice_page_size]
        self.pending_invoices_table.setRowCount(len(page_rows))
        for row_index, (invoice_no, invoice_date, customer, total, received) in enumerate(page_rows):
            balance_due = max(Decimal("0.00"), money(Decimal(str(total or 0)) - Decimal(str(received or 0))))
            values = (
                str(invoice_no),
                QDate.fromString(invoice_date, "yyyy-MM-dd").toString("dd/MM/yyyy"),
                customer or "Unnamed buyer",
                f"₹{money(total or 0):,.2f}",
                f"₹{money(received or 0):,.2f}",
                f"₹{balance_due:,.2f}",
            )
            for column, value in enumerate(values):
                self.pending_invoices_table.setItem(
                    row_index, column, QTableWidgetItem(value)
                )

        self.pending_page_label.setText(
            f"Page {self.pending_invoice_page + 1} of {page_count} · "
            f"{len(rows)} {'private ' if self.pending_invoice_private else ''}invoices"
        )
        self.pending_previous_button.setEnabled(self.pending_invoice_page > 0)
        self.pending_next_button.setEnabled(self.pending_invoice_page + 1 < page_count)
        self.open_pending_invoice_button.setEnabled(bool(page_rows))


    def open_selected_pending_invoice(self):

        row = self.pending_invoices_table.currentRow()
        if row < 0:
            return

        invoice_number = self.pending_invoices_table.item(row, 0)
        if not invoice_number:
            return

        self.document_kind.setCurrentIndex(0)
        self.invoice_no.setText(invoice_number.text())
        self.load_invoice()
        self.main_tabs.setCurrentWidget(self.invoice_tab)


    def create_quotation_register_tab(self):

        self.quotation_register_tab = QWidget()
        layout = QVBoxLayout(self.quotation_register_tab)
        toolbar = QHBoxLayout()
        self.quotation_search = QLineEdit()
        self.quotation_search.setPlaceholderText(
            "Search customer or quotation number"
        )
        self.quotation_search.textChanged.connect(
            lambda value: self.reset_quotation_page()
        )
        toolbar.addWidget(self.quotation_search, 1)
        self.open_quotation_button = QPushButton("Open Selected Quotation")
        self.open_quotation_button.clicked.connect(self.open_selected_quotation)
        toolbar.addWidget(self.open_quotation_button)
        self.convert_selected_quotation_button = QPushButton(
            "Order Confirmed · Add to Invoice"
        )
        self.convert_selected_quotation_button.clicked.connect(
            self.convert_selected_quotation_to_invoice
        )
        self.convert_selected_quotation_button.setEnabled(False)
        toolbar.addWidget(self.convert_selected_quotation_button)
        layout.addLayout(toolbar)

        self.quotation_table = QTableWidget(0, 4)
        self.quotation_table.setHorizontalHeaderLabels(
            ["Quotation No", "Date", "Customer", "Total"]
        )
        self.quotation_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.quotation_table.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.quotation_table.setSelectionMode(
            QTableWidget.SelectionMode.SingleSelection
        )
        self.quotation_table.itemSelectionChanged.connect(
            self.update_quotation_actions
        )
        self.quotation_table.verticalHeader().setVisible(False)
        self.quotation_table.horizontalHeader().setSectionResizeMode(
            2, QHeaderView.ResizeMode.Stretch
        )
        self.quotation_table.cellDoubleClicked.connect(
            lambda row, column: self.open_selected_quotation()
        )
        layout.addWidget(self.quotation_table, 1)

        pagination = QHBoxLayout()
        self.quotation_previous_button = QPushButton("Previous")
        self.quotation_previous_button.clicked.connect(
            lambda checked=False: self.change_quotation_page(-1)
        )
        self.quotation_page_label = QLabel("Page 1 of 1")
        self.quotation_page_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.quotation_next_button = QPushButton("Next")
        self.quotation_next_button.clicked.connect(
            lambda checked=False: self.change_quotation_page(1)
        )
        pagination.addWidget(self.quotation_previous_button)
        pagination.addWidget(self.quotation_page_label, 1)
        pagination.addWidget(self.quotation_next_button)
        layout.addLayout(pagination)

        self.quotation_page = 0
        self.quotation_page_size = 20
        self.refresh_quotation_register()


    def reset_quotation_page(self):

        self.quotation_page = 0
        self.refresh_quotation_register()


    def change_quotation_page(self, offset):

        self.quotation_page += offset
        self.refresh_quotation_register()


    def refresh_quotation_register(self):

        if not hasattr(self, "quotation_table"):
            return

        search_pattern = f"%{self.quotation_search.text().strip()}%"
        with sqlite3.connect(DB_PATH) as con:
            rows = con.execute(
                """
                SELECT invoice_no, invoice_date, customer_name, grand_total
                FROM invoices
                WHERE document_kind = 'QUOTATION'
                  AND (COALESCE(customer_name, '') LIKE ?
                       OR CAST(invoice_no AS TEXT) LIKE ?)
                ORDER BY invoice_date DESC, invoice_no DESC
                """,
                (search_pattern, search_pattern),
            ).fetchall()

        page_count = max(
            1,
            (len(rows) + self.quotation_page_size - 1)
            // self.quotation_page_size,
        )
        self.quotation_page = min(max(0, self.quotation_page), page_count - 1)
        start = self.quotation_page * self.quotation_page_size
        page_rows = rows[start:start + self.quotation_page_size]
        self.quotation_table.setRowCount(len(page_rows))
        for row_index, (number, date, customer, total) in enumerate(page_rows):
            values = (
                str(number),
                QDate.fromString(date, "yyyy-MM-dd").toString("dd/MM/yyyy"),
                customer or "Unnamed buyer",
                f"₹{money(total or 0):,.2f}",
            )
            for column, value in enumerate(values):
                self.quotation_table.setItem(
                    row_index, column, QTableWidgetItem(value)
                )

        self.quotation_page_label.setText(
            f"Page {self.quotation_page + 1} of {page_count} · {len(rows)} quotations"
        )
        self.quotation_previous_button.setEnabled(self.quotation_page > 0)
        self.quotation_next_button.setEnabled(self.quotation_page + 1 < page_count)
        self.update_quotation_actions()


    def update_quotation_actions(self):

        selected = self.quotation_table.currentRow() >= 0
        self.open_quotation_button.setEnabled(selected)
        self.convert_selected_quotation_button.setEnabled(selected)


    def open_selected_quotation(self):

        row = self.quotation_table.currentRow()
        if row < 0:
            return
        quotation_number = self.quotation_table.item(row, 0)
        if not quotation_number:
            return

        self.document_kind.setCurrentIndex(1)
        self.invoice_no.setText(quotation_number.text())
        self.load_invoice()
        self.main_tabs.setCurrentWidget(self.invoice_tab)


    def convert_selected_quotation_to_invoice(self):

        row = self.quotation_table.currentRow()
        if row < 0:
            return
        quotation_number = self.quotation_table.item(row, 0)
        if not quotation_number:
            return

        self.document_kind.setCurrentIndex(1)
        self.invoice_no.setText(quotation_number.text())
        self.load_invoice()
        if self.loaded_invoice_no != int(quotation_number.text()):
            return
        self.main_tabs.setCurrentWidget(self.invoice_tab)
        self.convert_quotation_to_invoice()


    def convert_quotation_to_invoice(self):

        if document_kind_for_selection(self.document_kind.currentIndex()) != "QUOTATION":
            return

        quotation_state = self.capture_document_draft()
        if not quotation_state["items"]:
            QMessageBox.warning(
                self, "Quotation Has No Items", "Add at least one item before creating an invoice."
            )
            return

        self.save_document_draft("QUOTATION")
        self.document_kind.setCurrentIndex(0)
        quotation_state.update({
            "invoice_no": str(
                self.next_invoice_number("INVOICE", is_private=False)
            ),
            "invoice_date": QDate.currentDate().toString("yyyy-MM-dd"),
            "payment_percent": 50,
            "private_invoice": False,
        })
        self.reset_receipt_amount_autofill()
        self.loaded_invoice_no = None
        self.apply_document_draft_state(quotation_state)
        self.save_document_draft("INVOICE")
        self.refresh_invoice_payments()
        self.refresh_invoice_request_history()
        self.main_tabs.setCurrentWidget(self.invoice_tab)
        self.invoice_scroll_area.verticalScrollBar().setValue(0)
        self.statusBar().showMessage(
            f"Quotation confirmed. Invoice {self.invoice_no.text()} is ready to edit."
        )


    def capture_document_draft(self):

        return {
            "invoice_no": self.invoice_no.text(),
            "invoice_date": self.invoice_date.date().toString("yyyy-MM-dd"),
            "customer_name": self.customer_name.text(),
            "customer_address": self.customer_address.text(),
            "customer_address_2": self.customer_address_2.text(),
            "customer_mobile": self.customer_mobile.text(),
            "customer_gstin": self.customer_gstin.text(),
            "buyer_profile": self.buyer_selector.currentData() or "",
            "po_number": self.po_number.text(),
            "invoice_type": self.invoice_type.currentIndex(),
            "gst_rate": self.gst_rate.value(),
            "items": [
                {
                    "description": self.cell(row, 1),
                    "hsn": self.cell(row, 2),
                    "quantity": self.cell(row, 3),
                    "rate": self.cell(row, 4),
                    "per": self.cell(row, 5),
                }
                for row in range(self.items.rowCount())
                if self.cell(row, 1)
            ],
            "hsn_code": self.hsn_code.text(),
            "bundle_count": self.bundle_count.text(),
            "through": self.through.text(),
            "bank_name": self.bank_name.text(),
            "bank_account": self.bank_account.text(),
            "bank_ifsc": self.bank_ifsc.text(),
            "bank_upi_id": self.bank_upi_id.text(),
            "bank_profile": self.bank_selector.currentData() or "",
            "payment_percent": self.payment_percent.value(),
            "payment_request_amount": self.payment_request_amount.value(),
            "declaration": self.declaration.toPlainText(),
            "seller_signatory": self.seller_signatory.text(),
            "jurisdiction": self.jurisdiction.text(),
            "private_invoice": self.private_invoice_checkbox.isChecked(),
            "custom_output_dir": str(self.custom_output_dir or ""),
        }


    def save_document_draft(self, document_kind):

        state_json = json.dumps(self.capture_document_draft())
        with sqlite3.connect(DB_PATH) as con:
            con.execute(
                """
                INSERT INTO document_drafts (document_kind, state_json)
                VALUES (?, ?)
                ON CONFLICT(document_kind) DO UPDATE SET
                    state_json = excluded.state_json
                """,
                (document_kind, state_json),
            )


    def delete_document_draft(self, document_kind):

        with sqlite3.connect(DB_PATH) as con:
            con.execute(
                "DELETE FROM document_drafts WHERE document_kind = ?",
                (document_kind,),
            )


    def restore_document_draft(self, document_kind):

        with sqlite3.connect(DB_PATH) as con:
            row = con.execute(
                "SELECT state_json FROM document_drafts WHERE document_kind = ?",
                (document_kind,),
            ).fetchone()
        if not row:
            return False

        state = json.loads(row[0])
        return self.apply_document_draft_state(state)


    def apply_document_draft_state(self, state):

        self.invoice_no.setText(state.get("invoice_no", ""))
        self.invoice_date.setDate(
            QDate.fromString(state.get("invoice_date", ""), "yyyy-MM-dd")
        )
        for key, field in (
            ("customer_name", self.customer_name),
            ("customer_address", self.customer_address),
            ("customer_address_2", self.customer_address_2),
            ("customer_mobile", self.customer_mobile),
            ("customer_gstin", self.customer_gstin),
            ("po_number", self.po_number),
            ("hsn_code", self.hsn_code),
            ("bundle_count", self.bundle_count),
            ("through", self.through),
            ("bank_name", self.bank_name),
            ("bank_account", self.bank_account),
            ("bank_ifsc", self.bank_ifsc),
            ("bank_upi_id", self.bank_upi_id),
            ("seller_signatory", self.seller_signatory),
            ("jurisdiction", self.jurisdiction),
        ):
            field.setText(state.get(key, ""))
        self.buyer_selector.blockSignals(True)
        self.buyer_selector.setCurrentIndex(
            max(0, self.buyer_selector.findData(state.get("buyer_profile", "")))
        )
        self.buyer_selector.blockSignals(False)
        self.bank_selector.blockSignals(True)
        self.bank_selector.setCurrentIndex(
            max(0, self.bank_selector.findData(state.get("bank_profile", "")))
        )
        self.bank_selector.blockSignals(False)
        self.invoice_type.setCurrentIndex(state.get("invoice_type", 0))
        self.gst_rate.setValue(state.get("gst_rate", 5))
        self.payment_percent.setValue(state.get("payment_percent", 50))
        if "payment_request_amount" in state:
            self.payment_request_amount.setValue(
                state["payment_request_amount"]
            )
        self.declaration.setPlainText(state.get("declaration", ""))
        self.private_invoice_checkbox.blockSignals(True)
        self.private_invoice_checkbox.setChecked(
            state.get("private_invoice", False)
        )
        self.private_invoice_checkbox.blockSignals(False)
        custom_output_dir = state.get("custom_output_dir", "")
        self.custom_output_dir = Path(custom_output_dir) if custom_output_dir else None

        self.items.setRowCount(0)
        for item_data in state.get("items", []):
            self.add_item_row()
            row_index = self.items.rowCount() - 1
            for column, key in (
                (1, "description"), (2, "hsn"), (3, "quantity"),
                (4, "rate"), (5, "per"),
            ):
                self.items.item(row_index, column).setText(item_data.get(key, ""))
        self.refresh()
        self.refresh_invoice_request_history()
        return True


    def refresh_invoice_history(self):

        if not hasattr(self, "invoice_request_history_table"):
            return

        self.prune_missing_invoice_records()
        self.refresh_invoice_request_history()


    def refresh_invoice_request_history(self):

        if not hasattr(self, "invoice_request_history_table"):
            return

        self.invoice_request_history_table.setRowCount(0)
        is_invoice = (
            document_kind_for_selection(self.document_kind.currentIndex()) == "INVOICE"
        )
        self.invoice_request_history_box.setVisible(is_invoice)
        if not is_invoice:
            self.update_invoice_request_actions()
            return

        try:
            invoice_no = int(self.invoice_no.text())
        except ValueError:
            self.update_invoice_request_actions()
            return

        with sqlite3.connect(DB_PATH) as con:
            requests = con.execute(
                """
                  SELECT id, requested_at, request_percent, request_amount,
                      pdf_path, payment_update_path, final_pdf_path
                FROM invoice_requests
                WHERE invoice_no = ?
                ORDER BY requested_at DESC, id DESC
                """
            , (invoice_no,)).fetchall()
            invoice = con.execute(
                "SELECT grand_total, pdf_path FROM invoices "
                "WHERE invoice_no = ? AND document_kind = 'INVOICE'",
                (invoice_no,),
            ).fetchone()
            paid_total = con.execute(
                "SELECT COALESCE(SUM(amount), 0) FROM invoice_payments "
                "WHERE invoice_no = ? AND document_kind = 'INVOICE'",
                (invoice_no,),
            ).fetchone()[0]

        received_percent = (
            money(money(paid_total) / money(invoice[0]) * 100)
            if invoice and money(invoice[0] or 0) > 0 else Decimal("0")
        )

        history_rows = list(requests)
        if invoice and invoice[1] and Path(invoice[1]).is_file():
            saved_pdf = Path(invoice[1])
            saved_at = datetime.fromtimestamp(
                saved_pdf.stat().st_mtime
            ).isoformat(timespec="seconds")
            history_rows.append((None, saved_at, None, None, str(saved_pdf), "", ""))
        history_rows.sort(
            key=lambda history_row: (
                history_row[1], history_row[0] or 0
            ),
            reverse=True,
        )

        for row, (
            request_id, requested_at, percent, amount, original_pdf_path,
            payment_update_path, final_pdf_path,
        ) in enumerate(history_rows):
            requested_time = datetime.fromisoformat(requested_at)
            request_status = f"{received_percent:.2f}% paid"
            display_pdf_path = payment_update_path or original_pdf_path
            values = (
                requested_time.strftime("%d/%m/%Y %H:%M:%S"),
                "Saved invoice PDF"
                if request_id is None
                else f"{percent:g}% · ₹{money(amount):,.2f}",
                request_status,
                Path(display_pdf_path).name,
            )
            self.invoice_request_history_table.insertRow(row)
            for column, value in enumerate(values):
                item = QTableWidgetItem(value)
                if column == 0:
                    item.setData(Qt.ItemDataRole.UserRole, request_id)
                if column == 2 and received_percent >= 100:
                    item.setBackground(QBrush(QColor("#dff2e3")))
                    item.setForeground(QBrush(QColor("#245c37")))
                elif column == 2 and received_percent > 0:
                    item.setBackground(QBrush(QColor("#fff0d5")))
                    item.setForeground(QBrush(QColor("#80521b")))
                self.invoice_request_history_table.setItem(row, column, item)
        self.update_invoice_request_actions()


    def update_invoice_request_actions(self):

        row = self.invoice_request_history_table.currentRow()
        selected = row >= 0
        self.open_invoice_request_button.setEnabled(selected)
        self.download_invoice_request_button.setEnabled(selected)
        self.send_paid_invoice_button.setEnabled(selected)


    def selected_invoice_request_path(self):

        row = self.invoice_request_history_table.currentRow()
        if row < 0:
            return None
        request_id = self.invoice_request_history_table.item(
            row, 0
        ).data(Qt.ItemDataRole.UserRole)
        if request_id is None:
            with sqlite3.connect(DB_PATH) as con:
                result = con.execute(
                    "SELECT pdf_path FROM invoices "
                    "WHERE invoice_no = ? AND document_kind = 'INVOICE'",
                    (int(self.invoice_no.text()),),
                ).fetchone()
            path = Path(result[0]) if result and result[0] else None
            return path if path and path.is_file() else None

        with sqlite3.connect(DB_PATH) as con:
            result = con.execute(
                """
                SELECT COALESCE(
                    NULLIF(payment_update_path, ''),
                    pdf_path
                )
                FROM invoice_requests
                """
                "WHERE id = ? AND invoice_no = ?",
                (request_id, int(self.invoice_no.text())),
            ).fetchone()
        path = Path(result[0]) if result else None
        return path if path and path.is_file() else None


    def open_selected_invoice_request(self):

        pdf_path = self.selected_invoice_request_path()
        if not pdf_path:
            QMessageBox.warning(self, "PDF Missing", "This request PDF could not be found.")
            return
        if not QDesktopServices.openUrl(QUrl.fromLocalFile(str(pdf_path.resolve()))):
            QMessageBox.warning(
                self, "PDF Could Not Be Opened", f"Open this PDF manually:\n{pdf_path}"
            )


    def download_selected_invoice_request(self):

        pdf_path = self.selected_invoice_request_path()
        if not pdf_path:
            QMessageBox.warning(self, "PDF Missing", "This request PDF could not be found.")
            return
        destination, _ = QFileDialog.getSaveFileName(
            self, "Download Request PDF", str(Path.home() / pdf_path.name),
            "PDF files (*.pdf)",
        )
        if not destination:
            return
        try:
            shutil.copy2(pdf_path, destination)
        except OSError as error:
            QMessageBox.warning(self, "Download Failed", str(error))
            return
        QMessageBox.information(self, "PDF Downloaded", destination)


    def send_selected_paid_invoice(self):

        if self.document_kind.currentIndex() != 0:
            return
        pdf_path = self.selected_invoice_request_path()
        if not pdf_path:
            QMessageBox.warning(self, "PDF Missing", "This paid invoice PDF could not be found.")
            return

        invoice_no = int(self.invoice_no.text())
        mobile = normalize_indian_mobile(self.customer_mobile.text())
        if not mobile:
            QMessageBox.warning(
                self,
                "Invalid Buyer Mobile",
                "Enter a valid Indian mobile number before sharing.",
            )
            return

        mime_data = QMimeData()
        mime_data.setUrls([QUrl.fromLocalFile(str(pdf_path.resolve()))])
        QApplication.clipboard().setMimeData(mime_data)
        opened = self.open_whatsapp_chat(mobile)
        if opened:
            QMessageBox.information(
                self,
                "Attach Invoice PDF",
                f"The PDF is copied. Paste it into the WhatsApp chat (Ctrl+V):\n\n{pdf_path.name}",
            )
        else:
            QMessageBox.warning(
                self,
                "WhatsApp Could Not Be Opened",
                f"Open WhatsApp and paste the completed invoice PDF (Ctrl+V):\n{pdf_path.name}",
            )


    def refresh_invoice_payments(self):

        if not hasattr(self, "invoice_payments_table"):
            return

        self.invoice_payments_table.setRowCount(0)
        if document_kind_for_selection(self.document_kind.currentIndex()) != "INVOICE":
            self.record_payment_button.setEnabled(False)
            self.invoice_payment_summary.setText(
                "Payments are recorded against invoices, not quotations."
            )
            self.refresh_payment_request()
            return

        try:
            invoice_no = int(self.invoice_no.text())
        except ValueError:
            self.record_payment_button.setEnabled(False)
            self.invoice_payment_summary.setText(
                "Save the invoice before recording payments."
            )
            self.refresh_payment_request()
            return

        with sqlite3.connect(DB_PATH) as con:
            invoice = con.execute(
                "SELECT grand_total FROM invoices "
                "WHERE invoice_no = ? AND document_kind = 'INVOICE'",
                (invoice_no,),
            ).fetchone()
            payments = con.execute(
                """
                SELECT id, paid_on, amount, method, reference
                FROM invoice_payments
                WHERE document_kind = 'INVOICE' AND invoice_no = ?
                ORDER BY paid_on, id
                """,
                (invoice_no,),
            ).fetchall()

        if not invoice:
            self.record_payment_button.setEnabled(False)
            self.invoice_payment_summary.setText(
                "Save the invoice before recording payments."
            )
            self.refresh_payment_request()
            return

        grand_total = money(invoice[0] or 0)
        paid_total = money(
            sum((Decimal(str(payment[2])) for payment in payments), Decimal("0"))
        )
        balance_due = max(Decimal("0.00"), money(grand_total - paid_total))
        for row, (payment_id, paid_on, amount, method, reference) in enumerate(payments):
            self.invoice_payments_table.insertRow(row)
            date_text = QDate.fromString(paid_on, "yyyy-MM-dd").toString("dd/MM/yyyy")
            for column, value in enumerate((
                date_text or paid_on,
                f"₹{money(amount):,.2f}",
                method,
                reference,
            )):
                self.invoice_payments_table.setItem(
                    row, column, QTableWidgetItem(str(value))
                )
            self.invoice_payments_table.item(row, 0).setData(
                Qt.ItemDataRole.UserRole, payment_id
            )

        self.invoice_payment_summary.setText(
            f"Invoice total: ₹{grand_total:,.2f}    "
            f"Received: ₹{paid_total:,.2f}    "
            f"Balance due: ₹{balance_due:,.2f}"
        )
        self.record_payment_button.setEnabled(balance_due > 0)
        self.update_invoice_payment_actions()
        self.adjust_payment_request_to_balance()
        self.refresh_payment_request()


    def refresh_payment_request(self, source=None):

        if not hasattr(self, "payment_request_amount"):
            return
        if self._updating_payment_request:
            return
        if source is not None:
            self._payment_request_source = source

        total = self.calculate()[3] if hasattr(self, "items") else Decimal("0.00")
        self._updating_payment_request = True
        if self._payment_request_source == "amount":
            requested_amount = money(self.payment_request_amount.value())
            request_percent = (
                requested_amount / total * 100 if total > 0 else Decimal("0")
            )
            self.payment_percent.setValue(float(request_percent))
        else:
            request_percent = Decimal(str(self.payment_percent.value()))
            requested_amount = money(total * request_percent / 100)
            self.payment_request_amount.setValue(float(requested_amount))
        self._updating_payment_request = False

        paid_total = Decimal("0.00")
        if (
            document_kind_for_selection(self.document_kind.currentIndex()) == "INVOICE"
            and self.invoice_no.text().strip().isdigit()
        ):
            paid_total = self.invoice_paid_total(int(self.invoice_no.text()))

        balance_due = max(Decimal("0.00"), money(total - paid_total))
        self._payment_request_balance_due = balance_due
        self._payment_request_exceeds_balance = requested_amount > balance_due
        if self._payment_request_exceeds_balance:
            self.payment_request_warning.setText(
                f"Request amount ₹{requested_amount:,.2f} is above the "
                f"balance due of ₹{balance_due:,.2f}. Reduce the request "
                "amount before saving or sending."
            )
        self.payment_request_warning.setVisible(
            self._payment_request_exceeds_balance
        )
        remaining_percent = (
            balance_due / total * 100 if total > 0 else Decimal("0")
        )
        self.remaining_percent_label.setText(
            f"Balance remaining: ₹{balance_due:,.2f} ({remaining_percent:.2f}%)"
        )
        current_amount = money(self.receipt_amount.value())
        if (
            not self._receipt_amount_manual_override
            or current_amount == self._auto_receipt_amount
        ):
            self._updating_receipt_amount = True
            self.receipt_amount.setValue(float(requested_amount))
            self._updating_receipt_amount = False
            self._auto_receipt_amount = requested_amount


    def adjust_payment_request_to_balance(self):

        if document_kind_for_selection(self.document_kind.currentIndex()) != "INVOICE":
            return
        invoice_no = self.invoice_no.text().strip()
        if not invoice_no.isdigit():
            return

        total = self.calculate()[3]
        paid_total = self.invoice_paid_total(int(invoice_no))
        balance_due = max(Decimal("0.00"), money(total - paid_total))
        if money(self.payment_request_amount.value()) > balance_due:
            self.payment_request_amount.setValue(float(balance_due))


    def set_payment_fields_to_balance(self):

        invoice_no = self.invoice_no.text().strip()
        if not invoice_no.isdigit():
            return
        total = self.calculate()[3]
        paid_total = self.invoice_paid_total(int(invoice_no))
        balance_due = max(Decimal("0.00"), money(total - paid_total))
        self._payment_request_source = "amount"
        self.payment_request_amount.setValue(float(balance_due))
        request_percent = (
            money(balance_due / total * 100) if total > 0 else Decimal("0.00")
        )
        self.payment_percent.setValue(float(request_percent))
        self.reset_receipt_amount_autofill()
        self._updating_receipt_amount = True
        self.receipt_amount.setValue(float(balance_due))
        self._updating_receipt_amount = False
        self._auto_receipt_amount = balance_due
        self.refresh_payment_request()


    def validate_payment_request(self):

        if document_kind_for_selection(self.document_kind.currentIndex()) != "INVOICE":
            return True
        self.refresh_payment_request()
        if not self._payment_request_exceeds_balance:
            return True

        requested_amount = money(self.payment_request_amount.value())
        QMessageBox.warning(
            self,
            "Request Amount Exceeds Balance",
            f"The request amount is ₹{requested_amount:,.2f}, but only "
            f"₹{self._payment_request_balance_due:,.2f} remains due. "
            "Reduce the request amount before saving or sending this invoice.",
        )
        self.payment_request_amount.setFocus()
        self.payment_request_amount.selectAll()
        return False


    def track_receipt_amount_override(self, value):

        if self._updating_receipt_amount:
            return
        if money(value) != self._auto_receipt_amount:
            self._receipt_amount_manual_override = True


    def reset_receipt_amount_autofill(self):

        self._receipt_amount_manual_override = False
        self._auto_receipt_amount = Decimal("0.00")
        self._updating_receipt_amount = True
        self.receipt_amount.setValue(0)
        self._updating_receipt_amount = False


    def invoice_paid_total(self, invoice_no):

        with sqlite3.connect(DB_PATH) as con:
            paid_total = con.execute(
                """
                SELECT COALESCE(SUM(amount), 0)
                FROM invoice_payments
                WHERE document_kind = 'INVOICE' AND invoice_no = ?
                """,
                (invoice_no,),
            ).fetchone()[0]
        return money(paid_total)


    def update_invoice_payment_actions(self):

        has_selection = self.invoice_payments_table.currentRow() >= 0
        self.edit_payment_button.setEnabled(has_selection)
        self.delete_payment_button.setEnabled(has_selection)


    def refresh_invoice_after_payment_change(self):

        self.set_payment_fields_to_balance()
        if not self.save_invoice(show_dialog=False, show_success=False):
            QMessageBox.warning(
                self,
                "Invoice Refresh Failed",
                "The receipt was changed, but the invoice workbook could not be refreshed.",
            )
            return False
        if not self.export_pdf(show_dialog=False):
            QMessageBox.warning(
                self,
                "PDF Refresh Failed",
                "The receipt was changed, but the invoice PDF could not be refreshed.",
            )
            return False
        if not self.sync_invoice_request_pdfs(int(self.invoice_no.text())):
            QMessageBox.warning(
                self,
                "Request PDF Refresh Failed",
                "The invoice PDF was updated, but the latest request history PDF could not be refreshed.",
            )
        self.remove_moved_private_files()
        self.refresh_invoice_payments()
        self.refresh_invoice_history()
        self.refresh_pending_invoices()
        return True


    def open_whatsapp_chat(self, mobile, message=None):

        web_url = f"https://web.whatsapp.com/send?phone={mobile}"
        app_url = f"whatsapp://send?phone={mobile}"
        if message:
            web_url += f"&text={quote(message)}"
            app_url += f"&text={quote(message)}"
        whatsapp_web_url = QUrl(web_url)
        whatsapp_url = QUrl(app_url)
        if sys.platform == "win32":
            try:
                import winreg

                with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "whatsapp") as protocol_key:
                    winreg.QueryValueEx(protocol_key, "URL Protocol")
            except (ImportError, OSError):
                return QDesktopServices.openUrl(whatsapp_web_url)
            if QDesktopServices.openUrl(whatsapp_url):
                return True
            return QDesktopServices.openUrl(whatsapp_web_url)
        if sys.platform.startswith("linux"):
            return QDesktopServices.openUrl(whatsapp_web_url)
        if QDesktopServices.openUrl(whatsapp_url):
            return True
        return QDesktopServices.openUrl(whatsapp_web_url)


    def export_xlsx_to_pdf(self, xlsx_path, pdf_path):

        xlsx_path = Path(xlsx_path)
        pdf_path = Path(pdf_path)
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        if sys.platform == "win32":
            powershell = shutil.which("powershell.exe")
            if not powershell:
                raise RuntimeError("Windows PowerShell was not found.")
            workbook_path = str(xlsx_path.resolve()).replace("'", "''")
            output_path = str(pdf_path.resolve()).replace("'", "''")
            script = (
                "$excel = $null; $book = $null; "
                "try { "
                "$excel = New-Object -ComObject Excel.Application; "
                "$excel.Visible = $false; "
                f"$book = $excel.Workbooks.Open('{workbook_path}'); "
                f"$book.ExportAsFixedFormat(0, '{output_path}'); "
                "} finally { "
                "if ($book) { $book.Close($false) }; "
                "if ($excel) { $excel.Quit() } "
                "}"
            )
            result = subprocess.run(
                [powershell, "-NoProfile", "-NonInteractive", "-Command", script],
                capture_output=True,
                text=True,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            if result.returncode:
                raise RuntimeError(result.stderr or result.stdout)
        else:
            libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
            if not libreoffice:
                raise RuntimeError("LibreOffice is required to refresh request PDFs.")
            with tempfile.TemporaryDirectory(dir=DATA_DIR) as temp_dir:
                result = subprocess.run(
                    [
                        libreoffice,
                        "--headless",
                        "--convert-to", "pdf",
                        "--outdir", temp_dir,
                        str(xlsx_path),
                    ],
                    capture_output=True,
                    text=True,
                )
                converted_pdf = Path(temp_dir) / f"{xlsx_path.stem}.pdf"
                if not converted_pdf.is_file():
                    raise RuntimeError(result.stderr or result.stdout)
                shutil.copy2(converted_pdf, pdf_path)
        if not pdf_path.is_file():
            raise RuntimeError(f"PDF was not created: {pdf_path}")
        return True


    def export_xlsx_batch_to_pdf(self, jobs):

        jobs = [(Path(xlsx), Path(pdf)) for xlsx, pdf in jobs]
        if not jobs:
            return True

        if sys.platform == "win32":
            powershell = shutil.which("powershell.exe")
            if not powershell:
                raise RuntimeError("Windows PowerShell was not found.")
            job_values = ", ".join(
                "@{Workbook='" + str(xlsx.resolve()).replace("'", "''")
                + "'; Pdf='" + str(pdf.resolve()).replace("'", "''") + "'}"
                for xlsx, pdf in jobs
            )
            script = (
                "$excel = $null; "
                "try { "
                "$excel = New-Object -ComObject Excel.Application; "
                "$excel.Visible = $false; $excel.DisplayAlerts = $false; "
                f"$jobs = @({job_values}); "
                "foreach ($job in $jobs) { "
                "$book = $null; "
                "try { "
                "$book = $excel.Workbooks.Open($job.Workbook); "
                "$book.ExportAsFixedFormat(0, $job.Pdf); "
                "} finally { if ($book) { $book.Close($false) } } "
                "} "
                "} finally { if ($excel) { $excel.Quit() } }"
            )
            result = subprocess.run(
                [powershell, "-NoProfile", "-NonInteractive", "-Command", script],
                capture_output=True,
                text=True,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            if result.returncode:
                raise RuntimeError(result.stderr or result.stdout)
        else:
            libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
            if not libreoffice:
                raise RuntimeError("LibreOffice is required to refresh request PDFs.")
            output_dir = jobs[0][1].parent
            if any(pdf.parent != output_dir for _, pdf in jobs):
                raise ValueError("Batch PDF outputs must share one folder.")
            result = subprocess.run(
                [
                    libreoffice,
                    "--headless",
                    "--convert-to", "pdf",
                    "--outdir", str(output_dir),
                    *(str(xlsx) for xlsx, _ in jobs),
                ],
                capture_output=True,
                text=True,
            )
            if result.returncode:
                raise RuntimeError(result.stderr or result.stdout)

        missing_pdfs = [str(pdf) for _, pdf in jobs if not pdf.is_file()]
        if missing_pdfs:
            raise RuntimeError(
                "PDF export did not create: " + ", ".join(missing_pdfs)
            )
        return True


    def sync_invoice_request_pdfs(self, invoice_no):

        with sqlite3.connect(DB_PATH) as con:
            invoice = con.execute(
                """
                SELECT xlsx_path, pdf_path, grand_total,
                    (SELECT COALESCE(SUM(amount), 0) FROM invoice_payments
                     WHERE document_kind = 'INVOICE' AND invoice_no = invoices.invoice_no)
                FROM invoices
                WHERE invoice_no = ? AND document_kind = 'INVOICE'
                """,
                (invoice_no,),
            ).fetchone()
            requests = con.execute(
                """
                  SELECT id, request_percent, request_amount, pdf_path, qr_path,
                      final_pdf_path
                FROM invoice_requests
                WHERE invoice_no = ?
                ORDER BY requested_at DESC, id DESC
                """,
                (invoice_no,),
            ).fetchall()

        if not invoice or not requests:
            return True
        source_xlsx = Path(invoice[0]) if invoice[0] else None
        source_pdf = Path(invoice[1]) if invoice[1] else None
        if not source_xlsx or not source_xlsx.is_file() or not source_pdf or not source_pdf.is_file():
            return False

        invoice_total = money(invoice[2] or 0)
        paid_total = money(invoice[3] or 0)
        balance_due = max(Decimal("0.00"), invoice_total - paid_total)
        recognized_paid = min(invoice_total, paid_total)
        received_percent = (
            money(recognized_paid / invoice_total * 100)
            if invoice_total > 0 else Decimal("0")
        )
        updated_paths = {}

        try:
            with tempfile.TemporaryDirectory(dir=DATA_DIR) as temp_dir:
                jobs = []
                job_targets = []
                for (
                    request_id, request_percent, request_amount, request_pdf,
                    qr_path, _previous_final_pdf_path,
                ) in requests:
                    if not request_pdf:
                        continue
                    request_pdf = Path(request_pdf)
                    payment_update_path = request_pdf.with_name(
                        f"{request_pdf.stem}_PAYMENT_STATUS.pdf"
                    )
                    request_qr = (
                        Path(qr_path)
                        if qr_path and Path(qr_path).is_file() else None
                    )
                    request_xlsx = Path(temp_dir) / f"request_{request_id}.xlsx"
                    request_output = Path(temp_dir) / f"request_{request_id}.pdf"
                    shutil.copy2(source_xlsx, request_xlsx)
                    workbook = load_workbook(request_xlsx)
                    sheet = workbook.active
                    item_row_shift = invoice_item_row_shift(sheet)
                    summary_row = 44 + item_row_shift
                    payment_column = (
                        "G" if f"G{summary_row}:H{summary_row}" in {
                            str(cell_range)
                            for cell_range in sheet.merged_cells.ranges
                        } else "H"
                    )
                    sheet[f"{payment_column}{summary_row}"] = (
                        f"Request: {request_percent:g}% / "
                        f"INR {money(request_amount):.2f}"
                    )
                    sheet[f"{payment_column}{summary_row + 1}"] = (
                        f"Total paid: {received_percent:.2f}% / "
                        f"INR {recognized_paid:.2f}"
                    )
                    sheet[f"{payment_column}{summary_row + 2}"] = (
                        f"Balance due: INR {balance_due:.2f}"
                    )
                    qr_row = 43 + item_row_shift

                    qr_image_index = next(
                        (
                            index for index, image in enumerate(sheet._images)
                            if getattr(image.anchor, "_from", None)
                            and image.anchor._from.col in (5, 6)
                            and image.anchor._from.row == qr_row
                        ),
                        None,
                    )
                    if qr_image_index is not None:
                        if request_qr:
                            original_image = sheet._images[qr_image_index]
                            request_image = SpreadsheetImage(str(request_qr))
                            request_image.width = original_image.width
                            request_image.height = original_image.height
                            request_image.anchor = original_image.anchor
                            sheet._images[qr_image_index] = request_image
                        else:
                            sheet._images.pop(qr_image_index)
                    elif request_qr:
                        request_image = SpreadsheetImage(str(request_qr))
                        request_image.width = 60
                        request_image.height = 60
                        request_image.anchor = f"F{qr_row + 1}"
                        sheet.add_image(request_image)

                    workbook.save(request_xlsx)
                    workbook.close()
                    jobs.append((request_xlsx, request_output))
                    job_targets.append((request_id, request_output))
                    updated_paths[request_id] = str(payment_update_path)

                self.export_xlsx_batch_to_pdf(jobs)
                for request_id, request_output in job_targets:
                    payment_update_path = Path(updated_paths[request_id])
                    payment_update_path.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(request_output, payment_update_path)

            legacy_paid_pdf = source_pdf.with_name(
                f"Invoice_{source_pdf.stem.removeprefix('INV-')}_Paid.pdf"
            )
            if legacy_paid_pdf.is_file():
                legacy_paid_pdf.unlink()
            for request in requests:
                previous_final_pdf = Path(request[5]) if request[5] else None
                if previous_final_pdf and previous_final_pdf.is_file():
                    previous_final_pdf.unlink()

            with sqlite3.connect(DB_PATH) as con:
                for request_id, payment_update_path in updated_paths.items():
                    con.execute(
                        """
                        UPDATE invoice_requests
                        SET payment_update_path = ?, final_pdf_path = ''
                        WHERE id = ?
                        """,
                        (
                            payment_update_path,
                            request_id,
                        ),
                    )
        except Exception:
            return False
        return True


    def edit_selected_invoice_payment(self):

        row = self.invoice_payments_table.currentRow()
        if row < 0:
            return
        payment_id = self.invoice_payments_table.item(row, 0).data(
            Qt.ItemDataRole.UserRole
        )
        with sqlite3.connect(DB_PATH) as con:
            payment = con.execute(
                "SELECT paid_on, amount, method, reference "
                "FROM invoice_payments WHERE id = ? AND invoice_no = ?",
                (payment_id, int(self.invoice_no.text())),
            ).fetchone()
        if not payment:
            self.refresh_invoice_payments()
            return

        dialog = QDialog(self)
        dialog.setWindowTitle("Edit Receipt")
        form = QFormLayout(dialog)
        amount_field = NoWheelDoubleSpinBox()
        amount_field.setRange(0.01, 999999999.99)
        amount_field.setDecimals(2)
        amount_field.setPrefix("₹ ")
        amount_field.setValue(float(payment[1]))
        method_field = NoWheelComboBox()
        method_field.addItems(["Cash", "UPI", "Bank Transfer", "Cheque", "Other"])
        method_field.setCurrentText(payment[2])
        date_field = NoWheelDateEdit(
            QDate.fromString(payment[0], "yyyy-MM-dd")
        )
        date_field.setCalendarPopup(True)
        date_field.setDisplayFormat("dd/MM/yyyy")
        reference_field = QLineEdit(payment[3])
        form.addRow("Amount", amount_field)
        form.addRow("Method", method_field)
        form.addRow("Date", date_field)
        form.addRow("Reference", reference_field)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Save
            | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        form.addRow(buttons)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return

        new_amount = money(amount_field.value())
        invoice_no = int(self.invoice_no.text())
        with sqlite3.connect(DB_PATH) as con:
            invoice_total = con.execute(
                "SELECT grand_total FROM invoices "
                "WHERE invoice_no = ? AND document_kind = 'INVOICE'",
                (invoice_no,),
            ).fetchone()[0]
            others_total = con.execute(
                "SELECT COALESCE(SUM(amount), 0) FROM invoice_payments "
                "WHERE invoice_no = ? AND id != ?",
                (invoice_no, payment_id),
            ).fetchone()[0]
        allowed_amount = money(
            Decimal(str(invoice_total or 0)) - Decimal(str(others_total))
        )
        if new_amount > allowed_amount:
            QMessageBox.warning(
                self,
                "Payment Exceeds Balance",
                f"This receipt cannot exceed ₹{allowed_amount:,.2f}.",
            )
            return

        answer = QMessageBox.question(
            self,
            "Confirm Receipt Edit",
            "Apply these changes to the saved receipt?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return
        with sqlite3.connect(DB_PATH) as con:
            con.execute(
                """
                UPDATE invoice_payments
                SET paid_on = ?, amount = ?, method = ?, reference = ?
                WHERE id = ? AND invoice_no = ?
                """,
                (
                    date_field.date().toString("yyyy-MM-dd"),
                    float(new_amount),
                    method_field.currentText(),
                    reference_field.text().strip(),
                    payment_id,
                    invoice_no,
                ),
            )
        self.refresh_invoice_after_payment_change()


    def delete_selected_invoice_payment(self):

        row = self.invoice_payments_table.currentRow()
        if row < 0:
            return
        item = self.invoice_payments_table.item(row, 0)
        payment_id = item.data(Qt.ItemDataRole.UserRole)
        answer = QMessageBox.question(
            self,
            "Delete Receipt",
            "Delete this payment receipt? This cannot be undone.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer != QMessageBox.StandardButton.Yes:
            return
        confirmation = QMessageBox.question(
            self,
            "Confirm Permanent Deletion",
            "Confirm again to permanently delete this receipt.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if confirmation != QMessageBox.StandardButton.Yes:
            return
        with sqlite3.connect(DB_PATH) as con:
            con.execute(
                "DELETE FROM invoice_payments WHERE id = ? AND invoice_no = ?",
                (payment_id, int(self.invoice_no.text())),
            )
        self.refresh_invoice_after_payment_change()


    def record_invoice_payment(self):

        if document_kind_for_selection(self.document_kind.currentIndex()) != "INVOICE":
            return

        if not self.save_invoice(show_dialog=True, show_success=False):
            return

        invoice_no = int(self.invoice_no.text())
        amount = money(self.receipt_amount.value())
        if amount <= 0:
            QMessageBox.warning(
                self, "Invalid Payment", "Enter an amount greater than zero."
            )
            return

        with sqlite3.connect(DB_PATH) as con:
            invoice = con.execute(
                "SELECT grand_total FROM invoices "
                "WHERE invoice_no = ? AND document_kind = 'INVOICE'",
                (invoice_no,),
            ).fetchone()
            paid_total = con.execute(
                """
                SELECT COALESCE(SUM(amount), 0)
                FROM invoice_payments
                WHERE document_kind = 'INVOICE' AND invoice_no = ?
                """,
                (invoice_no,),
            ).fetchone()[0]
            balance_due = money(Decimal(str(invoice[0] or 0)) - Decimal(str(paid_total)))
            if amount > balance_due:
                QMessageBox.warning(
                    self,
                    "Payment Exceeds Balance",
                    f"The balance due is ₹{balance_due:,.2f}.",
                )
                return

            con.execute(
                """
                INSERT INTO invoice_payments (
                    document_kind, invoice_no, paid_on, amount, method, reference
                ) VALUES ('INVOICE', ?, ?, ?, ?, ?)
                """,
                (
                    invoice_no,
                    self.receipt_date.date().toString("yyyy-MM-dd"),
                    float(amount),
                    self.receipt_method.currentText(),
                    self.receipt_reference.text().strip(),
                ),
            )

        self.set_payment_fields_to_balance()
        if not self.save_invoice(show_dialog=False, show_success=False):
            self.refresh_invoice_payments()
            QMessageBox.warning(
                self,
                "Payment Recorded",
                "The payment was recorded, but the invoice workbook could not be refreshed.",
            )
            return
        if not self.export_pdf(show_dialog=False):
            self.refresh_invoice_payments()
            QMessageBox.warning(
                self,
                "Payment Recorded",
                "The payment and workbook were saved, but the PDF could not be refreshed.",
            )
            return
        if not self.sync_invoice_request_pdfs(invoice_no):
            QMessageBox.warning(
                self,
                "Request PDF Refresh Failed",
                "The invoice PDF was updated, but request history PDFs could not be refreshed.",
            )

        self.reset_receipt_amount_autofill()
        self.receipt_date.setDate(QDate.currentDate())
        self.receipt_reference.clear()
        self.refresh_invoice_payments()
        self.refresh_invoice_request_history()
        self.remove_moved_private_files()
        QMessageBox.information(
            self,
            "Payment Recorded",
            f"₹{amount:,.2f} received by {self.receipt_method.currentText()}.",
        )


    def document_kind_changed(self, index):

        is_invoice_selection = index == 0
        document_kind = document_kind_for_selection(index)
        if self.active_document_kind is None:
            self.active_document_kind = document_kind
        elif document_kind != self.active_document_kind:
            self.save_document_draft(self.active_document_kind)
            self.active_document_kind = document_kind
            if not self.restore_document_draft(document_kind):
                self.new_invoice()
        self.document_number_label.setText(
            "Invoice No" if is_invoice_selection else "Quotation No"
        )
        self.document_details_box.setTitle(
            "Invoice Details" if is_invoice_selection else "Quotation Details"
        )
        self.save_document_button.setText(
            "Save Invoice" if is_invoice_selection else "Save Quotation"
        )
        self.invoice_heading.setText(
            "TAX INVOICE" if is_invoice_selection else "QUOTATION"
        )
        self.new_document_button.setText(
            "+ New Invoice" if is_invoice_selection else "+ New Quotation"
        )
        self.private_invoice_checkbox.setText("Private invoice")
        self.private_invoice_checkbox.setVisible(is_invoice_selection)
        if not is_invoice_selection:
            self.private_invoice_checkbox.setChecked(False)
        self.invoice_payments_box.setVisible(is_invoice_selection)
        self.bank_box.setVisible(is_invoice_selection)
        self.invoice_request_history_box.setVisible(is_invoice_selection)
        self.attach_invoice_checkbox.setVisible(is_invoice_selection)
        self.whatsapp_button.setText(
            "Send WhatsApp" if is_invoice_selection else "Send WhatsApp"
        )
        self.whatsapp_button.setVisible(True)
        self.refresh_invoice_payments()
        self.refresh_invoice_request_history()


    def private_invoice_changed(self, is_private):

        if (
            document_kind_for_selection(self.document_kind.currentIndex()) != "INVOICE"
            or self.loaded_invoice_no is not None
        ):
            return
        self.invoice_no.setText(
            str(self.next_invoice_number("INVOICE", is_private))
        )
        self.refresh_invoice_payments()
        self.refresh_invoice_request_history()


    def choose_save_folder(self):

        folder = QFileDialog.getExistingDirectory(
            self,
            "Choose a folder for this document",
            str(self.custom_output_dir or Path.home()),
        )
        if folder:
            self.custom_output_dir = Path(folder)
            self.statusBar().showMessage(
                f"Personal save folder selected: {self.custom_output_dir}"
            )


    def share_document_whatsapp(self):

        document_name = self.document_kind.currentText()
        document_kind = document_kind_for_selection(
            self.document_kind.currentIndex()
        )
        is_invoice = document_kind == "INVOICE"
        if is_invoice and not self.validate_payment_request():
            return
        invoice_total = self.calculate()[3] if is_invoice else Decimal("0.00")
        paid_total = self.invoice_paid_total(int(self.invoice_no.text())) if is_invoice and self.invoice_no.text().isdigit() else Decimal("0.00")
        balance_due = max(Decimal("0.00"), money(invoice_total - paid_total))
        request_percent = self.payment_percent.value() if is_invoice else 0
        request_amount = (
            money(self.payment_request_amount.value())
            if is_invoice else Decimal("0.00")
        )
        attach_invoice = is_invoice and self.attach_invoice_checkbox.isChecked()
        mobile = re.sub(r"\D", "", self.customer_mobile.text())
        if len(mobile) == 10:
            mobile = "91" + mobile
        if not mobile:
            QMessageBox.warning(
                self, "Buyer Mobile Required",
                "Enter the buyer's mobile number before sharing the document.",
            )
            return

        try:
            invoice_no = int(self.invoice_no.text())
        except ValueError:
            QMessageBox.warning(
                self, f"{document_name} Number Required",
                f"Enter a valid {document_name.lower()} number before sharing.",
            )
            return
        existing_request = None
        if is_invoice:
                        with sqlite3.connect(DB_PATH) as con:
                                existing_request = con.execute(
                                        """
                                        SELECT r.id, r.pdf_path, r.qr_path
                                        FROM invoice_requests AS r
                                        WHERE r.invoice_no = ?
                                            AND (
                                                ROUND(r.paid_total_at_request, 2) = ROUND(?, 2)
                                                OR (
                                                    r.paid_total_at_request IS NULL
                                                    AND ROUND((
                                                        SELECT COALESCE(SUM(p.amount), 0)
                                                        FROM invoice_payments AS p
                                                        WHERE p.invoice_no = r.invoice_no
                                                            AND p.document_kind = 'INVOICE'
                                                            AND datetime(p.created_at) <= datetime(r.requested_at)
                                                    ), 2) = ROUND(?, 2)
                                                )
                                            )
                                        ORDER BY r.requested_at DESC, r.id DESC
                                        LIMIT 1
                                        """,
                                        (invoice_no, float(paid_total), float(paid_total)),
                                ).fetchone()
        if not self.save_document_bundle(show_dialog=False):
            QMessageBox.warning(
                self,
                f"{document_name} Not Sent",
                f"Save the {document_name.lower()} and PDF before opening WhatsApp.",
            )
            return
        if is_invoice and self.payment_percent.value() != request_percent:
            self.payment_percent.setValue(request_percent)
            if not self.save_document_bundle(show_dialog=False):
                QMessageBox.warning(
                    self,
                    f"{document_name} Not Sent",
                    "The selected payment request could not be saved to the invoice PDF.",
                )
                return
        with sqlite3.connect(DB_PATH) as con:
            row = con.execute(
                """SELECT pdf_path FROM invoices
                   WHERE invoice_no = ? AND document_kind = ?""",
                (
                    invoice_no,
                    document_kind_for_selection(self.document_kind.currentIndex()),
                ),
            ).fetchone()
        pdf_path = row[0] if row else ""
        if not pdf_path or not Path(pdf_path).is_file():
            QMessageBox.warning(
                self,
                "PDF Missing",
                "Save the document PDF before attaching it in WhatsApp.",
            )
            return
        attachment_paths = (
            [Path(pdf_path).resolve()]
            if not is_invoice or attach_invoice else []
        )
        qr_path = None
        if is_invoice and request_amount > 0 and self.bank_upi_id.text().strip():
            qr_path = self.quotation_payment_qr_path(
                invoice_no, self.invoice_date.date().toPython()
            )
            if not qr_path.is_file():
                QMessageBox.warning(
                    self,
                    "Payment QR Missing",
                    f"The invoice scanner could not be found:\n{qr_path}",
                )
                return
            attachment_paths.append(qr_path.resolve())

        if is_invoice:
            request_time = datetime.now()
            timestamp = request_time.strftime("%Y%m%d_%H%M%S_%f")
            request_folder = Path(pdf_path).parent / "Requests"
            request_folder.mkdir(parents=True, exist_ok=True)
            if existing_request:
                request_id, previous_request_pdf, previous_request_qr = existing_request
                snapshot_pdf_path = Path(previous_request_pdf)
            else:
                request_id = None
                previous_request_qr = ""
                snapshot_pdf_path = request_folder / (
                    f"{Path(pdf_path).stem}_REQUEST_{timestamp}.pdf"
                )
            shutil.copy2(pdf_path, snapshot_pdf_path)
            snapshot_qr_path = None
            if qr_path and qr_path.is_file():
                snapshot_qr_path = (
                    Path(previous_request_qr)
                    if previous_request_qr else request_folder / (
                        f"{Path(pdf_path).stem}_REQUEST_{timestamp}_QR.png"
                    )
                )
                shutil.copy2(qr_path, snapshot_qr_path)
            elif previous_request_qr:
                Path(previous_request_qr).unlink(missing_ok=True)
            with sqlite3.connect(DB_PATH) as con:
                if request_id is None:
                    con.execute(
                        """
                        INSERT INTO invoice_requests (
                            invoice_no, requested_at, request_percent, request_amount,
                            pdf_path, qr_path, paid_total_at_request
                        ) VALUES (?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            invoice_no,
                            request_time.isoformat(timespec="microseconds"),
                            request_percent,
                            float(request_amount),
                            str(snapshot_pdf_path),
                            str(snapshot_qr_path or ""),
                            float(paid_total),
                        ),
                    )
                else:
                    con.execute(
                        """
                        UPDATE invoice_requests
                        SET requested_at = ?, request_percent = ?, request_amount = ?,
                            pdf_path = ?, qr_path = ?, paid_total_at_request = ?
                        WHERE id = ?
                        """,
                        (
                            request_time.isoformat(timespec="microseconds"),
                            request_percent,
                            float(request_amount),
                            str(snapshot_pdf_path),
                            str(snapshot_qr_path or ""),
                            float(paid_total),
                            request_id,
                        ),
                    )
            if not self.sync_invoice_request_pdfs(invoice_no):
                QMessageBox.warning(
                    self,
                    "Request PDF Refresh Failed",
                    "The request was created, but its payment-history PDF could not be refreshed.",
                )
            attachment_paths = (
                [snapshot_pdf_path.resolve()] if attach_invoice else []
            )
            if snapshot_qr_path:
                attachment_paths.append(snapshot_qr_path.resolve())
            self.refresh_invoice_request_history()

        pdf_mime_data = QMimeData()
        pdf_mime_data.setUrls(
            [QUrl.fromLocalFile(str(path)) for path in attachment_paths]
        )
        QApplication.clipboard().setMimeData(pdf_mime_data)
        greeting = self.customer_name.text().strip() or "Customer"
        if not is_invoice:
            message = (
                f"Hello {greeting},\n\n"
                "Please find attached our quotation for your review. "
                "We would appreciate your feedback and look forward to the opportunity to work with you. \n\n"
                "Please feel free to reach out if you have any questions or require any changes.\n\n"
                f"Regards,\n{self.seller_name.text()}"
            )
        else:
            message_lines = [
                f"Hello {greeting},",
                self.seller_name.text(),
                "",
                f"Invoice total: INR {invoice_total:.2f}",
                f"Payment request ({request_percent:g}%): INR {request_amount:.2f}",
                f"Received: INR {paid_total:.2f}",
            ]
            if balance_due > 0:
                message_lines.append(f"Balance due: INR {balance_due:.2f}")
            else:
                message_lines.append("Payment complete")
            message = "\n".join(message_lines)
        opened = self.open_whatsapp_chat(mobile, message)
        attachment_names = " and ".join(path.name for path in attachment_paths)
        if opened:
            QMessageBox.information(
                self,
                "Attach Files",
                f"The following file(s) are copied. Paste them into the WhatsApp chat (Ctrl+V):\n\n{attachment_names}",
            )
        else:
            QMessageBox.warning(
                self,
                "WhatsApp Could Not Be Opened",
                f"The following file(s) are copied. Open WhatsApp and paste them into the chat (Ctrl+V):\n{attachment_names}",
            )


    def upload_signature(self):

        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Select signature image",
            str(Path.home()),
            "PNG images (*.png)",
        )
        if not file_path:
            return

        image_reader = QImageReader(file_path)
        image_format = bytes(image_reader.format()).lower()
        if not image_reader.canRead() or image_format != b"png":
            QMessageBox.warning(
                self,
                "Invalid Signature Image",
                "Select a valid PNG image.",
            )
            return

        try:
            shutil.copy2(file_path, SIGNATURE_PATH)
        except OSError as error:
            QMessageBox.warning(
                self,
                "Signature Could Not Be Saved",
                str(error),
            )
            return

        self.signature_image_path = SIGNATURE_PATH
        self.refresh_signature_display()


    def refresh_signature_display(self):

        if not self.signature_image_path or not self.signature_image_path.is_file():
            self.signature_preview.setPixmap(QPixmap())
            self.signature_preview.setText("No signature")
            self.signature_image_label.setText("No signature uploaded")
            self.upload_signature_button.setText("Upload signature PNG")
            return

        image = QImage(str(self.signature_image_path))
        if image.isNull():
            self.signature_preview.setPixmap(QPixmap())
            self.signature_preview.setText("Preview unavailable")
        else:
            self.signature_preview.setText("")
            self.signature_preview.setPixmap(
                QPixmap.fromImage(image).scaled(
                    self.signature_preview.size(),
                    Qt.AspectRatioMode.KeepAspectRatio,
                    Qt.TransformationMode.SmoothTransformation,
                )
            )
        self.signature_image_label.setText("Signature ready for invoices")
        self.upload_signature_button.setText("Replace signature")


    def add_item_row(self):

        if self.items.rowCount() >= MAX_INVOICE_ITEMS:
            QMessageBox.information(
                self,
                "Item Limit",
                f"An invoice or quotation can contain up to "
                f"{MAX_INVOICE_ITEMS} item rows."
            )
            return

        row = self.items.rowCount()
        self.items.insertRow(row)
        self.items.blockSignals(True)
        for column in range(self.items.columnCount()):
            item = QTableWidgetItem("")
            item.setBackground(QBrush(QColor("#ffffff")))
            item.setForeground(QBrush(QColor("#1f2937")))
            if column == 5:
                item.setText("PCS")
            if column not in (1, 2, 3, 4, 5):
                item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
            self.items.setItem(row, column, item)
        self.items.blockSignals(False)
        self.refresh_item_row_controls()
        self.refresh()


    def refresh_item_row_controls(self):

        for row in range(self.items.rowCount()):
            remove_button = QPushButton("Remove")
            remove_button.clicked.connect(
                lambda checked=False, item_row=row: self.remove_item_row(item_row)
            )
            self.items.setCellWidget(row, 7, remove_button)


    def remove_item_row(self, row=None):

        if row is None:
            row = self.items.currentRow()
        if row >= 0:
            self.items.removeRow(row)
            self.refresh_item_row_controls()
            self.refresh()


    # ========================================================
    # NEW INVOICE
    # ========================================================

    def new_invoice(self):

        document_kind = document_kind_for_selection(
            self.document_kind.currentIndex()
        )
        self.delete_document_draft(document_kind)
        self.loaded_invoice_no = None
        self.custom_output_dir = None
        self.reset_receipt_amount_autofill()
        self.private_invoice_checkbox.blockSignals(True)
        self.private_invoice_checkbox.setChecked(False)
        self.private_invoice_checkbox.blockSignals(False)

        self.invoice_no.setText(
            str(
                self.next_invoice_number(document_kind, is_private=False)
            )
        )


        self.invoice_date.setDate(
            QDate.currentDate()
        )


        self.customer_name.clear()
        self.buyer_selector.setCurrentIndex(0)

        self.customer_gstin.clear()

        self.customer_address.clear()
        self.customer_address_2.clear()

        self.customer_mobile.clear()


        self.customer_address.clear()
        self.customer_address_2.clear()
        self.customer_mobile.clear()
        self.po_number.clear()
        self.amount_words.clear()
        self.tax_words.clear()
        self.bundle_count.clear()
        self.through.clear()
        self.items.setRowCount(0)
        self.invoice_type.setCurrentIndex(0)
        self.gst_rate.setValue(5)
        self.payment_percent.setValue(50)


        self.refresh()
        self.refresh_invoice_payments()
        self.refresh_invoice_request_history()
        self.invoice_scroll_area.verticalScrollBar().setValue(0)


        self.statusBar().showMessage(
            f"New {self.document_kind.currentText().lower()} created"
        )


    # ========================================================
    # READ ITEMS
    # ========================================================

    def cell(
        self,
        row,
        column
    ):

        item = self.items.item(
            row,
            column
        )

        if item:

            return item.text().strip()

        return ""


    def decimal(
        self,
        value
    ):

        try:

            return Decimal(
                str(value or "0")
            )

        except:

            return Decimal("0")


    def get_items(self):

        items = []


        for row in range(
            self.items.rowCount()
        ):

            description = self.cell(
                row,
                1
            )


            if not description:

                continue


            quantity = self.decimal(
                self.cell(
                    row,
                    3
                )
            )


            rate = self.decimal(
                self.cell(
                    row,
                    4
                )
            )


            per = self.cell(
                row,
                5
            )

            if not per:

                per = "PCS"


            hsn = self.cell(
                row,
                2
            )


            amount = money(
                quantity * rate
            )


            items.append({

                "description":
                    description,

                "quantity":
                    quantity,

                "rate":
                    rate,

                "per":
                    per,

                "hsn":
                    hsn,

                "amount":
                    amount

            })


        return items


    # ========================================================
    # CALCULATE
    # ========================================================

    def calculate(self):

        items = self.get_items()


        subtotal = money(
            sum(
                (
                    item["amount"]
                    for item in items
                ),
                Decimal("0")
            )
        )


        gst = money(

            subtotal
            *
            Decimal(
                str(
                    self.gst_rate.value()
                )
            )
            /
            Decimal("100")

        )


        total = money(
            subtotal + gst
        )


        return (
            items,
            subtotal,
            gst,
            total
        )


    def split_gst(self, gst):

        cgst = money(gst / 2)
        sgst = money(gst - cgst)

        return cgst, sgst


    # ========================================================
    # REFRESH TOTAL
    # ========================================================

    def refresh(self):

        if not hasattr(
            self,
            "summary"
        ):

            return


        (
            items,
            subtotal,
            gst,
            total
        ) = self.calculate()

        self.items.blockSignals(True)

        for row in range(self.items.rowCount()):
            self.items.item(row, 0).setText(str(row + 1))
            amount = money(
                self.decimal(self.cell(row, 3))
                * self.decimal(self.cell(row, 4))
            ) if self.cell(row, 1) else Decimal("0.00")
            self.items.item(row, 6).setText(f"{amount:.2f}")

        self.items.blockSignals(False)


        if self.invoice_type.currentIndex() == 0:

            tax_text = (
                f"IGST "
                f"{self.gst_rate.value():g}%"
                f": ₹{gst:.2f}"
            )

        else:
            cgst, sgst = self.split_gst(gst)

            tax_text = (

                f"CGST "
                f"{self.gst_rate.value()/2:g}%"
                f": ₹{cgst:.2f}    "

                f"SGST "
                f"{self.gst_rate.value()/2:g}%"
                f": ₹{sgst:.2f}"

            )


        self.summary.setText(

            f"Taxable Value: "
            f"₹{subtotal:.2f}     "

            f"{tax_text}     "

            f"Grand Total: "
            f"₹{total:.2f}"

        )
        self.refresh_payment_request()

        self.amount_words.setText(amount_to_words(total))
        self.tax_words.setText(amount_to_words(gst))

        cgst, sgst = self.split_gst(gst)
        if self.invoice_type.currentIndex() == 0:
            tax_headers = [
                "HSN/SAC", "", "Taxable Value", "Integrated Rate",
                "Integrated Amount", "", "", "Total Tax Amount"
            ]
            tax_values = [
                self.hsn_code.text(), "", f"{subtotal:.2f}",
                f"{self.gst_rate.value():g}%", f"{gst:.2f}", "", "",
                f"{gst:.2f}"
            ]
        else:
            tax_headers = [
                "HSN/SAC", "", "Taxable Value", "Central Rate",
                "Central Amount", "State Rate", "State Amount",
                "Total Tax Amount"
            ]
            tax_values = [
                self.hsn_code.text(), "", f"{subtotal:.2f}",
                f"{self.gst_rate.value() / 2:g}%", f"{cgst:.2f}",
                f"{self.gst_rate.value() / 2:g}%", f"{sgst:.2f}",
                f"{gst:.2f}"
            ]

        self.hsn_table.setHorizontalHeaderLabels(tax_headers)
        for column, value in enumerate(tax_values):
            self.hsn_table.setItem(0, column, QTableWidgetItem(value))
        self.hsn_table.setItem(1, 0, QTableWidgetItem("TOTAL"))
        self.hsn_table.setItem(1, 2, QTableWidgetItem(f"{subtotal:.2f}"))
        for column in (4, 6, 7):
            item = self.hsn_table.item(0, column)
            if item and item.text():
                self.hsn_table.setItem(1, column, QTableWidgetItem(item.text()))


    # ========================================================
    # SAVE XLSX
    # ========================================================

    def create_invoice_workbook(
        self,
        items,
        subtotal,
        gst,
        total,
        invoice_no,
        invoice_date,
        invoice_type,
        document_kind
    ):

        workbook = load_workbook(TEMPLATE_DIR / "C GST.xlsx")
        sheet = workbook.active
        paid_total = (
            self.invoice_paid_total(invoice_no)
            if document_kind == "INVOICE" else Decimal("0.00")
        )
        balance_due = max(Decimal("0.00"), money(total - paid_total))
        received_percent = (
            money(paid_total / total * 100) if total else Decimal("0.00")
        )
        payment_request_amount = (
            money(self.payment_request_amount.value())
            if document_kind == "INVOICE" else Decimal("0.00")
        )
        authorised_signatory = sheet["F50"].value or "Authorised Signatory"
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.value = None
        sheet.title = "MAIN"
        sheet.sheet_view.showGridLines = False
        windows_pdf_layout = sys.platform == "win32"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        if windows_pdf_layout:
            sheet.page_setup.fitToWidth = 0
            sheet.page_setup.fitToHeight = 0
            sheet.page_setup.scale = 103
            sheet.page_margins.left = 0.18
            sheet.page_margins.right = 0.18
            sheet.print_options.horizontalCentered = False
        else:
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 1 if len(items) <= 15 else 0
            sheet.page_setup.scale = None
            sheet.page_margins.left = 0.45
            sheet.page_margins.right = 0.45
            sheet.print_options.horizontalCentered = True
        sheet.print_area = "A1:H51"
        sheet.page_margins.top = 0.2
        sheet.page_margins.bottom = 0.3
        sheet.print_options.verticalCentered = False
        existing_merges = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
        for merge_range in (
            "A8:H8", "A8:D8", "E8:H8",
            "A9:H9", "A9:D9", "E9:H9", "E9:F9", "G9:H9", "F9:H9",
        ):
            if merge_range in existing_merges:
                sheet.unmerge_cells(merge_range)
                existing_merges.remove(merge_range)
        for address in ("A8", "A9"):
            current_border = sheet[address].border
            sheet[address].border = Border(
                left=current_border.left,
                right=Side(),
                top=current_border.top,
                bottom=current_border.bottom,
            )
        for merge_range in (
            "A8:H8", "A9:H9",
        ):
            sheet.merge_cells(merge_range)
            existing_merges.add(merge_range)
        contact_edge = Side(style="medium")
        for row in (8, 9):
            cell = sheet.cell(row, 8)
            current_border = cell.border
            cell.border = Border(
                left=current_border.left,
                right=contact_edge,
                top=current_border.top,
                bottom=current_border.bottom,
            )
        for column in range(1, 9):
            cell = sheet.cell(9, column)
            current_border = cell.border
            cell.border = Border(
                left=current_border.left,
                right=current_border.right,
                top=current_border.top,
                bottom=contact_edge,
            )
        if document_kind == "INVOICE":
            for merge_range in (
                "D43:H43", "C43:F43", "D44:E44", "F44:H44", "C44:F44",
                "D45:E45", "F45:H45", "C45:F45", "D46:E46", "F46:H46",
                "C44:E44", "C45:E45", "C46:E46", "C46:F46", "G43:H43",
                "G44:G46", "H44:H46", "A47:C47", "F47:H47",
                "F48:H49", "F50:H50",
            ):
                if merge_range in existing_merges:
                    sheet.unmerge_cells(merge_range)
                    existing_merges.remove(merge_range)
            for merge_range in (
                "A43:B43", "A44:B44", "A45:B45", "C43:E43",
                "D44:E44", "D45:E45", "D46:E46", "F43:H43",
                "F44:F46", "G44:H44", "G45:H45", "G46:H46",
                "A47:E47", "A48:E49", "F47:H47",
                "F48:H49", "F50:H50", "A51:H51",
            ):
                sheet.merge_cells(merge_range)
                existing_merges.add(merge_range)
        if document_kind == "QUOTATION":
            for merge_range in (
                "A44:C44", "A45:C45", "D44:E44", "F44:H44",
                "D45:E45", "F45:H45", "D46:E46", "F46:H46",
                "D43:H43", "A47:C47", "A48:E49", "F47:H47",
                "F48:H49", "F50:H50", "A51:H51",
            ):
                if merge_range in existing_merges:
                    sheet.unmerge_cells(merge_range)
                    existing_merges.remove(merge_range)
            for merge_range in (
                "A43:C43", "A44:C44", "A45:C45", "D43:F43",
                "D44:F44", "D45:F45", "D46:F46", "G43:H43",
                "G44:G46",
                "A47:C47", "A48:E49", "F47:H47", "F48:H49",
                "F50:H50", "A51:H51",
            ):
                sheet.merge_cells(merge_range)
                existing_merges.add(merge_range)
        for merge_range in (
            *(f"B{row}:C{row}" for row in range(23, 26)),
            "E12:F12", "G12:H12", "E13:F13", "G13:H13",
            "E14:F14", "G14:H14",
            "A34:H34", "A35:H35", "A36:B37", "A41:H41", "A42:H42",
            "A44:C44", "A45:C45",
            ("A47:E47" if document_kind == "INVOICE" else "A47:C47"),
            "A48:E49", "A51:H51",
            "A14:D14",
        ):
            if document_kind == "INVOICE" and merge_range in (
                "A44:C44", "A45:C45"
            ):
                continue
            if document_kind == "QUOTATION" and merge_range in (
                "A47:C47", "A48:E49", "A51:H51"
            ):
                continue
            if merge_range not in existing_merges:
                sheet.merge_cells(merge_range)

        detail_row, tax_total_row, eoe_row = 38, 39, 33
        amount_label_row, amount_words_row = 34, 35
        tax_words_row = 41

        website_text = self.seller_website.text().strip()
        if website_text and not website_text.lower().startswith("website:"):
            website_text = f"Website: {website_text}"
        seller_state_text = re.sub(r"\s+", " ", self.seller_state.text().strip())
        website_phone = self.seller_phone.text().strip()
        phone_numbers = []
        for phone in re.findall(
            r"(?:\+?\s*91[\s().-]*)?[6-9](?:[\s().-]*\d){9}",
            website_phone,
        ):
            digits = re.sub(r"\D", "", phone)
            if len(digits) == 12 and digits.startswith("91"):
                digits = digits[2:]
            if len(digits) == 10:
                phone_numbers.append(f"+91 {digits}")
        if phone_numbers:
            website_phone_text = "PHONE:\n" + "\n".join(phone_numbers)
        else:
            fallback_phones = [
                part.strip()
                for part in re.split(r"[,;|&]+", website_phone)
                if part.strip()
            ]
            website_phone_text = (
                "PHONE:\n" + "\n".join(fallback_phones)
                if fallback_phones else ""
            )
        email_text = self.seller_email_phone.text().strip()
        email_match = re.search(
            r"\s+((?:\+\s*)?\d[\d\s().-]*\d)\s*$",
            email_text,
        )
        email_phone = email_match.group(1).strip() if email_match else ""
        if email_match:
            email_text = email_text[:email_match.start()].strip()
        if email_text and not email_text.lower().startswith("email:"):
            email_text = f"Email: {email_text}"
        known_phone_digits = {
            re.sub(r"\D", "", phone)[-10:]
            for phone in phone_numbers
        }
        email_phone_digits = re.sub(r"\D", "", email_phone)[-10:]
        if email_phone and email_phone_digits not in known_phone_digits:
            phone_numbers.append(email_phone)
        phone_values = (
            "\n".join(phone_numbers)
            if phone_numbers else website_phone_text.partition("\n")[2]
        )
        phone_label = "PHONE:" if phone_values else ""
        values = {
            "A1": self.invoice_heading.text(),
            "A2": self.seller_name.text(),
            "A3": self.seller_business.text(),
            "A4": self.seller_address_1.text(),
            "A5": self.seller_address_2.text(),
            "A6": f"GSTIN: {self.seller_gstin.text()}",
            "A7": seller_state_text,
            "A8": "\n".join(text for text in (website_text, email_text) if text),
            "A9": f"{phone_label} {phone_values}".strip(),
            "A10": "BUYER DETAILS",
            "E10": f"{self.document_kind.currentText().upper()} DETAILS",
            "A11": self.customer_name.text(),
            "A12": "\n".join(
                line for line in (
                    self.customer_address.text().strip(),
                    self.customer_address_2.text().strip(),
                ) if line
            ),
            "E12": f"{self.document_kind.currentText().upper()} NO",
            "G12": invoice_no,
            "E13": "DATE",
            "G13": invoice_date.strftime("%d/%m/%Y"),
            "E14": "P.O NO",
            "G14": self.po_number.text(),
            "A13": f"Mobile: {self.customer_mobile.text()}",
            "A14": f"Gst no: {self.customer_gstin.text()}",
            "A16": "SL.NO",
            "B16": "Description of Goods",
            "D16": "HSN/SAC",
            "E16": "Quantity",
            "F16": "Rate",
            "G16": "Per",
            "H16": "Amount",
            "D43": "BANK DETAILS:",
            "A44": f"NO OF BUNDLE'S : {self.bundle_count.text()}",
            "A45": f"THROUGH : {self.through.text()}",
            "D44": "Bank Name:",
            "F44": self.bank_name.text(),
            "D45": "A/C. No:",
            "F45": self.bank_account.text(),
            "D46": "Branch & IFSC Code:",
            "F46": self.bank_ifsc.text(),
            "A47": "Declaration",
            "F47": authorised_signatory,
            "A48": self.declaration.toPlainText().strip(),
            "F50": self.seller_signatory.text(),
        }
        if document_kind == "INVOICE":
            for address in (
                "F44", "F45", "D43", "D44", "D45", "F46", "D46", "F47", "F50",
                "A47", "A48", "A44", "A45",
            ):
                values.pop(address, None)
            values.update({
                "A43": "BUNDLES / THROUGH",
                "A44": f"NO OF BUNDLE'S: {self.bundle_count.text()}",
                "A45": f"THROUGH: {self.through.text()}",
                "C43": "BANK DETAILS",
                "C44": "Bank Name:",
                "D44": self.bank_name.text(),
                "C45": "A/C No:",
                "D45": self.bank_account.text(),
                "C46": "IFSC:",
                "D46": self.bank_ifsc.text(),
                "F43": "PAYMENT SUMMARY",
                "G44": (
                    "PAID IN FULL"
                    if balance_due <= 0 and payment_request_amount <= 0
                    else f"Request: {self.payment_percent.value():g}% / "
                         f"INR {payment_request_amount:.2f}"
                ),
                "G45": f"Total paid: {received_percent:.2f}% / INR {paid_total:.2f}",
                "G46": f"Balance due: INR {balance_due:.2f}",
                "A47": "Declaration",
                "A48": self.declaration.toPlainText().strip(),
                "F47": "Authorised Signatory",
                "F50": self.seller_signatory.text(),
            })
        else:
            for address in (
                "A43", "A44", "A45", "D43", "D44", "F44", "D45",
                "F45", "D46", "F46", "G43", "G44", "G45", "G46",
            ):
                values.pop(address, None)
            for row in range(43, 47):
                sheet.row_dimensions[row].hidden = True
            values.update({
                "A47": "Declaration",
                "A48": self.declaration.toPlainText().strip(),
                "F47": "Authorised Signatory",
                "F50": self.seller_signatory.text(),
            })
        for address, value in values.items():
            sheet[address] = value
        sheet.row_dimensions[15].hidden = True
        for address in ("A2", "A3"):
            sheet[address].alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for address, row in (("A4", 4), ("A5", 5)):
            sheet[address].alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for address, row in (("A6", 6), ("A7", 7)):
            sheet[address].alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
            )
            address_text = str(sheet[address].value or "")
            if address_text:
                line_capacity = max(
                    32,
                    int(sum(
                        sheet.column_dimensions[column].width or 8.43
                        for column in "ABCDEFGH"
                    ) * 0.85),
                )
                wrapped_lines = max(
                    1,
                    len(textwrap.wrap(
                        address_text,
                        width=line_capacity,
                        break_long_words=True,
                        break_on_hyphens=True,
                    )),
                )
                sheet.row_dimensions[row].height = max(36, wrapped_lines * 24)
        for row in (3, 6, 7):
            sheet.row_dimensions[row].height = max(
                sheet.row_dimensions[row].height or 15, 26
            )
        for row in range(3, 8):
            current_height = sheet.row_dimensions[row].height or 15
            sheet.row_dimensions[row].height = round(current_height * 1.1, 1)
        sheet.row_dimensions[6].height = 22
        sheet.row_dimensions[7].height = 22
        sheet.row_dimensions[8].height = max(
            sheet.row_dimensions[8].height or 15, 34
        )
        contact_width = max(
            24,
            int(sum(
                sheet.column_dimensions[column].width or 8.43
                for column in "ABCDEFGH"
            ) * 0.85),
        )
        contact_lines = sum(
            max(1, len(textwrap.wrap(line, width=contact_width)))
            for line in phone_values.splitlines()
        ) or 1
        sheet.row_dimensions[9].height = max(24, contact_lines * 16 + 10)
        seller_contact_lines = sum(
            max(
                1,
                len(textwrap.wrap(
                    line,
                    width=contact_width,
                    break_long_words=True,
                    break_on_hyphens=True,
                )),
            )
            for line in str(sheet["A8"].value or "").splitlines()
        ) or 1
        sheet.row_dimensions[8].height = max(24, seller_contact_lines * 18 + 8)
        if document_kind == "INVOICE":
            for cell_range in (
                "A43:B43", "C43:E43", "F43:H43",
                "A44:B44", "A45:B45", "C44:E44", "C45:E45",
                "C46:E46", "G44:H44", "G45:H45", "G46:H46",
                "A47:C47", "A48:E49",
            ):
                for row_cells in sheet[cell_range]:
                    for cell in row_cells:
                        cell.alignment = Alignment(
                            horizontal="left", vertical="center", wrap_text=True
                        )
            for header in ("C43", "F43"):
                sheet[header].font = Font(
                    name="Arial", size=12, bold=True, color="FF000000"
                )
                sheet[header].alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            payment_summary_font = Font(
                name="Arial", size=10, bold=False, color="FF000000"
            )
            for address in ("G44", "G45", "G46"):
                sheet[address].font = payment_summary_font
            bank_details_font = Font(
                name="Arial", size=10, bold=True, color="FF000000"
            )
            for address in ("C44", "C45", "C46"):
                sheet[address].font = bank_details_font
            bank_value_font = Font(
                name="Arial", size=11, bold=True, color="FF000000"
            )
            for address in ("D44", "D45", "D46"):
                sheet[address].font = bank_value_font
            for address in ("F47", "F50"):
                sheet[address].alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            for row in range(44, 47):
                sheet.row_dimensions[row].height = 32
            sheet.column_dimensions["F"].width = 16
            sheet.column_dimensions["G"].width = 16.2
            sheet.column_dimensions["H"].width = 30.3
            if windows_pdf_layout:
                sheet.column_dimensions["F"].width = 17.0
                sheet.column_dimensions["G"].width = 17.5
                sheet.column_dimensions["H"].width = 32.0
            sheet.column_dimensions["B"].width = 46
            sheet.column_dimensions["C"].width = 20.99
            sheet.row_dimensions[47].height = 20
            sheet.row_dimensions[48].height = 40
            sheet.row_dimensions[49].height = 40
            sheet.row_dimensions[51].height = 21
            footer_side = Side(style="thin")
            footer_border = Border(
                left=footer_side,
                right=footer_side,
                top=footer_side,
                bottom=footer_side,
            )
            signature_area_border = Border(
                left=Side(),
                right=Side(),
                top=Side(),
                bottom=Side(),
            )
            for address in (
                "A43", "C43", "F43", "A44", "A45", "C44", "C45", "C46",
                "G44", "G45", "G46", "A47", "A48",
            ):
                sheet[address].border = footer_border
            sheet["F48"].border = signature_area_border
            for cell in sheet["A47:E47"][0]:
                cell.border = Border(
                    left=cell.border.left,
                    right=cell.border.right,
                    top=cell.border.top,
                    bottom=Side(),
                )
            for row_cells in sheet["A48:E49"]:
                for cell in row_cells:
                    cell.border = Border(
                        left=cell.border.left,
                        right=Side(),
                        top=Side(),
                        bottom=Side(),
                    )
            for row in range(47, 51):
                cell = sheet.cell(row, 6)
                cell.border = Border(
                    left=Side(),
                    right=cell.border.right,
                    top=cell.border.top,
                    bottom=cell.border.bottom,
                )
            for row in range(47, 51):
                cell = sheet.cell(row, 8)
                cell.border = Border(
                    left=cell.border.left,
                    right=footer_side,
                    top=cell.border.top,
                    bottom=cell.border.bottom,
                )
            for column in range(1, 6):
                cell = sheet.cell(50, column)
                cell.border = Border(
                    left=cell.border.left,
                    right=cell.border.right,
                    top=cell.border.top,
                    bottom=footer_side,
                )
            for row in range(47, 51):
                for column in range(4, 7):
                    cell = sheet.cell(row, column)
                    cell.border = Border(
                        left=Side(),
                        right=Side(),
                        top=cell.border.top,
                        bottom=cell.border.bottom,
                    )
            for column in range(1, 9):
                sheet.cell(51, column).border = Border(
                    left=Side(), right=Side(), top=Side(), bottom=Side()
                )
            for row in range(43, 47):
                for column in range(1, 9):
                    sheet.cell(row, column).border = footer_border
            for column in range(1, 3):
                sheet.cell(46, column).border = Border(
                    left=Side(),
                    right=Side(),
                    top=Side(),
                    bottom=Side(),
                )
            sheet["A2"].alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        sheet["A12"].alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True, indent=1
        )
        for address in ("A11", "A13", "A14"):
            sheet[address].alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True,
                indent=1,
            )
        address_text = str(sheet["A12"].value or "")
        if address_text:
            merged_width = sum(
                sheet.column_dimensions[column].width or 8.43
                for column in "ABCD"
            )
            line_capacity = max(20, int(merged_width * 0.9))
            wrapped_address_lines = sum(
                max(
                    1,
                    len(textwrap.wrap(
                        line,
                        width=line_capacity,
                        break_long_words=True,
                        break_on_hyphens=True,
                    )),
                )
                for line in address_text.splitlines()
            )
            sheet.row_dimensions[12].height = max(
                30, wrapped_address_lines * 15 + 6
            )
        for row in range(10, 16):
            current_height = sheet.row_dimensions[row].height or 15
            sheet.row_dimensions[row].height = round(current_height, 1)
        if windows_pdf_layout:
            for row in (12, 13, 14):
                sheet.row_dimensions[row].height = max(
                    sheet.row_dimensions[row].height or 15,
                    30,
                )
        sheet[f"A{amount_label_row}"] = (
            f"Amount Chargeable (in words): {amount_to_words(total)}"
        )
        sheet[f"A{amount_words_row}"] = None
        sheet[f"A{amount_label_row}"].alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=False,
            shrink_to_fit=True,
        )
        jurisdiction_cell = "A51"
        sheet[jurisdiction_cell].alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=False
        )
        sheet.row_dimensions[51].hidden = False
        sheet[jurisdiction_cell] = self.jurisdiction.text().strip()

        item_slot_count = max(10, len(items))
        item_line_height = 19 if windows_pdf_layout else 20
        if windows_pdf_layout and len(items) <= 15:
            minimum_item_row_height = 24
        elif windows_pdf_layout:
            minimum_item_row_height = 24
        else:
            minimum_item_row_height = (
                48 if len(items) <= 15 else 90 if len(items) <= 20 else 24
            )
        base_item_heights = []
        for item in items:
            wrapped_lines = sum(
                max(1, len(textwrap.wrap(line, width=44, break_long_words=True)))
                for line in item["description"].splitlines()
            ) or 1
            base_item_heights.append(
                max(minimum_item_row_height, wrapped_lines * item_line_height)
            )
        for index, item in enumerate(items[:10]):
            row = 18 + index
            sheet.row_dimensions[row].hidden = False
            sheet.cell(row, 1, index + 1)
            sheet.cell(row, 2, item["description"])
            sheet.cell(row, 4, item["hsn"])
            sheet.cell(row, 5, Decimal(str(item["quantity"])))
            sheet.cell(row, 5).number_format = "General"
            sheet.cell(row, 6, float(item["rate"]))
            sheet.cell(row, 7, item["per"])
            sheet.cell(row, 8, float(item["amount"]))
            sheet.row_dimensions[row].height = base_item_heights[index]
        for row in range(18 + len(items), 28):
            empty_slot = row >= 18 + item_slot_count
            sheet.row_dimensions[row].hidden = windows_pdf_layout and empty_slot
            sheet.row_dimensions[row].height = (
                0 if windows_pdf_layout and empty_slot
                else minimum_item_row_height
            )
        sheet.row_dimensions[16].height = 23

        sheet["H28"] = float(subtotal)
        tax_rate = self.gst_rate.value()
        cgst, sgst = self.split_gst(gst)
        if invoice_type == "IGST":
            sheet["B29"] = "IGST"
            sheet["H29"] = float(gst)
            sheet["B32"] = "TOTAL"
            sheet["H32"] = float(total)
            sheet["A36"] = "HSN/SAC"
            sheet["C36"] = "Taxable value"
            sheet["D36"] = "Integrated Tax"
            sheet["H36"] = "Total Tax Amount"
            sheet["D37"] = "Rate"
            sheet["E37"] = "Amount"
            sheet.cell(detail_row, 1, self.hsn_code.text())
            sheet.cell(detail_row, 3, float(subtotal))
            sheet.cell(detail_row, 4, tax_rate / 100)
            sheet.cell(detail_row, 5, float(gst))
            sheet.cell(detail_row, 8, float(gst))
            sheet.cell(tax_total_row, 1, "TOTAL")
            sheet.cell(tax_total_row, 3, float(subtotal))
            sheet.cell(tax_total_row, 5, float(gst))
            sheet.cell(tax_total_row, 8, float(gst))
            sheet[f"A{tax_words_row}"] = (
                f"Tax Amount (words): {amount_to_words(gst)}"
            )
            sheet[f"A{tax_words_row + 1}"] = None
        else:
            sheet["B29"] = "CGST"
            sheet["H29"] = float(cgst)
            sheet["B30"] = "SGST"
            sheet["H30"] = float(sgst)
            sheet["B32"] = "TOTAL"
            sheet["H32"] = float(total)
            sheet["A36"] = "HSN/SAC"
            sheet["C36"] = "Taxable value"
            sheet["D36"] = "Central Tax"
            sheet["F36"] = "State Tax"
            sheet["H36"] = "Total Tax Amount"
            sheet["D37"] = "Rate"
            sheet["E37"] = "Amount"
            sheet["F37"] = "Rate"
            sheet["G37"] = "Amount"
            sheet.cell(detail_row, 1, self.hsn_code.text())
            sheet.cell(detail_row, 3, float(subtotal))
            sheet.cell(detail_row, 4, tax_rate / 200)
            sheet.cell(detail_row, 5, float(cgst))
            sheet.cell(detail_row, 6, tax_rate / 200)
            sheet.cell(detail_row, 7, float(sgst))
            sheet.cell(detail_row, 8, float(gst))
            sheet.cell(tax_total_row, 1, "TOTAL")
            sheet.cell(tax_total_row, 3, float(subtotal))
            sheet.cell(tax_total_row, 5, float(cgst))
            sheet.cell(tax_total_row, 7, float(sgst))
            sheet.cell(tax_total_row, 8, float(gst))
            sheet[f"A{tax_words_row}"] = (
                f"Tax Amount (words): {amount_to_words(gst)}"
            )
            sheet[f"A{tax_words_row + 1}"] = None

        sheet[f"D{eoe_row}"] = "E. & O.E"
        sheet["A41"].alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=False,
            shrink_to_fit=True,
        )
        sheet["A48"].alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True, indent=1
        )
        sheet.row_dimensions[34].height = 21
        sheet.row_dimensions[35].height = 4
        sheet.row_dimensions[41].height = 21
        sheet.row_dimensions[40].height = 0
        sheet.row_dimensions[40].hidden = True
        sheet.row_dimensions[42].height = 0
        sheet.row_dimensions[42].hidden = True
        sheet.row_dimensions[48].height = 40
        sheet.row_dimensions[49].height = 40
        sheet.row_dimensions[50].height = 45
        for row in (*range(1, 10), *range(28, 52)):
            if not sheet.row_dimensions[row].hidden:
                current_height = sheet.row_dimensions[row].height or 15
                sheet.row_dimensions[row].height = round(current_height * 1.15, 1)
        if document_kind == "INVOICE":
            for row in range(43, 47):
                current_height = sheet.row_dimensions[row].height or 15
                sheet.row_dimensions[row].height = round(
                    current_height * 1.15, 1
                )
        if len(items) <= 10:
            compact_heights = {
                1: 30, 4: 22, 5: 22, 6: 19, 7: 19, 8: 38,
                28: 22, 29: 18, 30: 18, 32: 24, 33: 16, 34: 21,
                35: 3, 36: 18, 37: 16, 38: 18, 39: 18, 41: 21,
                43: 20, 44: 30, 45: 30, 46: 30, 47: 20,
                48: 32, 49: 32, 50: 40, 51: 18,
            }
            for row, height in compact_heights.items():
                if not sheet.row_dimensions[row].hidden:
                    sheet.row_dimensions[row].height = height
            for row in (*range(1, 17), *range(28, 52)):
                if not sheet.row_dimensions[row].hidden:
                    sheet.row_dimensions[row].height += 6
        sheet.row_dimensions[8].height = max(
            sheet.row_dimensions[8].height or 15,
            seller_contact_lines * 18 + 12,
            38,
        )

        for row in range(18, 28):
            for column in (1, 5, 6, 7, 8):
                sheet.cell(row, column).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            sheet.cell(row, 2).alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True,
                indent=1,
            )
        for address in (
            "G12", "G13", "G14", "F47", "F50",
        ):
            sheet[address].alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        for address in ("E12", "E13", "E14"):
            sheet[address].alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True,
                indent=1,
            )
        for address in ("A8",):
            sheet[address].alignment = Alignment(
                horizontal="center", vertical="top", wrap_text=True
            )
        sheet["A9"].alignment = Alignment(
            horizontal="right", vertical="top", wrap_text=True
        )
        bank_value_addresses = (
            ("D44", "D45", "D46")
            if document_kind == "INVOICE" else ("F44", "F45", "F46")
        )
        for address in bank_value_addresses:
            sheet[address].alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True
            )
        thin_side = Side(style="thin")
        bank_cell_border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )
        medium_side = Side(style="medium")

        def set_cell_border(cell, left=None, right=None, top=None, bottom=None):
            current = cell.border
            cell.border = Border(
                left=left or current.left,
                right=right or current.right,
                top=top or current.top,
                bottom=bottom or current.bottom,
            )

        for column in range(1, 9):
            set_cell_border(sheet.cell(10, column), top=medium_side)
        for row in range(10, 16):
            set_cell_border(sheet.cell(row, 1), left=medium_side)
            set_cell_border(sheet.cell(row, 4), right=medium_side)
            set_cell_border(sheet.cell(row, 5), left=medium_side)
            set_cell_border(sheet.cell(row, 8), right=medium_side)
        set_cell_border(sheet.cell(14, 1), bottom=medium_side)
        set_cell_border(sheet.cell(14, 4), bottom=medium_side)

        for row in range(12, 15):
            for column in range(5, 9):
                set_cell_border(
                    sheet.cell(row, column),
                    left=thin_side if column in (5, 7) else None,
                    right=thin_side if column in (6, 8) else None,
                    top=thin_side,
                    bottom=thin_side,
                )
        for row in range(44, 47):
            for column in (("C", "F") if document_kind == "INVOICE" else ("D", "F")):
                sheet[f"{column}{row}"].border = bank_cell_border

        if self.signature_image_path:
            signature = SpreadsheetImage(str(self.signature_image_path))
            signature_columns = ("F", "G", "H")
            column_width_scale = 1
            signature_column_widths = [
                int(
                    (sheet.column_dimensions[column].width or 8.43)
                    * column_width_scale * 7 + 5
                )
                for column in signature_columns
            ]
            signatory_width = sum(signature_column_widths)
            signature_width = max(1, signatory_width)
            signature_height = 130 if len(items) <= 10 else 170
            scale = min(
                signature_width / signature.width,
                signature_height / signature.height,
            )
            signature.width = max(1, int(signature.width * scale))
            signature.height = max(1, int(signature.height * scale))
            horizontal_offset = max(
                0,
                (signatory_width - signature.width) // 2
                + signature_column_widths[0] // 2,
            )
            start_column = 5
            signature_area_height = sum(
                (sheet.row_dimensions[row].height or 15) * 96 / 72
                for row in (48, 49)
            )
            signature.anchor = OneCellAnchor(
                _from=AnchorMarker(
                    col=start_column,
                    row=47,
                    colOff=pixels_to_EMU(horizontal_offset),
                    rowOff=pixels_to_EMU(max(
                        0, int((signature_area_height - signature.height) / 2)
                    )),
                ),
                ext=XDRPositiveSize2D(
                    cx=pixels_to_EMU(signature.width),
                    cy=pixels_to_EMU(signature.height),
                ),
            )
            sheet.add_image(signature)

        if (
            document_kind == "INVOICE"
            and self.bank_upi_id.text().strip()
        ):
            payment_uri = self.payment_uri(
                invoice_no,
                payment_request_amount
                if payment_request_amount > 0 else None,
            )
            qr_path = self.quotation_payment_qr_path(
                invoice_no, invoice_date, total
            )
            qr_path.parent.mkdir(parents=True, exist_ok=True)
            payment_code = qrcode.QRCode(border=2)
            payment_code.add_data(payment_uri)
            payment_code.make(fit=True)
            payment_code.make_image().save(qr_path)
            payment_qr = SpreadsheetImage(str(qr_path))
            qr_cell_width = int(
                (sheet.column_dimensions["F"].width or 8.43) * 7 + 5
            )
            qr_cell_height = sum(
                (sheet.row_dimensions[row].height or 15) * 96 / 72
                for row in range(44, 47)
            )
            qr_padding = 4
            qr_max_size = min(
                qr_cell_width - 2 * qr_padding,
                int(qr_cell_height - 2 * qr_padding),
            )
            scale = min(
                qr_max_size / payment_qr.width,
                qr_max_size / payment_qr.height,
                1,
            )
            payment_qr.width = max(1, int(payment_qr.width * scale))
            payment_qr.height = max(1, int(payment_qr.height * scale))
            qr_left_offset = max(0, (qr_cell_width - payment_qr.width) // 2)
            qr_top_offset = max(
                0, int((qr_cell_height - payment_qr.height) / 2)
            )
            payment_qr.anchor = OneCellAnchor(
                _from=AnchorMarker(
                    col=5,
                    row=43,
                    colOff=pixels_to_EMU(qr_left_offset),
                    rowOff=pixels_to_EMU(qr_top_offset),
                ),
                ext=XDRPositiveSize2D(
                    cx=pixels_to_EMU(payment_qr.width),
                    cy=pixels_to_EMU(payment_qr.height),
                ),
            )
            sheet.add_image(payment_qr)

        uniform_border_side = Side(style="medium", color="FF000000")
        final_row = 55 if document_kind == "QUOTATION" else 51
        for row in sheet.iter_rows(min_row=1, max_row=final_row, min_col=1, max_col=8):
            for cell in row:
                current = cell.border
                cell.border = Border(
                    left=uniform_border_side if current.left.style else Side(),
                    right=uniform_border_side if current.right.style else Side(),
                    top=uniform_border_side if current.top.style else Side(),
                    bottom=uniform_border_side if current.bottom.style else Side(),
                )
        for row in (28, 29, 30, 31, 32, 37, 38, 39):
            for column in range(1, 9):
                sheet.cell(row, column).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
        for row in sheet.iter_rows(min_row=1, max_row=51, min_col=1, max_col=8):
            for cell in row:
                if cell.value is not None:
                    larger_font = copy(cell.font)
                    if cell.row >= 10:
                        larger_font.name = "Arial"
                        larger_font.bold = True
                        larger_font.color = "FF000000"
                    larger_font.sz = (larger_font.sz or 10) + 2
                    cell.font = larger_font
        heading_addresses = (
            "A10", "E10", "A16", "B16", "D16", "E16", "F16", "G16",
            "H16", "A43", "D33", "A47", "F47",
        )
        for address in heading_addresses:
            if sheet[address].value is not None:
                heading_font = copy(sheet[address].font)
                heading_font.bold = True
                sheet[address].font = heading_font
        for row in (36, 37):
            for cell in sheet[row]:
                if cell.value is not None:
                    heading_font = copy(cell.font)
                    heading_font.bold = True
                    cell.font = heading_font
        company_font = copy(sheet["A2"].font)
        company_font.sz = 46
        sheet["A2"].font = company_font
        sheet.row_dimensions[2].height = 62
        sheet.row_dimensions[3].height = 32
        if len(items) >= 8:
            sheet["B25"].font = copy(sheet["B24"].font)
        for row in sheet.iter_rows(
            min_row=1, max_row=final_row, min_col=1, max_col=8
        ):
            for cell in row:
                if not isinstance(cell.value, str) or not cell.value.strip():
                    continue
                alignment = copy(cell.alignment)
                if alignment.horizontal in (None, "general"):
                    alignment.horizontal = "left"
                if alignment.horizontal in ("left", "right"):
                    alignment.indent = 1
                    cell.alignment = alignment
        extra_item_count = max(0, item_slot_count - 10)
        if extra_item_count:
            shifted_merges = [
                CellRange(str(cell_range))
                for cell_range in sheet.merged_cells.ranges
                if cell_range.min_row >= 28
            ]
            for cell_range in shifted_merges:
                sheet.unmerge_cells(str(cell_range))

            shifted_dimensions = {
                int(row): copy(dimension)
                for row, dimension in sheet.row_dimensions.items()
                if int(row) >= 28
            }
            shifted_cell_styles = {
                (row, column): copy(sheet.cell(row, column)._style)
                for row in range(28, final_row + 1)
                for column in range(1, 9)
            }
            sheet.insert_rows(28, extra_item_count)
            for row in shifted_dimensions:
                del sheet.row_dimensions[row]
            for old_row, dimension in shifted_dimensions.items():
                new_row = old_row + extra_item_count
                dimension.index = new_row
                sheet.row_dimensions[new_row] = dimension

            for cell_range in shifted_merges:
                cell_range.min_row += extra_item_count
                cell_range.max_row += extra_item_count
                sheet.merge_cells(str(cell_range))

            for (old_row, column), style in shifted_cell_styles.items():
                sheet.cell(old_row + extra_item_count, column)._style = copy(style)

            if document_kind == "INVOICE":
                declaration_title_row = 47 + extra_item_count
                for cell, clear_left, clear_right in (
                    (sheet.cell(declaration_title_row, 1), False, True),
                    (sheet.cell(declaration_title_row, 5), False, True),
                    (sheet.cell(declaration_title_row, 6), True, False),
                ):
                    current_border = cell.border
                    cell.border = Border(
                        left=Side() if clear_left else current_border.left,
                        right=Side() if clear_right else current_border.right,
                        top=current_border.top,
                        bottom=current_border.bottom,
                    )
            item_row_styles = [
                copy(sheet.cell(27, column)._style)
                for column in range(1, 9)
            ]
            for item_index in range(10, item_slot_count):
                row = 18 + item_index
                for column, style in enumerate(item_row_styles, start=1):
                    sheet.cell(row, column)._style = copy(style)
                sheet.merge_cells(
                    start_row=row, start_column=2, end_row=row, end_column=3
                )
                item_alignment = copy(sheet.cell(row, 2).alignment)
                item_alignment.horizontal = "left"
                item_alignment.vertical = "center"
                item_alignment.wrap_text = True
                item_alignment.indent = 1
                sheet.cell(row, 2).alignment = item_alignment
                if item_index < len(items):
                    item = items[item_index]
                    sheet.cell(row, 1, item_index + 1)
                    sheet.cell(row, 2, item["description"])
                    sheet.cell(row, 4, item["hsn"])
                    sheet.cell(row, 5, Decimal(str(item["quantity"])))
                    sheet.cell(row, 5).number_format = "General"
                    sheet.cell(row, 6, float(item["rate"]))
                    sheet.cell(row, 7, item["per"])
                    sheet.cell(row, 8, float(item["amount"]))
                    sheet.row_dimensions[row].height = base_item_heights[item_index]
                else:
                    sheet.row_dimensions[row].height = minimum_item_row_height

            for image in sheet._images:
                if image.anchor._from.row >= 27:
                    image.anchor._from.row += extra_item_count
            sheet.print_area = f"A1:H{final_row + extra_item_count}"
            sheet.print_title_rows = "16:16"
        item_last_row = 17 + item_slot_count
        if windows_pdf_layout and len(items) <= 15:
            item_rows = range(18, 18 + item_slot_count)
            current_item_height = sum(
                sheet.row_dimensions[row].height or 15
                for row in item_rows
            )
            extra_height = max(0, 450 - current_item_height)
            empty_item_rows = range(18 + len(items), 18 + item_slot_count)
            if len(empty_item_rows):
                extra_per_empty_row = extra_height / len(empty_item_rows)
                for row in empty_item_rows:
                    sheet.row_dimensions[row].height += extra_per_empty_row
        item_bottom = Side(style="medium", color="FF000000")
        for row in range(18, item_last_row + 1):
            bottom = (
                item_bottom
                if row == item_last_row or (len(items) > 15 and row == 32)
                else None
            )
            for column in range(1, 9):
                cell = sheet.cell(row, column)
                current = cell.border
                cell.border = Border(
                    left=current.left,
                    right=current.right,
                    top=current.top,
                    bottom=bottom or current.bottom,
                )
        if len(items) > 15:
            sheet.row_breaks.append(Break(id=32))
        right_edge_rows = (
            eoe_row,
            amount_label_row,
            tax_words_row,
            *range(43, 51),
        )
        right_edge = Side(style="medium", color="FF000000")
        for row in right_edge_rows:
            cell = sheet.cell(row + extra_item_count, 8)
            current_border = cell.border
            cell.border = Border(
                left=current_border.left,
                right=right_edge,
                top=current_border.top,
                bottom=current_border.bottom,
            )
        amount_words_row = amount_label_row + extra_item_count
        amount_words_edge = Side(style="medium", color="FF000000")
        for column in range(1, 9):
            cell = sheet.cell(amount_words_row, column)
            current_border = cell.border
            cell.border = Border(
                left=(amount_words_edge if column == 1 else current_border.left),
                right=(amount_words_edge if column == 8 else current_border.right),
                top=amount_words_edge,
                bottom=amount_words_edge,
            )
        for address in ("A2",):
            cell = sheet[address]
            current_border = cell.border
            cell.border = Border(
                left=current_border.left,
                right=Side(),
                top=current_border.top,
                bottom=current_border.bottom,
            )
        return workbook

    def remove_moved_private_files(self):

        paths = self.pending_private_move_paths
        self.pending_private_move_paths = []
        for path in paths:
            try:
                Path(path).unlink(missing_ok=True)
            except OSError as error:
                QMessageBox.warning(
                    self,
                    "Private File Cleanup Failed",
                    f"The new document was saved, but this old private file could not be removed:\n{path}\n\n{error}",
                )


    def save_document_bundle(self, show_dialog=True):

        if not self.validate_payment_request():
            return False

        if not self.save_invoice(
            show_dialog=show_dialog,
            show_success=False,
        ):
            if self.private_move_cancelled:
                self.private_move_cancelled = False
            return False

        if not self.export_pdf(show_dialog=False):
            QMessageBox.warning(
                self,
                "PDF Export Failed",
                "The XLSX was saved, but the PDF could not be created. Check that LibreOffice is installed.",
            )
            return False

        self.remove_moved_private_files()

        invoice_no = int(self.invoice_no.text())
        document_kind = document_kind_for_selection(
            self.document_kind.currentIndex()
        )
        with sqlite3.connect(DB_PATH) as con:
            xlsx_path, pdf_path = con.execute(
                """SELECT xlsx_path, pdf_path FROM invoices
                   WHERE invoice_no = ? AND document_kind = ?""",
                (invoice_no, document_kind),
            ).fetchone()
        downloads_location = QStandardPaths.writableLocation(
            QStandardPaths.StandardLocation.DownloadLocation
        )
        downloads_dir = Path(downloads_location) if downloads_location else Path.home() / "Downloads"
        downloads_dir.mkdir(parents=True, exist_ok=True)
        download_copy = downloads_dir / Path(pdf_path).name
        try:
            if Path(pdf_path).resolve() != download_copy.resolve():
                shutil.copy2(pdf_path, download_copy)
        except OSError as error:
            if show_dialog:
                QMessageBox.warning(
                    self,
                    "Downloads Copy Failed",
                    f"The PDF was saved to its normal folder, but could not be copied to Downloads:\n{error}",
                )
            return True
        if show_dialog:
            QMessageBox.information(
                self,
                "Document Saved",
                f"XLSX saved to:\n{xlsx_path}\n\nPDF saved to:\n{pdf_path}"
                f"\n\nPDF copy in Downloads:\n{download_copy}",
            )
        return True


    def save_invoice(self, show_dialog=True, show_success=None):

        try:
            if show_success is None:
                show_success = show_dialog
            self.pending_private_move_paths = []
            self.private_move_cancelled = False
            invoice_no = int(self.invoice_no.text())
            invoice_date = self.invoice_date.date().toPython()
            items, subtotal, gst, total = self.calculate()
            if not items:
                if show_dialog:
                    QMessageBox.warning(self, "No Items", "Please enter at least one item.")
                return False
            if not self.validate_payment_request():
                return False

            mobile = self.customer_mobile.text().strip()
            if mobile and not normalize_indian_mobile(mobile):
                if show_dialog:
                    QMessageBox.warning(
                        self,
                        "Invalid Buyer Mobile",
                        "Enter a valid Indian mobile number with 10 digits, optionally prefixed by +91.",
                    )
                self.customer_mobile.setFocus()
                return False

            self.save_seller_profile()
            invoice_type = "IGST" if self.invoice_type.currentIndex() == 0 else "CGST + SGST"
            document_kind = document_kind_for_selection(
                self.document_kind.currentIndex()
            )
            scanner_path = None
            document_root = self.invoice_storage_root(
                document_kind, invoice_no, total
            )
            document_folder = document_folder_for_selection(
                self.document_kind.currentIndex()
            )
            workbook_root = document_root
            if IS_FROZEN and not self.custom_output_dir:
                workbook_root = DATA_DIR / "workbooks"
                if document_kind == "INVOICE" and self.private_invoice_checkbox.isChecked():
                    workbook_root /= "Private"
            year_folder = (
                workbook_root / document_folder / "XLSX" / str(invoice_date.year)
            )
            year_folder.mkdir(parents=True, exist_ok=True)
            output = year_folder / invoice_filename(
                invoice_no, self.customer_name.text(), ".xlsx", document_kind
            )

            with sqlite3.connect(DB_PATH) as con:
                existing_invoice = con.execute(
                    "SELECT xlsx_path, pdf_path, customer_name, document_kind, is_private "
                    "FROM invoices WHERE invoice_no = ? AND document_kind = ?",
                    (invoice_no, document_kind),
                ).fetchone()
            paid_total = (
                self.invoice_paid_total(invoice_no)
                if document_kind == "INVOICE" else Decimal("0.00")
            )
            if paid_total > money(total):
                if show_dialog:
                    QMessageBox.warning(
                        self,
                        "Invoice Total Below Amount Received",
                        f"The recorded payments total ₹{paid_total:,.2f}, which is "
                        f"more than the new invoice total of ₹{money(total):,.2f}.",
                    )
                return False
            if existing_invoice and existing_invoice[0]:
                previous_xlsx = Path(existing_invoice[0])
                was_private = bool(existing_invoice[4])
                will_be_private = (
                    document_kind == "INVOICE"
                    and self.private_invoice_checkbox.isChecked()
                )
                if was_private and not will_be_private:
                    answer = QMessageBox.question(
                        self,
                        "Move Private Document?",
                        f"This document is already saved in Private/{document_folder}. "
                        f"Move it to {document_folder}? The old private XLSX and PDF "
                        "will be removed after the new PDF is saved.",
                        QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                        QMessageBox.StandardButton.No,
                    )
                    if answer != QMessageBox.StandardButton.Yes:
                        self.private_move_cancelled = True
                        return False
                    self.pending_private_move_paths = [
                        existing_invoice[0],
                        existing_invoice[1],
                    ]
                    if existing_invoice[3] in ("INVOICE", "QUOTATION"):
                        private_qr_path = (
                            PRIVATE_INVOICE_DIR
                            / document_folder
                            / "QR"
                            / str(invoice_date.year)
                            / invoice_filename(
                                invoice_no,
                                existing_invoice[2],
                                "_payment.png",
                                existing_invoice[3],
                            )
                        )
                        self.pending_private_move_paths.append(private_qr_path)
            if existing_invoice and existing_invoice[0]:
                previous_xlsx = Path(existing_invoice[0])
                previous_pdf = Path(existing_invoice[1]) if existing_invoice[1] else None
                target_pdf_folder = self.invoice_pdf_folder(
                    document_kind, invoice_no, invoice_date
                )
                target_pdf = target_pdf_folder / f"{output.stem}.pdf"
                cleanup_paths = [previous_xlsx, previous_pdf]
                if previous_pdf:
                    cleanup_paths.append(
                        previous_pdf.with_name(
                            f"Invoice_{previous_pdf.stem.removeprefix('INV-')}_Paid.pdf"
                        )
                    )
                for old_path in cleanup_paths:
                    if (
                        old_path
                        and old_path.resolve() not in {
                            output.resolve(), target_pdf.resolve()
                        }
                        and old_path not in self.pending_private_move_paths
                    ):
                        self.pending_private_move_paths.append(old_path)
                if previous_xlsx.parent.parent.name == "XLSX":
                    previous_root = previous_xlsx.parent.parent.parent
                    old_qr = (
                        previous_root / "QR" / str(invoice_date.year)
                        / invoice_filename(
                            invoice_no,
                            existing_invoice[2],
                            "_payment.png",
                            document_kind,
                        )
                    )
                    new_qr = (
                        document_root / document_folder / "QR"
                        / str(invoice_date.year)
                        / invoice_filename(
                            invoice_no,
                            self.customer_name.text(),
                            "_payment.png",
                            document_kind,
                        )
                    )
                    if old_qr.resolve() != new_qr.resolve():
                        self.pending_private_move_paths.append(old_qr)
            if output.exists() and not existing_invoice:
                if show_dialog:
                    QMessageBox.warning(
                        self, "Already Exists", f"Invoice {invoice_no} already exists."
                    )
                return False

            workbook = self.create_invoice_workbook(
                items, subtotal, gst, total, invoice_no, invoice_date,
                invoice_type, document_kind,
            )
            workbook.save(output)

            with sqlite3.connect(DB_PATH) as con:
                values = (
                    invoice_date.isoformat(),
                    invoice_type,
                    self.customer_name.text(),
                    float(total),
                    str(output),
                    document_kind,
                    scanner_path,
                    self.bank_upi_id.text().strip(),
                    int(
                        document_kind == "INVOICE"
                        and self.private_invoice_checkbox.isChecked()
                    ),
                    invoice_no,
                )
                if existing_invoice:
                    con.execute(
                        """
                        UPDATE invoices
                        SET invoice_date = ?, invoice_type = ?, customer_name = ?,
                            grand_total = ?, xlsx_path = ?, document_kind = ?,
                            qr_scanner_path = ?, payment_address = ?, is_private = ?
                        WHERE invoice_no = ? AND document_kind = ?
                        """,
                        (*values, document_kind),
                    )
                else:
                    con.execute(
                        """
                        INSERT INTO invoices (
                            invoice_date, invoice_type, customer_name, grand_total,
                            xlsx_path, document_kind, qr_scanner_path, payment_address,
                            is_private, invoice_no
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        values,
                    )

            self.loaded_invoice_no = invoice_no
            self.save_document_draft(document_kind)
            self.refresh_invoice_history()
            self.refresh_pending_invoices()
            self.refresh_quotation_register()
            self.refresh_invoice_payments()
            self.statusBar().showMessage(f"{document_kind.title()} saved: {output}")
            if show_dialog and show_success:
                QMessageBox.information(
                    self,
                    f"{document_kind.title()} Saved",
                    f"{document_kind.title()} {invoice_no} saved successfully.\n\n{output}",
                )
            return True

        except sqlite3.IntegrityError:
            if show_dialog:
                QMessageBox.warning(
                    self, "Already Saved", "This invoice number already exists."
                )
            return False
        except Exception as error:
            if show_dialog:
                QMessageBox.critical(self, "Error", str(error))
            return False


    # ========================================================
    # PDF
    # ========================================================

    def export_pdf(self, show_dialog=True):

        try:

            invoice_no = int(
                self.invoice_no.text()
            )


            invoice_date = (
                self.invoice_date
                .date()
                .toPython()
            )


            xlsx_file, saved_buyer_name = self.saved_invoice_file(invoice_no)
            document_kind = document_kind_for_selection(
                self.document_kind.currentIndex()
            )
            document_folder = document_folder_for_selection(
                self.document_kind.currentIndex()
            )
            workbook_root = self.invoice_storage_root(
                document_kind, invoice_no
            )
            if IS_FROZEN and not self.custom_output_dir:
                workbook_root = DATA_DIR / "workbooks"
                if document_kind == "INVOICE" and self.private_invoice_checkbox.isChecked():
                    workbook_root /= "Private"
            default_year_folder = (
                workbook_root / document_folder / "XLSX" / str(invoice_date.year)
            )
            buyer_name = saved_buyer_name or self.customer_name.text()
            if not xlsx_file or not xlsx_file.exists():
                preferred_xlsx = default_year_folder / invoice_filename(
                    invoice_no,
                    buyer_name,
                    ".xlsx",
                    document_kind,
                )
                xlsx_file = preferred_xlsx


            # Save XLSX first
            if not xlsx_file.exists():
                if not self.save_invoice(show_dialog=show_dialog):
                    return False
                saved_xlsx, saved_buyer_name = self.saved_invoice_file(invoice_no)
                if saved_xlsx:
                    xlsx_file = saved_xlsx
                    buyer_name = saved_buyer_name or buyer_name


            if not xlsx_file.exists():
                return False


            pdf_folder = self.invoice_pdf_folder(
                document_kind, invoice_no, invoice_date
            )
            pdf_folder.mkdir(parents=True, exist_ok=True)
            pdf_file = pdf_folder / f"{xlsx_file.stem}.pdf"


            if sys.platform == "win32":
                generated_pdf = pdf_file
                powershell = shutil.which("powershell.exe")
                if not powershell:
                    raise RuntimeError(
                        "Windows PowerShell was not found. The Excel invoice "
                        f"is saved at:\n{xlsx_file}"
                    )

                workbook_path = str(xlsx_file.resolve()).replace("'", "''")
                pdf_path = str(pdf_file.resolve()).replace("'", "''")
                script = (
                    "$excel = $null; $book = $null; "
                    "try { "
                    "$excel = New-Object -ComObject Excel.Application; "
                    "$excel.Visible = $false; "
                    f"$book = $excel.Workbooks.Open('{workbook_path}'); "
                    f"$book.ExportAsFixedFormat(0, '{pdf_path}'); "
                    "} finally { "
                    "if ($book) { $book.Close($false) }; "
                    "if ($excel) { $excel.Quit() } "
                    "}"
                )
                result = subprocess.run(
                    [powershell, "-NoProfile", "-NonInteractive", "-Command", script],
                    capture_output=True,
                    text=True,
                    creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                )
                if result.returncode:
                    raise RuntimeError(
                        "Microsoft Excel could not export the invoice to PDF. "
                        "Confirm Excel is installed and activated.\n\n"
                        f"{result.stderr or result.stdout}"
                    )
            else:
                libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
                if not libreoffice:
                    if show_dialog:
                        QMessageBox.warning(
                            self,
                            "PDF Export Unavailable",
                            "LibreOffice is required to export PDF on Ubuntu. "
                            "The XLSX invoice is already saved:\n"
                            f"{xlsx_file}"
                        )
                    return False

                with tempfile.TemporaryDirectory(dir=DATA_DIR) as temp_dir:
                    result = subprocess.run(
                        [
                            libreoffice,
                            "--headless",
                            "--convert-to", "pdf",
                            "--outdir", temp_dir,
                            str(xlsx_file)
                        ],
                        capture_output=True,
                        text=True
                    )
                    temporary_pdf = Path(temp_dir) / f"{xlsx_file.stem}.pdf"
                    if temporary_pdf.exists():
                        shutil.copy2(temporary_pdf, pdf_file)
                generated_pdf = pdf_file


            if not generated_pdf.exists():

                raise RuntimeError(

                    "PDF was not created.\n\n"
                    +
                    result.stderr

                )


            # Database
            with sqlite3.connect(
                DB_PATH
            ) as con:

                con.execute(

                    """
                    UPDATE invoices

                    SET pdf_path = ?

                    WHERE invoice_no = ? AND document_kind = ?
                    """,

                    (
                        str(
                            generated_pdf
                        ),
                            invoice_no,
                            document_kind,
                    )

                )

            self.mirror_external_final_pdf(
                generated_pdf, document_kind, invoice_no, invoice_date
            )

            if document_kind == "INVOICE":
                self.refresh_invoice_request_history()


            self.statusBar().showMessage(

                f"PDF saved: "
                f"{generated_pdf.name}"

            )

            if show_dialog:
                QMessageBox.information(
                    self,
                    "PDF Saved",
                    f"PDF successfully saved to:\n\n{generated_pdf}"
                )
            return True


        except Exception as error:
            if show_dialog:
                QMessageBox.critical(self, "PDF Error", str(error))
            return False


# ============================================================
# START
# ============================================================

if __name__ == "__main__":

    app = QApplication(
        sys.argv
    )

    app.setApplicationName(
        "Invoice Manager"
    )


    window = InvoiceApp()

    window.show()


    sys.exit(
        app.exec()
    )